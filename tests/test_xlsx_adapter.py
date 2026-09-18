"""tests/test_xlsx_adapter.py — Excel 适配器重写版测试（D2 设计 §9 用例清单 T1-T25）

engine 策略：core/engine.py（I1a）合入前，用符合 D1 §2.2-2.4 契约的最小
stub 挂到 masker.engine 上（仅测试文件内，不进 src）；合入后自动改用真实
引擎。另设 attach_engine=False 用例覆盖"引擎未注入"的过渡 shim 路径
（当前生产代码在 I1a 合入前实际走的路径）。
"""

import logging
import re
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from mask_tool.adapters.xlsx_adapter import XlsxAdapter
from mask_tool.core.detector import Detector
from mask_tool.core.masker import Masker
from mask_tool.core.policy import PolicyEngine
from mask_tool.core.tokenizer import TokenGenerator
from mask_tool.models.config import MaskConfig
from mask_tool.models.detection import (
    DetectionStatus,
    DetectionType,
    Location,
)
from mask_tool.models.mapping import TokenMapping

# ---------------------------------------------------------------------------
# engine 加载：真实引擎优先，未合入则契约 stub
# ---------------------------------------------------------------------------

try:
    from mask_tool.core.engine import ReplacementEngine as _RealEngine  # noqa: F401
    _REAL_ENGINE_AVAILABLE = True
except ImportError:
    _RealEngine = None
    _REAL_ENGINE_AVAILABLE = False

TOKEN_RE = re.compile(
    r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]"
)

_ENGINE_SOURCE_NOTE: list = []


@dataclass(frozen=True)
class _StubSpan:
    start: int
    end: int
    text: str
    result: Any


@dataclass
class _StubMaskOutcome:
    text: str
    replacements: list
    dropped_overlaps: list
    new_mappings: list


class _StubReplacementEngine:
    """D1 §2.2-2.4 契约的最小实现：区间定位 -> 重叠消解（长度优先）-> 一次重建。"""

    def __init__(self, token_generator, masker, irreversible=False):
        self.token_gen = token_generator
        self.masker = masker
        self.irreversible = irreversible

    def mask_plain_text(self, text, results, statuses=None, allowed_originals=None):
        if statuses is None:
            statuses = frozenset({DetectionStatus.AUTO_MASK})
        spans = []
        for r in results:
            if r.status not in statuses:
                continue
            if allowed_originals is not None and r.text not in allowed_originals:
                continue
            if not r.text:
                continue
            cursor = 0
            while True:
                idx = text.find(r.text, cursor)
                if idx == -1:
                    break
                spans.append(_StubSpan(idx, idx + len(r.text), r.text, r))
                cursor = idx + len(r.text)
        order = sorted(
            spans,
            key=lambda s: (-(s.end - s.start), -s.result.confidence, s.start),
        )
        chosen, dropped = [], []
        for s in order:
            if any(s.start < c.end and c.start < s.end for c in chosen):
                dropped.append(s)
            else:
                chosen.append(s)
        chosen.sort(key=lambda s: s.start)
        replacements, memo, new_mappings = [], {}, []
        for s in chosen:
            if s.text in memo:
                replacements.append(memo[s.text])
                continue
            if self.irreversible:
                new, token = "***", None
            else:
                token = self.token_gen.generate(s.text, s.result.text_type)
                mapping = TokenMapping(
                    token=token,
                    original=s.text,
                    text_type=s.result.text_type,
                    confidence=s.result.confidence,
                )
                self.masker.mappings.append(mapping)
                new_mappings.append(mapping)
                new = token
            memo[s.text] = (s, new, token)
            replacements.append(memo[s.text])
        out, cursor = [], 0
        for s, new, _token in replacements:
            out.append(text[cursor:s.start])
            out.append(new)
            cursor = s.end
        out.append(text[cursor:])
        return _StubMaskOutcome("".join(out), replacements, dropped, new_mappings)


def _attach_engine(token_gen, masker, irreversible) -> str:
    """按 D1 契约为 masker 挂引擎：真实引擎优先，未合入则用契约 stub。"""
    if getattr(masker, "engine", None) is not None:
        return "masker(built-in)"
    if _REAL_ENGINE_AVAILABLE:
        try:
            masker.engine = _RealEngine(token_gen, irreversible=irreversible)
            return "real"
        except TypeError:
            pass  # 签名暂不匹配：退回 stub，集成期再对齐
    masker.engine = _StubReplacementEngine(token_gen, masker, irreversible)
    return "stub"


DEFAULT_LEXICON = {
    "company": ["某某建设集团有限公司", "公司"],
    "person": ["张三", "李四"],
}


def _make_adapter(mode="smart", irreversible=False, lexicon=None,
                  attach_engine=True):
    detector = Detector(
        {k: list(v) for k, v in (lexicon or DEFAULT_LEXICON).items()}, set(),
    )
    policy = PolicyEngine(MaskConfig(mode=mode))
    token_gen = TokenGenerator()
    masker = Masker(token_gen, irreversible=irreversible)
    source = "shim"
    if attach_engine:
        source = _attach_engine(token_gen, masker, irreversible)
    _ENGINE_SOURCE_NOTE.append(source)
    return XlsxAdapter(detector, policy, masker)


def _norm_tokens(text: str) -> str:
    """把 token 编号归一化为 _N，屏蔽引擎间编号顺序差异。"""
    return re.sub(TOKEN_RE, lambda m: m.group(0).split("_")[0] + "_N]", text)


@pytest.fixture(scope="session", autouse=True)
def _engine_source_report():
    yield
    sources = ", ".join(sorted(set(_ENGINE_SOURCE_NOTE))) or "(none)"
    print(f"\n[xlsx tests] engine path: {sources} (real available="
          f"{_REAL_ENGINE_AVAILABLE})")


# ---------------------------------------------------------------------------
# 手写最小 xlsx（用于公式缓存 / 19 位整数精度用例）
# ---------------------------------------------------------------------------

_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
       'package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
       'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
       '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/'
       'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
         '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
         'relationships"><Relationship Id="rId1" Type="http://schemas.'
         'openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
         'Target="xl/workbook.xml"/></Relationships>')
_WB = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
       'main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
       'relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/>'
       '</sheets></workbook>')
_WBRELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
           'relationships"><Relationship Id="rId1" Type="http://schemas.'
           'openxmlformats.org/officeDocument/2006/relationships/worksheet" '
           'Target="worksheets/sheet1.xml"/></Relationships>')


def _write_raw_xlsx(path, sheet_xml):
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _RELS)
        z.writestr("xl/workbook.xml", _WB)
        z.writestr("xl/_rels/workbook.xml.rels", _WBRELS)
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


# ---------------------------------------------------------------------------
# T1-T5 / T25：数字单元格检测（M1）
# ---------------------------------------------------------------------------

class TestNumericCells:

    def _run(self, tmp_path, value, fmt=None):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = value
        if fmt:
            ws["A1"].number_format = fmt
        src = tmp_path / "num.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        return load_workbook(out).active, adapter

    def test_t01_currency_amount_auto(self, tmp_path):
        """T1：货币格式大额（≥1万）AUTO 替换为 [AMOUNT_xxx]，number_format 保留。"""
        ws, adapter = self._run(tmp_path, 12000000, "¥#,##0.00")
        assert ws["A1"].value == "[AMOUNT_001]"
        assert ws["A1"].number_format == "¥#,##0.00"
        m = adapter.masker.mappings[0]
        assert m.original == "12000000" and m.text_type == DetectionType.AMOUNT
        assert getattr(m, "kind", "text") == "number"

    def test_t02_currency_small_noop(self, tmp_path):
        """T2：货币小额（¥5.00）不产生结果、不改动。"""
        ws, adapter = self._run(tmp_path, 5, "¥#,##0.00")
        assert ws["A1"].value == 5
        assert adapter.masker.mappings == []

    def test_t03_big_int_suggest_only(self, tmp_path):
        """T3：9 位整数 SUGGEST 仅报告不替换（H6 对齐）。"""
        ws, adapter = self._run(tmp_path, 120000000)
        assert ws["A1"].value == 120000000
        assert adapter.masker.mappings == []
        results = adapter._synth_numeric_results("120000000", "General",
                                                 Location(file="x"))
        assert len(results) == 1
        r = results[0]
        assert r.source == "xlsx_numeric"
        assert r.text_type == DetectionType.AMOUNT and r.confidence == 0.70
        adapter.policy.apply(results)
        assert results[0].status == DetectionStatus.SUGGEST_MASK

    def test_t04_card_luhn_pass_auto(self, tmp_path):
        """T4：Luhn 通过的 16 位卡号 AUTO 替换，source=xlsx_numeric。"""
        ws, adapter = self._run(tmp_path, 4111111111111111)
        assert ws["A1"].value == "[CUSTOM_001]"
        m = adapter.masker.mappings[0]
        assert m.text_type == DetectionType.CUSTOM
        assert getattr(m, "kind", "text") == "number"
        results = adapter._synth_numeric_results("4111111111111111", "General",
                                                 Location(file="x"))
        assert results[0].source == "xlsx_numeric"
        assert results[0].confidence == 0.85
        assert XlsxAdapter._luhn_ok("4111111111111111")

    def test_t05_card_luhn_fail_suggest(self, tmp_path):
        """T5：Luhn 失败的 16 位数字 SUGGEST 不替换。
        写侧保真：未脱敏的 ≥16位整数预转文本（openpyxl save 会丢 int 精度）。"""
        ws, adapter = self._run(tmp_path, 1234567890123456)
        assert ws["A1"].value == "1234567890123456"      # 转文本保真（原 int 会在 save 时丢精度）
        assert ws["A1"].number_format == "@"
        assert adapter.masker.mappings == []
        assert not XlsxAdapter._luhn_ok("1234567890123456")
        results = adapter._synth_numeric_results("1234567890123456", "General",
                                                 Location(file="x"))
        assert results[0].confidence == 0.65
        adapter.policy.apply(results)
        assert results[0].status == DetectionStatus.SUGGEST_MASK

    def test_t25_zero_values_reach_pipeline(self, tmp_path):
        """T25：cell.value==0 / 0.0 不因真值判断被跳过（N4 回归）。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 0
        ws["A2"] = 0.0
        ws["A3"] = 0
        ws["A3"].number_format = "¥#,##0.00"
        src = tmp_path / "zero.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        called = []
        original = adapter._process_numeric_cell

        def spy(cell, file_path, sheet):
            called.append(cell.coordinate)
            return original(cell, file_path, sheet)

        adapter._process_numeric_cell = spy
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert {"A1", "A2", "A3"} <= set(called)   # 流程可达
        assert ws2["A1"].value == 0 and ws2["A2"].value == 0
        assert adapter.masker.mappings == []       # 无规则命中（小额控制）


# ---------------------------------------------------------------------------
# T6 / T13 / T22：文本单元格
# ---------------------------------------------------------------------------

class TestTextCells:

    def test_t06_overlap_replacement_no_nesting(self, tmp_path):
        """T6：词典长短词重叠（公司 ⊂ 全称）区间替换不丢字、不嵌套（N1 回归）。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "合同金额12000.5万元，乙方张三，本公司已确认"
        ws["A2"] = "某某建设集团有限公司"
        src = tmp_path / "text.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        v1 = ws2["A1"].value
        # 张三 / 公司 均被替换；SUGGEST 档的金额保持原文（H6）
        assert "张三" not in v1 and "公司" not in v1
        assert "12000.5万元" in v1
        assert v1.count("[") == v1.count("]") == 2
        assert TOKEN_RE.search(v1)
        # 全称整词替换，无残片、无嵌套
        assert TOKEN_RE.fullmatch(ws2["A2"].value)
        assert "有限" not in ws2["A2"].value

    def test_t13_shared_string_isolation(self, tmp_path):
        """T13：共享字符串陷阱——同串两格各自替换、互不破坏，非敏感格不动。"""
        wb = Workbook()
        ws = wb.active
        ws["B1"] = "张三"
        ws["B2"] = "张三"
        ws["C1"] = "普通文本"
        src = tmp_path / "shared.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["B1"].value == "[PERSON_001]"
        assert ws2["B2"].value == "[PERSON_001]"    # 同原文复用同一 token
        assert ws2["C1"].value == "普通文本"
        person_mappings = [m for m in adapter.masker.mappings
                           if m.original == "张三"]
        assert len(person_mappings) == 1

    def test_t22_datetime_bool_error_untouched(self, tmp_path):
        """T22：datetime / bool / error 单元格完全不受影响。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = datetime(2024, 3, 15)
        ws["A2"] = True
        ws["A3"] = "#N/A"
        ws["A4"] = "#DIV/0!"
        src = tmp_path / "misc.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        for coord, expect in (("A1", datetime(2024, 3, 15)), ("A2", True),
                              ("A3", "#N/A"), ("A4", "#DIV/0!")):
            assert ws2[coord].value == expect, coord
        assert ws2["A1"].data_type == "d" and ws2["A1"].is_date
        assert ws2["A2"].data_type == "b"
        assert ws2["A3"].data_type == "e"
        assert adapter.masker.mappings == []


# ---------------------------------------------------------------------------
# T7-T9：公式单元格
# ---------------------------------------------------------------------------

class TestFormulaCells:

    def test_t07_formula_literal_masked(self, tmp_path):
        """T7：字面量内敏感词替换且公式语法保持。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = '="客户张三"'
        ws["A2"] = '=CONCAT("李四","与张三")'
        src = tmp_path / "f.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["A1"].value == '="客户[PERSON_001]"'
        assert ws2["A1"].data_type == "f"
        assert ws2["A1"].value.startswith("=")
        v2 = ws2["A2"].value
        assert "李四" not in v2 and "张三" not in v2
        assert v2.startswith('=CONCAT("') and v2.count('"') == 4  # 引号结构完整

    def test_t08_outside_literal_untouched(self, tmp_path):
        """T8：字面量外命中（表名/引用含词典词）仅提示，公式原样。"""
        wb = Workbook()
        ws = wb.active
        formula = "=某某建设集团有限公司!A1+SUM(本公司!B1:B2)"
        ws["A1"] = formula
        src = tmp_path / "f2.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["A1"].value == formula
        assert adapter.masker.mappings == []

    def test_t09_formula_cache_cleared(self, tmp_path):
        """T9：公式缓存值在输出中被清空（<v></v>），缓存不泄漏。"""
        sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<worksheet xmlns="http://schemas.openxmlformats.org/'
                 'spreadsheetml/2006/main"><sheetData>'
                 '<row r="1"><c r="A1"><v>100</v></c><c r="B1"><v>200</v></c></row>'
                 '<row r="2"><c r="A2"><f>SUM(A1:B1)</f><v>300</v></c></row>'
                 '</sheetData></worksheet>')
        src = _write_raw_xlsx(tmp_path / "cache.xlsx", sheet)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        xml = zipfile.ZipFile(str(out)).read(
            "xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<v></v>" in xml           # 公式缓存被清空
        assert ">300<" not in xml         # 缓存值不泄漏
        assert "SUM(A1:B1)" in xml        # 公式本体保留
        ws2 = load_workbook(out).active
        assert ws2["A2"].value == "=SUM(A1:B1)"


# ---------------------------------------------------------------------------
# T10-T12 / T20-T21：批注 / 页眉页脚 / 合并单元格 / 工作表名 / 数据验证
# ---------------------------------------------------------------------------

class TestCoverageParts:

    def test_t10_comment_masked_author_cleared(self, tmp_path):
        """T10（R3-B7 修订）：批注文本脱敏，author 单向清除为匿名占位
        （不再保留原作者名，与 docx 侧语义对齐；unmask 不还原）。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "金额说明"
        ws["A1"].comment = Comment("请张三确认", "审计员")
        src = tmp_path / "c.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert _norm_tokens(ws2["A1"].comment.text) == "请[PERSON_N]确认"
        assert ws2["A1"].comment.author == "匿名"
        assert ws2["A1"].value == "金额说明"

    def test_t11_header_footer_masked(self, tmp_path):
        """T11：页眉页脚 6 部位脱敏，&P 控制码保留。"""
        wb = Workbook()
        ws = wb.active
        ws.oddHeader.left.text = "某某建设集团有限公司 &P"
        ws.oddFooter.center.text = "制表：张三"
        ws.firstHeader.right.text = "李四复核"
        src = tmp_path / "hf.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert _norm_tokens(ws2.oddHeader.left.text) == "[COMPANY_N] &P"
        assert _norm_tokens(ws2.oddFooter.center.text) == "制表：[PERSON_N]"
        assert _norm_tokens(ws2.firstHeader.right.text) == "[PERSON_N]复核"
        assert "&P" in ws2.oddHeader.left.text

    def test_t12_merged_cell_anchor_only(self, tmp_path):
        """T12：合并单元格锚点替换，MergedCell 不抛异常，范围与锚点样式保留。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "某某建设集团有限公司"
        ws["A1"].font = Font(bold=True)
        ws.merge_cells("A1:B2")
        src = tmp_path / "m.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert TOKEN_RE.fullmatch(ws2["A1"].value)
        assert {str(r) for r in ws2.merged_cells.ranges} == {"A1:B2"}
        assert ws2["A1"].font.bold is True

    def test_t20_sheet_and_defined_name_hint_only(self, tmp_path, caplog):
        """T20：工作表名/定义名称名只检测提示，名称串不被改动。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "张三数据"
        ws["A1"] = 1
        wb.defined_names["李四合计"] = DefinedName("李四合计", attr_text="张三数据!$A$1")
        src = tmp_path / "names.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        with caplog.at_level(logging.WARNING, logger="mask_tool"):
            out = adapter.process(src, tmp_path / "out")
        wb2 = load_workbook(out)
        assert wb2.sheetnames == ["张三数据"]          # 表名未被改动
        assert "李四合计" in wb2.defined_names         # 名称未被改动
        assert wb2.defined_names["李四合计"].value == "张三数据!$A$1"
        assert "工作表名" in caplog.text
        assert "定义名称" in caplog.text

    def test_t21_data_validation_list_items(self, tmp_path):
        """T21：数据验证 list 项级替换，sqref 不变。"""
        wb = Workbook()
        ws = wb.active
        dv = DataValidation(type="list", formula1='"张三,李四,公开"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("J1:J5")
        src = tmp_path / "dv.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        dvs = ws2.data_validations.dataValidation
        assert len(dvs) == 1
        assert dvs[0].type == "list"
        assert _norm_tokens(dvs[0].formula1) == '"[PERSON_N],[PERSON_N],公开"'
        assert str(dvs[0].sqref) == "J1:J5"


# ---------------------------------------------------------------------------
# T14-T15：富文本
# ---------------------------------------------------------------------------

class TestRichTextCells:

    def test_t14_block_hit_format_preserved(self, tmp_path):
        """T14：块内命中只改该块文本，bold 内联格式与其余块保留。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = CellRichText(TextBlock(InlineFont(b=True), "张三"), " 签字留存")
        src = tmp_path / "r.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out, rich_text=True).active
        value = ws2["A1"].value
        assert isinstance(value, CellRichText)
        blocks = list(value)
        hit = [b for b in blocks
               if not isinstance(b, str) and TOKEN_RE.fullmatch(b.text)]
        assert len(hit) == 1
        assert hit[0].font.b is True                    # 内联格式保留
        assert "签字留存" in str(value)                  # 其余块不动

    def test_t15_cross_block_flattened_with_warning(self, tmp_path, caplog):
        """T15：跨块命中整格展平为 token 串 + warning 日志。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = CellRichText(
            TextBlock(InlineFont(b=True), "某某建设"),
            TextBlock(InlineFont(i=True), "集团有限公司"),
        )
        src = tmp_path / "r2.xlsx"
        wb.save(src)
        # 词库去掉"公司"避免块内子词先命中，专注跨块路径
        adapter = _make_adapter(lexicon={"company": ["某某建设集团有限公司"],
                                         "person": ["张三", "李四"]})
        with caplog.at_level(logging.WARNING, logger="mask_tool"):
            out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out, rich_text=True).active
        assert ws2["A1"].value == "[COMPANY_001]"       # 展平为纯 token 文本
        assert not isinstance(ws2["A1"].value, CellRichText)
        assert "展平" in caplog.text


# ---------------------------------------------------------------------------
# T23：zip 预检
# ---------------------------------------------------------------------------

class TestZipPrecheck:

    @staticmethod
    def _inject(src, dst):
        extra = {
            "xl/threadedComments/threadedComment1.xml":
                b'<?xml version="1.0"?><threadedComments/>',
            "xl/pivotCache/pivotCacheDefinition1.xml":
                b'<?xml version="1.0"?><pivotCacheDefinition/>',
        }
        overrides = {
            "xl/threadedComments/threadedComment1.xml":
                "application/vnd.openxmlformats-officedocument.spreadsheetml."
                "threadedComments+xml",
            "xl/pivotCache/pivotCacheDefinition1.xml":
                "application/vnd.openxmlformats-officedocument.spreadsheetml."
                "pivotCacheDefinition+xml",
        }
        with zipfile.ZipFile(str(src)) as zin:
            entries = {n: zin.read(n) for n in zin.namelist()}
        ct = entries["[Content_Types].xml"].decode("utf-8")
        adds = "".join(
            f'<Override PartName="/{name}" ContentType="{ctype}"/>'
            for name, ctype in overrides.items())
        entries["[Content_Types].xml"] = ct.replace(
            "</Types>", adds + "</Types>").encode("utf-8")
        entries.update(extra)
        with zipfile.ZipFile(str(dst), "w", zipfile.ZIP_DEFLATED) as zout:
            for name, data in entries.items():
                zout.writestr(name, data)
        return dst

    def test_t23_at_risk_parts_warned(self, tmp_path, caplog):
        """T23：含线程批注/透视表缓存部件时产生 warning。"""
        wb = Workbook()
        wb.active["A1"] = "普通"
        base = tmp_path / "base.xlsx"
        wb.save(base)
        src = self._inject(base, tmp_path / "injected.xlsx")
        adapter = _make_adapter()
        with caplog.at_level(logging.WARNING, logger="mask_tool"):
            out = adapter.process(src, tmp_path / "out")
        assert "线程批注" in caplog.text
        assert "数据透视表" in caplog.text
        assert out.exists()
        # N6：openpyxl 保存时丢弃线程批注部件
        names = zipfile.ZipFile(str(out)).namelist()
        assert not any(n.startswith("xl/threadedComments/") for n in names)


# ---------------------------------------------------------------------------
# T24：不可逆模式
# ---------------------------------------------------------------------------

class TestIrreversible:

    def test_t24_all_hits_replaced_with_stars(self, tmp_path):
        """T24：不可逆模式全部命中替换为 ***，无 mapping 增长。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "乙方张三"
        ws["B1"] = 12000000
        ws["B1"].number_format = "¥#,##0.00"
        src = tmp_path / "irr.xlsx"
        wb.save(src)
        adapter = _make_adapter(irreversible=True)
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["A1"].value == "乙方***"
        assert ws2["B1"].value == "***"
        assert adapter.masker.mappings == []


# ---------------------------------------------------------------------------
# T19：unmask 值类型还原
# ---------------------------------------------------------------------------

class TestUnmaskRestore:

    def test_t19_restore_value_kind_matrix(self):
        """T19：kind=number 还原 int/float；kind=text / 缺失降级为文本。"""
        # 整数还原
        value, hits = XlsxAdapter.restore_value(
            "[AMOUNT_001]",
            {"[AMOUNT_001]": {"original": "120000000", "kind": "number"}})
        assert value == 120000000 and isinstance(value, int)
        assert hits == ["[AMOUNT_001]"]
        # 浮点还原
        value, _ = XlsxAdapter.restore_value(
            "[AMOUNT_002]",
            {"[AMOUNT_002]": {"original": "12000.5", "kind": "number"}})
        assert value == 12000.5 and isinstance(value, float)
        # kind=text 的"00123"保持文本，不被类型化
        value, _ = XlsxAdapter.restore_value(
            "[CUSTOM_001]",
            {"[CUSTOM_001]": {"original": "00123", "kind": "text"}})
        assert value == "00123" and isinstance(value, str)
        # kind 缺失（旧 mapping）降级为文本
        value, _ = XlsxAdapter.restore_value(
            "[AMOUNT_003]", {"[AMOUNT_003]": {"original": "120000000"}})
        assert value == "120000000" and isinstance(value, str)
        # TokenMapping 对象条目（含/不含 kind 属性）
        m = TokenMapping(token="[AMOUNT_004]", original="12000.5",
                         text_type=DetectionType.AMOUNT, confidence=0.85)
        m.kind = "number"
        value, _ = XlsxAdapter.restore_value("[AMOUNT_004]", {"[AMOUNT_004]": m})
        assert value == 12000.5
        m2 = TokenMapping(token="[AMOUNT_005]", original="12000.5",
                          text_type=DetectionType.AMOUNT, confidence=0.85)
        value, _ = XlsxAdapter.restore_value("[AMOUNT_005]", {"[AMOUNT_005]": m2})
        assert value == "12000.5" and isinstance(value, str)
        # 混排文本（非整格单 token）永远文本还原
        value, hits = XlsxAdapter.restore_value(
            "合计[AMOUNT_001]元",
            {"[AMOUNT_001]": {"original": "120000000", "kind": "number"}})
        assert value == "合计120000000元"
        # 无命中
        value, hits = XlsxAdapter.restore_value("普通文本", {"[AMOUNT_001]": {"original": "1"}})
        assert value == "普通文本" and hits == []

    def test_t19_sum_reference_recovers_numeric(self, tmp_path):
        """T19：数值还原后 SUM 引用格恢复数值类型。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 120000000
        ws["A1"].number_format = "¥#,##0.00"
        ws["A2"] = "=A1*2"
        src = tmp_path / "sum.xlsx"
        wb.save(src)
        adapter = _make_adapter(mode="aggressive")  # R2 0.70 在 aggressive 下 AUTO
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert TOKEN_RE.fullmatch(ws2["A1"].value)
        token_map = {m.token: m for m in adapter.masker.mappings}
        value, _ = XlsxAdapter.restore_value(ws2["A1"].value, token_map)
        ws2b = load_workbook(out).active  # 模拟 unmask 写回
        ws2b["A1"] = value
        assert isinstance(ws2b["A1"].value, int)
        assert ws2b["A1"].value == 120000000
        assert ws2b["A1"].data_type == "n"


# ---------------------------------------------------------------------------
# T16-T18：大数精度与往返一致性
# ---------------------------------------------------------------------------

def _unmask_file(masked_path, mappings, out_path):
    """测试侧 unmask：改走产品还原链 restore.restore_file_content（R3 审查
    修正——此前测试自实现 DV/定义名称还原逻辑，往返测试未覆盖产品还原
    面，绿灯掩盖了还原面缺口）。"""
    from mask_tool.adapters.restore import restore_file_content

    token_map = {m.token: m for m in mappings}
    restore_file_content(Path(masked_path), Path(out_path), token_map)
    return Path(out_path)


def _canon_cell(cell):
    value = cell.value
    if value is None:
        return ("empty", None)
    if isinstance(value, CellRichText):
        return ("s", str(value))
    if isinstance(value, bool):
        return ("b", value)
    if isinstance(value, (int, float)):
        return ("n", XlsxAdapter._numeric_canonical(value))
    return (cell.data_type, value)


def _style_key(cell):
    # StyleProxy 的 __eq__ 转发后变成 Font==StyleProxy（恒 False），
    # 需先 copy 取出真实样式对象再比较
    return (copy(cell.font), copy(cell.border), copy(cell.fill),
            copy(cell.alignment), copy(cell.protection), cell.number_format)


def _hf_key(ws):
    keys = []
    for part_name in ("oddHeader", "evenHeader", "firstHeader",
                      "oddFooter", "evenFooter", "firstFooter"):
        item = getattr(ws, part_name, None)
        for pos in ("left", "center", "right"):
            keys.append((part_name, pos,
                         getattr(getattr(item, pos, None), "text", None)))
    return keys


def _dv_key(ws):
    return sorted(
        (getattr(dv, "type", None), str(getattr(dv, "formula1", None)),
         str(getattr(dv, "sqref", "")))
        for dv in getattr(ws.data_validations, "dataValidation", []) or []
    )


def _cf_key(ws):
    return sorted(
        (str(fmt.sqref), tuple(
            (rule.type, rule.operator, tuple(rule.formula or []))
            for rule in fmt.rules))
        for fmt in ws.conditional_formatting)


def _dn_key(container):
    return sorted(
        (name, str(getattr(dn, "value", None)),
         getattr(dn, "localSheetId", None))
        for name, dn in container.items())


def _assert_roundtrip_equal(orig_path, restored_path):
    """D2 §6.2：值/类型/样式/批注/页眉/合并/验证/条件格式/定义名称全等比对。"""
    wb_a = load_workbook(str(orig_path), rich_text=True)
    wb_b = load_workbook(str(restored_path), rich_text=True)
    assert wb_a.sheetnames == wb_b.sheetnames
    for name in wb_a.sheetnames:
        sa, sb = wb_a[name], wb_b[name]
        for row in sa.iter_rows():
            for ca in row:
                cb = sb[ca.coordinate]
                assert _canon_cell(ca) == _canon_cell(cb), \
                    f"值差异 {name}!{ca.coordinate}"
                assert _style_key(ca) == _style_key(cb), \
                    f"样式差异 {name}!{ca.coordinate}"
                # R3-B7：批注 author 单向清除不还原——比对 text，
                # 并断言还原件 author 已被置为匿名占位（不再是原作者名）
                text_a = ca.comment.text if ca.comment is not None else None
                text_b = cb.comment.text if cb.comment is not None else None
                assert text_a == text_b, f"批注差异 {name}!{ca.coordinate}"
                assert cb.comment is None or cb.comment.author == "匿名", \
                    f"批注 author 未清除 {name}!{ca.coordinate}"
        assert _hf_key(sa) == _hf_key(sb), f"页眉页脚差异 {name}"
        assert ({str(r) for r in sa.merged_cells.ranges}
                == {str(r) for r in sb.merged_cells.ranges}), f"合并差异 {name}"
        assert _dv_key(sa) == _dv_key(sb), f"数据验证差异 {name}"
        assert _cf_key(sa) == _cf_key(sb), f"条件格式差异 {name}"
        assert _dn_key(sa.defined_names) == _dn_key(sb.defined_names), \
            f"sheet级定义名称差异 {name}"
    assert _dn_key(wb_a.defined_names) == _dn_key(wb_b.defined_names), \
        "全局定义名称差异"


class TestPrecisionRoundtrip:

    def test_t16_19digit_int_digit_exact(self, tmp_path):
        """T16：Excel 原生 19 位整数（手写 XML）mask->unmask 后逐位还原。"""
        sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<worksheet xmlns="http://schemas.openxmlformats.org/'
                 'spreadsheetml/2006/main"><sheetData>'
                 '<row r="1"><c r="A1"><v>4111111111111111110</v></c></row>'
                 '</sheetData></worksheet>')
        src = _write_raw_xlsx(tmp_path / "big.xlsx", sheet)
        wb = load_workbook(str(src))
        original = wb.active["A1"].value
        assert original == 4111111111111111110       # 加载即 Python int（无损）
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["A1"].value == "[CUSTOM_001]"     # Luhn 通过 -> AUTO
        token_map = {m.token: m for m in adapter.masker.mappings}
        value, hits = XlsxAdapter.restore_value(ws2["A1"].value, token_map)
        assert hits == ["[CUSTOM_001]"]
        # R3-B1：16-19 位在从严数值形态内 -> 还原为 int（下游公式引用不断裂），
        # canonical 逐位相等
        assert isinstance(value, int)
        assert value == original == 4111111111111111110
        assert XlsxAdapter._numeric_canonical(value) == str(original)

    def test_t17_float_shortest_roundtrip(self, tmp_path):
        """T17：float 最短往返——12000.5 / -98765.25 unmask 后 canonical 相等。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 12000.5
        ws["A1"].number_format = "¥#,##0.00"
        ws["A2"] = -98765.25
        ws["A2"].number_format = "¥#,##0.00"
        src = tmp_path / "fl.xlsx"
        wb.save(src)
        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert TOKEN_RE.fullmatch(ws2["A1"].value)
        assert TOKEN_RE.fullmatch(ws2["A2"].value)
        token_map = {m.token: m for m in adapter.masker.mappings}
        for coord, original in (("A1", 12000.5), ("A2", -98765.25)):
            value, hits = XlsxAdapter.restore_value(ws2[coord].value, token_map)
            assert hits, coord
            assert isinstance(value, float), coord
            assert XlsxAdapter._numeric_canonical(value) \
                == XlsxAdapter._numeric_canonical(original), coord
        assert token_map[ws2["A1"].value].original == "12000.5"
        assert getattr(token_map[ws2["A1"].value], "kind", "text") == "number"

    def test_t18_full_workbook_roundtrip(self, tmp_path):
        """T18：全特性工作簿 mask->unmask 往返全等（主用例）。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "合同数据"
        ws["A1"] = "甲方"
        ws["A1"].font = Font(bold=True)
        ws["B1"] = "某某建设集团有限公司"
        ws["A2"] = "乙方负责人"
        ws["B2"] = "张三"
        ws["A3"] = "合同金额"
        ws["B3"] = 12000000
        ws["B3"].number_format = "¥#,##0.00"
        ws["B4"] = 120000000            # R2 SUGGEST：默认不替换（H6）
        ws["A5"] = "公式"
        ws["B5"] = '="客户张三"'
        ws["C1"] = 1
        ws["C2"] = 2
        ws["C3"] = "=SUM(C1:C2)"
        ws["D1"] = "编号00123"
        ws["D1"].comment = Comment("请张三确认", "审计员")
        ws["E1"] = "李四"
        ws.merge_cells("E1:F2")
        ws["G1"] = CellRichText(
            TextBlock(InlineFont(b=True), "张三"), " 签字")
        ws["H1"] = datetime(2024, 3, 15)
        ws["H1"].number_format = "yyyy-mm-dd"
        ws["H2"] = True
        ws["I1"] = 12000.5
        ws["I1"].number_format = "¥#,##0.00"
        dv = DataValidation(type="list", formula1='"张三,李四,公开"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("J1:J5")
        ws.conditional_formatting.add(
            "K1:K3",
            CellIsRule(operator="greaterThan", formula=["10"],
                       fill=PatternFill(start_color="FFFF0000",
                                        end_color="FFFF0000",
                                        fill_type="solid")))
        ws.oddHeader.left.text = "某某建设集团有限公司 &P"
        ws.oddFooter.center.text = "制表：张三"
        ws.firstHeader.right.text = "李四复核"
        ws["L1"] = "链接"
        ws["L1"].hyperlink = "https://example.com/doc"
        wb.defined_names["备注常量"] = DefinedName(
            "备注常量", attr_text='="客户李四"')
        ws.defined_names["局部区域"] = DefinedName(
            "局部区域", attr_text="合同数据!$A$1")
        ws2 = wb.create_sheet("明细")
        ws2["A1"] = "普通数据"
        src = tmp_path / "full.xlsx"
        wb.save(src)

        adapter = _make_adapter()
        out = adapter.process(src, tmp_path / "out")

        # 脱敏确实发生：三类 token 均出现，敏感词在单元格值中消失
        wb_masked = load_workbook(str(out), rich_text=True)
        cell_text = " ".join(
            str(c.value) for wsm in wb_masked.worksheets
            for row in wsm.iter_rows() for c in row if c.value is not None)
        assert "[AMOUNT_" in cell_text and "[PERSON_" in cell_text
        assert "[COMPANY_" in cell_text
        assert "张三" not in cell_text and "李四" not in cell_text

        restored = _unmask_file(out, adapter.masker.mappings,
                                tmp_path / "restored.xlsx")
        _assert_roundtrip_equal(src, restored)

        # 还原后无残留 token
        wb_restored = load_workbook(str(restored), rich_text=True)
        restored_text = " ".join(
            str(c.value) for wsm in wb_restored.worksheets
            for row in wsm.iter_rows() for c in row if c.value is not None)
        assert not TOKEN_RE.search(restored_text)


# ---------------------------------------------------------------------------
# 引擎未注入的过渡 shim 路径（当前生产代码在 I1a 合入前的实际路径）
# ---------------------------------------------------------------------------

class TestShimFallback:

    def test_shim_overlap_replacement_and_mappings(self, tmp_path):
        """shim 路径：长度降序替换不嵌套，mapping 正常登记（kind=text）。"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "某某建设集团有限公司与本公司"
        ws["B1"] = 4111111111111111
        src = tmp_path / "shim.xlsx"
        wb.save(src)
        adapter = _make_adapter(attach_engine=False)
        out = adapter.process(src, tmp_path / "out")
        ws2 = load_workbook(out).active
        assert ws2["A1"].value == "[COMPANY_001]与本[COMPANY_002]"
        assert ws2["B1"].value == "[CUSTOM_001]"
        assert len(adapter.masker.mappings) == 3
        assert all(getattr(m, "kind", "text") == "text"
                   for m in adapter.masker.mappings[:2])
        assert getattr(adapter.masker.mappings[2], "kind", "text") == "number"

    def test_shim_output_name_kept(self, tmp_path):
        """process(output_name=...) 支持原名输出；None 时 {stem}_masked.xlsx。"""
        wb = Workbook()
        wb.active["A1"] = "张三"
        src = tmp_path / "原名文件.xlsx"
        wb.save(src)
        adapter = _make_adapter(attach_engine=False)
        out1 = adapter.process(src, tmp_path / "o1")
        assert out1.name == "原名文件_masked.xlsx"
        adapter2 = _make_adapter(attach_engine=False)
        out2 = adapter2.process(src, tmp_path / "o2", output_name="原名文件.xlsx")
        assert out2.name == "原名文件.xlsx"
