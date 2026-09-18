# -*- coding: utf-8 -*-
"""tests/test_web_zip_and_sync.py — I6 问题1/2/3：下载规则、目录 zip、勾选同步

覆盖：
- _safe_unzip：正常解压结构保留、路径穿越拒绝、zip 炸弹防护
  （成员数超限、声明总大小超限）
- 目录 zip 往返：检测（task_kind/相对路径登记/屏蔽类型登记）→ 脱敏
  （产物 zip 同结构 + 文件名 token + 非文档拷贝 + 屏蔽类型不进产物）
  → _run_restore_zip 还原（文件名/目录名/内容一致，mapping.json 不进产物）
- 单文件下载规则：单文件任务 download_kind=file（直接下同名单文件）；
  zip 任务 download_kind=zip
- 勾选同步（问题3）：_apply_grid_selection（取消勾选回传 → user_selections
  即时变化）+ _final_selected_indices（确认列表与勾选一致）
  + AgGrid update_mode 接线断言（SELECTION_CHANGED 触发同步回传）
"""

import io
import inspect
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from docx import Document

from mask_tool.core.templates import (
    DEFAULT_CONFIG_YAML, SAMPLE_LEXICON_YAML, WHITELIST_YAML,
)

LEXICON = {"person": ["李四"]}


@pytest.fixture(autouse=True)
def _cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _write_config_dir(lexicon=None):
    cfgdir = Path("config")
    cfgdir.mkdir(parents=True, exist_ok=True)
    default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")
    (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
    (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
    (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
    (cfgdir / "lexicon.yaml").write_text(
        yaml.dump(lexicon or LEXICON, allow_unicode=True), encoding="utf-8",
    )


def _make_docx_bytes(paragraphs) -> bytes:
    buf = io.BytesIO()
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def _upload(name: str, payload: bytes) -> SimpleNamespace:
    return SimpleNamespace(name=name, read=lambda: payload, size=len(payload))


def _patch_streamlit(monkeypatch, web):
    monkeypatch.setattr(web, "BATCHES_DIR", Path("batches"))
    monkeypatch.setattr(web, "_add_history", lambda rec: None)
    for fn in ("rerun", "error", "warning", "info", "success", "caption"):
        monkeypatch.setattr(web.st, fn, lambda *a, **k: None, raising=False)
    monkeypatch.setattr(
        web.st, "spinner", lambda *a, **k: _NullCtx(), raising=False,
    )
    for key in list(web.st.session_state.keys()):
        del web.st.session_state[key]


class _NullCtx:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _make_dir_zip_bytes(members: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


# ===========================================================================
# _safe_unzip：解压防护
# ===========================================================================

class TestSafeUnzip:

    def test_normal_unzip_preserves_structure(self, monkeypatch, tmp_path):
        """正常目录 zip：解压后相对结构完整保留。"""
        web = pytest.importorskip("mask_tool.web.app")
        payload = _make_dir_zip_bytes({
            "项目资料/合同.docx": b"fake-docx",
            "项目资料/子目录/说明.txt": "文本",
            "项目资料/顶层.xlsx": b"fake-xlsx",
        })
        root = web._safe_unzip(payload, tmp_path)
        assert (root / "项目资料" / "合同.docx").read_bytes() == b"fake-docx"
        assert (root / "项目资料" / "子目录" / "说明.txt").read_text(
            encoding="utf-8") == "文本"
        assert (root / "项目资料" / "顶层.xlsx").exists()

    def test_path_traversal_rejected(self, monkeypatch, tmp_path):
        """路径穿越成员（../escape.txt）→ 整包拒绝（ValueError）。"""
        web = pytest.importorskip("mask_tool.web.app")
        evil = _make_dir_zip_bytes({
            "正常.txt": "ok",
            "../越界.txt": "evil",
        })
        with pytest.raises(ValueError, match="越界|穿越"):
            web._safe_unzip(evil, tmp_path)
        # 绝对路径形态同样拒绝
        evil2 = _make_dir_zip_bytes({"/etc/passwd": "evil"})
        with pytest.raises(ValueError):
            web._safe_unzip(evil2, tmp_path)

    def test_zip_bomb_entry_count_rejected(self, monkeypatch, tmp_path):
        """成员数超过上限（500）→ 拒绝解压。"""
        web = pytest.importorskip("mask_tool.web.app")
        members = {f"d/f{i:03d}.txt": "x" for i in range(web.ZIP_MAX_ENTRIES + 1)}
        with pytest.raises(ValueError, match="成员数"):
            web._safe_unzip(_make_dir_zip_bytes(members), tmp_path)

    def test_zip_bomb_total_size_rejected(self, monkeypatch, tmp_path):
        """声明解压总大小超过上限 → 拒绝解压（逻辑验证：缩小上限注入）。"""
        web = pytest.importorskip("mask_tool.web.app")
        monkeypatch.setattr(web, "ZIP_MAX_TOTAL_BYTES", 16)
        payload = _make_dir_zip_bytes({"big.bin": b"x" * 32})
        with pytest.raises(ValueError, match="总大小"):
            web._safe_unzip(payload, tmp_path)


# ===========================================================================
# 目录 zip 往返（问题2）
# ===========================================================================

class TestDirectoryZipRoundtrip:

    def _run_full_zip_task(self, monkeypatch):
        """跑完整目录 zip 任务，返回 (web, mask_result, 上传 zip bytes)。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)

        zip_bytes = _make_dir_zip_bytes({
            "项目资料/合同-李四.docx": _make_docx_bytes(["负责人李四签约"]),
            "项目资料/子目录/说明.txt": "普通说明文件李四",
            "项目资料/机密.pptx": b"fake-pptx",
            "项目资料/汇报-李四.docx": _make_docx_bytes(["汇报人李四"]),
        })
        zip_upload = _upload("项目资料.zip", zip_bytes)

        web._run_detection(
            [], "smart", ner_enabled=False,
            manual_words=["李四"], manual_only=False, zip_file=zip_upload,
        )
        ss = web.st.session_state
        assert ss["task_kind"] == "zip"
        assert ss["zip_blocked_files"] == ["项目资料/机密.pptx"]
        assert set(ss["file_results"].keys()) == {
            "项目资料/合同-李四.docx", "项目资料/汇报-李四.docx",
        }

        results = ss["detection_results"]
        # 跨文件去重（既有语义）：同词全局一条；脱敏时 allowed_originals 全局生效
        assert [r.text for r in results] == ["李四"]

        web._run_masking(
            [], list(range(len(results))), results,
            mode="smart", ner_enabled=False, irreversible=False,
            learn_words=False, batch_id="MSK-ZIP", batch_name="t",
            mask_filenames=True, manual_words=["李四"], manual_only=False,
        )
        return web, web.st.session_state["mask_result"], zip_bytes

    def test_mask_zip_structure_names_and_blocked_excluded(self, monkeypatch):
        """产物 zip：同结构镜像树、文件名 token 化、非文档拷贝、
        屏蔽类型不进产物；mapping 单独可下载。"""
        web, mr, _ = self._run_full_zip_task(monkeypatch)

        assert mr["download_kind"] == "zip"
        assert mr["task_kind"] == "zip"
        out = zipfile.ZipFile(io.BytesIO(mr["zip_buffer"]))
        names = out.namelist()
        # 屏蔽类型不进产物（勿把未脱敏文件带进交付物）
        assert not any(n.endswith(".pptx") for n in names), names
        # 同结构 + 文件名脱敏（token 前缀）+ 非文档原样拷贝
        assert "项目资料/合同-[CUSTOM_001].docx" in names
        assert "项目资料/汇报-[CUSTOM_001].docx" in names
        assert "项目资料/子目录/说明.txt" in names
        assert "mapping.json" in names
        # mapping 数据独立存在（单独下载按钮用）
        mapping = json.loads(mr["mapping_data"])
        assert mapping["tokens"]["[CUSTOM_001]"]["original"] == "李四"
        assert len(mapping.get("paths", [])) >= 2  # 文件名改名记录

    def test_restore_zip_roundtrip(self, monkeypatch):
        """还原 zip：文件名/目录名/内容一致；mapping.json 不进还原产物。"""
        web, mr, _ = self._run_full_zip_task(monkeypatch)

        mapping = json.loads(mr["mapping_data"])
        restore_upload = _upload("masked.zip", mr["zip_buffer"])
        web._run_restore_zip(restore_upload, mapping)

        rr = web.st.session_state["restore_result"]
        assert rr is not None and rr["file_count"] == 2
        rz = zipfile.ZipFile(io.BytesIO(rr["zip_buffer"]))
        names = sorted(rz.namelist())
        # 目录结构 + 原文件名还原
        assert names == [
            "项目资料/合同-李四.docx",
            "项目资料/子目录/说明.txt",
            "项目资料/汇报-李四.docx",
        ]
        # 内容还原：docx 段落回到原文
        import tempfile, shutil as _sh
        td = Path(tempfile.mkdtemp())
        try:
            rz.extractall(td)
            texts = {}
            for p in td.rglob("*.docx"):
                texts[p.name] = [q.text for q in Document(str(p)).paragraphs]
            assert texts["合同-李四.docx"] == ["负责人李四签约"]
            assert texts["汇报-李四.docx"] == ["汇报人李四"]
            assert (td / "项目资料" / "子目录" / "说明.txt").read_text(
                encoding="utf-8") == "普通说明文件李四"
        finally:
            _sh.rmtree(td, ignore_errors=True)

    def test_single_file_task_downloads_bare_file(self, monkeypatch):
        """单文件任务（问题1下载规则）：download_kind=file，
        file_bytes 为脱敏后产物、file_name 为产物名。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)

        upload = _upload(
            "合同-李四.docx", _make_docx_bytes(["负责人李四签约"]),
        )
        web._run_detection(
            [upload], "smart", ner_enabled=False,
            manual_words=["李四"], manual_only=False,
        )
        results = web.st.session_state["detection_results"]
        web._run_masking(
            [upload], list(range(len(results))), results,
            mode="smart", ner_enabled=False, irreversible=False,
            learn_words=False, batch_id="MSK-ONE", batch_name="t",
            mask_filenames=True, manual_words=["李四"], manual_only=False,
        )
        mr = web.st.session_state["mask_result"]
        assert mr["download_kind"] == "file"
        assert mr["file_name"] == "合同-[CUSTOM_001]_masked.docx"
        # bytes 可直接作为下载 data 且是有效 docx
        doc = Document(io.BytesIO(mr["file_bytes"]))
        assert [p.text for p in doc.paragraphs] == ["负责人[CUSTOM_001]签约"]


# ===========================================================================
# 勾选与确认列表同步（问题3）
# ===========================================================================

class TestSelectionSync:

    def test_uncheck_syncs_final_list(self):
        """表格取消勾选（grid 回传）→ user_selections 即时变化 →
        \"即将脱敏\"确认列表与计数一致。"""
        web = pytest.importorskip("mask_tool.web.app")
        all_results = [f"r{i}" for i in range(4)]  # 4 条检测项
        selections = {0: True, 1: True, 2: True, 3: False}
        filtered = [0, 1, 2, 3]

        # 初始：3 项确认
        assert web._final_selected_indices(all_results, selections) == [0, 1, 2]

        # 用户在表格取消勾选第 0、2 项 → AgGrid 回传当前选中行
        changed = web._apply_grid_selection(
            filtered, [{"index": 1}], selections,
        )
        assert changed is True
        assert selections == {0: False, 1: True, 2: False, 3: False}
        # 确认列表立即只剩第 1 项
        final = web._final_selected_indices(all_results, selections)
        assert final == [1]
        assert sum(1 for i in filtered if selections.get(i, False)) == 1

        # 再次回传相同状态 → 无变化（幂等）
        assert web._apply_grid_selection(
            filtered, [{"index": 1}], selections,
        ) is False

    def test_grid_return_none_keeps_state(self):
        """selected_rows=None（无回传）→ 状态不变。"""
        web = pytest.importorskip("mask_tool.web.app")
        selections = {0: True}
        assert web._apply_grid_selection([0], None, selections) is False
        assert selections == {0: True}

    def test_filtered_rows_untouched(self):
        """被筛选隐藏的行勾选态不受当前视图回传影响。"""
        web = pytest.importorskip("mask_tool.web.app")
        selections = {0: True, 1: True}
        # 当前视图只显示第 0 行，且未勾选
        web._apply_grid_selection([0], [], selections)
        assert selections == {0: False, 1: True}  # 第 1 行保持原状

    def test_aggrid_wired_to_selection_changed(self):
        """接线断言：AgGrid update_mode 必须是 SELECTION_CHANGED
        （勾选变化即回传并触发 rerun，确认列表才能同步）。"""
        web = pytest.importorskip("mask_tool.web.app")
        source = inspect.getsource(web._render_masking_tab)
        assert "GridUpdateMode.SELECTION_CHANGED" in source
        assert "GridUpdateMode.NO_UPDATE" not in source
        # 同步与确认列表统一走提取后的函数
        assert "_apply_grid_selection(" in source
        assert "_final_selected_indices(" in source
