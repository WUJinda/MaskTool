# -*- coding: utf-8 -*-
"""tests/test_r3_fixes.py — R2/R3 双审查 A/B/C 级修复的回归测试

覆盖（每项至少 1 例）：
- A1 xlsx 三部件（数据验证 list / 定义名称值+备注 / 超链接 tooltip）
     mask -> 产品还原链 restore_file_content 往返：原文恢复 + token 零残留
- A2 ArrayFormula：检测面可见（detect_file_results / extract_texts）
     + mask 写入 + 产品还原往返
- A3 Web _run_masking 调 pipeline.prepare：自然 token 撞号让位
- A4 CLI --irreversible 单文件模式：产物文件名 ___ 样式（非可逆 token）
- B1 16-19 位 Luhn 卡号 kind=number 还原为 int（含 19 位边界与 20 位降级）
- B2 Web 批次 ID：mapping metadata.batch_id 与批次目录名对齐
- B3 CLI --confirm 文件名 SUGGEST（ner）勾选后改名；默认模式仍不改名
- B7 xlsx 批注 author 清除（单向，还原不恢复）
- B6/C11 19 位整值（int / float 形态）未脱敏保留时预转文本保真
- C10 unmask_tree 还原名被占用：加后缀 + warnings 记录
"""

import io
import json
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from typer.testing import CliRunner

import mask_tool.cli as cli_module
from mask_tool.cli import app
from mask_tool.core.templates import (
    DEFAULT_CONFIG_YAML, SAMPLE_LEXICON_YAML, WHITELIST_YAML,
)
from mask_tool.models.detection import (
    DetectionResult, DetectionStatus, DetectionType, Location,
)

TOKEN_RE = re.compile(
    r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]"
)

LEXICON = {
    "company": ["某某科技有限公司"],
    "person": ["张三", "李四"],
}


@pytest.fixture()
def runner():
    return CliRunner(env={"COLUMNS": "400"})


@pytest.fixture(autouse=True)
def _cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _flat(output: str) -> str:
    return "".join(output.split())


def _write_config_dir(lexicon=None):
    cfgdir = Path("config")
    cfgdir.mkdir(parents=True, exist_ok=True)
    default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")
    (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
    (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
    (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
    (cfgdir / "lexicon.yaml").write_text(
        yaml.dump(lexicon or LEXICON, allow_unicode=True), encoding="utf-8",
    )


def _make_docx(path: Path, paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _docx_paragraphs(path: Path):
    return [p.text for p in Document(str(path)).paragraphs]


def _batch_dirs(output: Path = Path("output")):
    return sorted(p for p in output.iterdir() if p.is_dir())


def _make_xlsx_adapter(lexicon=None):
    """xlsx adapter（词典/策略/引擎均用生产实现，仅词库最小化）。"""
    from mask_tool.adapters.xlsx_adapter import XlsxAdapter
    from mask_tool.core.detector import Detector
    from mask_tool.core.masker import Masker
    from mask_tool.core.policy import PolicyEngine
    from mask_tool.core.tokenizer import TokenGenerator
    from mask_tool.models.config import MaskConfig

    detector = Detector(
        {k: list(v) for k, v in (lexicon or LEXICON).items()}, set(),
    )
    policy = PolicyEngine(MaskConfig(mode="smart"))
    token_gen = TokenGenerator()
    masker = Masker(token_gen, irreversible=False)
    if getattr(masker, "engine", None) is None:  # 防御：未合入引擎的环境
        from mask_tool.core.engine import ReplacementEngine
        masker.engine = ReplacementEngine(token_gen, irreversible=False)
    return XlsxAdapter(detector, policy, masker)


def _token_map_of(adapter):
    return {m.token: m for m in adapter.masker.mappings}


# 手写最小 xlsx：绕过 openpyxl 的 %.16g 序列化，构造真 int 大数源
# （openpyxl 保存的 19 位 int 在 XML 中已是科学计数，重载即 float）
_RAW_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
           '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
           'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
           '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
           'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
           '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/'
           'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
_RAW_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
             'relationships"><Relationship Id="rId1" Type="http://schemas.'
             'openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
             'Target="xl/workbook.xml"/></Relationships>')
_RAW_WB = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
           'main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
           'relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/>'
           '</sheets></workbook>')
_RAW_WBRELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
               'relationships"><Relationship Id="rId1" Type="http://schemas.'
               'openxmlformats.org/officeDocument/2006/relationships/worksheet" '
               'Target="worksheets/sheet1.xml"/></Relationships>')


def _write_raw_xlsx(path, sheet_xml):
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _RAW_CT)
        z.writestr("_rels/.rels", _RAW_RELS)
        z.writestr("xl/workbook.xml", _RAW_WB)
        z.writestr("xl/_rels/workbook.xml.rels", _RAW_WBRELS)
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


# ===========================================================================
# A1：xlsx 三部件还原（产品链往返）
# ===========================================================================

class TestA1ThreePartsRoundtrip:

    def test_dv_defined_name_tooltip_roundtrip(self, tmp_path):
        """三部件 mask 写入 token -> 产品 restore_file_content 原文恢复，
        对账提取面零残留（此前 Web 还原静默残留 token 的缺口）。"""
        from mask_tool.adapters.extract import extract_texts
        from mask_tool.adapters.restore import restore_file_content

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "普通文本"
        dv = DataValidation(type="list", formula1='"张三,公开项"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("B1:B5")
        wb.defined_names["备注常量"] = DefinedName(
            "备注常量", attr_text='="客户李四"')
        ws["C1"] = "链接"
        ws["C1"].hyperlink = "https://example.com/doc"
        ws["C1"].hyperlink.tooltip = "某某科技有限公司合同"
        src = tmp_path / "parts.xlsx"
        wb.save(src)

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")

        # mask 断言：三部件 token 写入（处理顺序：全局定义名称先于 DV，
        # 故李四=PERSON_001、张三=PERSON_002）
        wbm = load_workbook(out)
        dv1 = wbm.active.data_validations.dataValidation[0]
        assert dv1.formula1 == '"[PERSON_002],公开项"'
        assert wbm.defined_names["备注常量"].value == '="客户[PERSON_001]"'
        assert wbm.active["C1"].hyperlink.tooltip == "[COMPANY_001]合同"

        # 产品还原链（与 Web/CLI unmask 同一入口）
        restored = tmp_path / "restored" / "parts.xlsx"
        restore_file_content(out, restored, _token_map_of(adapter))
        wbr = load_workbook(restored)
        assert wbr.active.data_validations.dataValidation[0].formula1 \
            == '"张三,公开项"'
        assert wbr.defined_names["备注常量"].value == '="客户李四"'
        assert wbr.active["C1"].hyperlink.tooltip == "某某科技有限公司合同"
        # 对账提取面零残留（extract_texts 返回拼接串，直接整串搜索）
        assert not TOKEN_RE.search(extract_texts(restored))


# ===========================================================================
# A2：ArrayFormula 三面对称
# ===========================================================================

class TestA2ArrayFormula:

    def test_detect_visible_and_roundtrip(self, tmp_path):
        """检测面可见（detect_file_results / extract_texts 含字面量实体）
        + mask 写入 .text + 产品还原往返。"""
        from openpyxl.worksheet.formula import ArrayFormula
        from mask_tool.adapters.extract import (
            detect_file_results, extract_texts,
        )
        from mask_tool.adapters.restore import restore_file_content
        from mask_tool.core.detector import Detector
        from mask_tool.core.policy import PolicyEngine
        from mask_tool.models.config import MaskConfig

        wb = Workbook()
        wb.active["B1"] = ArrayFormula("B1:B1", '="客户张三"')
        src = tmp_path / "af.xlsx"
        wb.save(src)

        cfg = MaskConfig(mode="smart")
        detector = Detector({k: list(v) for k, v in LEXICON.items()}, set())
        policy = PolicyEngine(cfg)
        results = detect_file_results(src, detector, policy)
        assert "张三" in {r.text for r in results}
        assert "张三" in extract_texts(src)   # 提取面（str）同覆盖

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")
        wbm = load_workbook(out)
        af = wbm.active["B1"].value
        assert type(af).__name__ == "ArrayFormula"
        assert "张三" not in af.text and "[PERSON_001]" in af.text

        restored = tmp_path / "restored" / "af.xlsx"
        restore_file_content(out, restored, _token_map_of(adapter))
        afr = load_workbook(restored).active["B1"].value
        assert type(afr).__name__ == "ArrayFormula"
        assert afr.text == '="客户张三"'
        assert not TOKEN_RE.search(extract_texts(restored))


# ===========================================================================
# A3 + B2：Web _run_masking 的 prepare 预扫描与批次 ID 对齐
# ===========================================================================

class TestA3B2WebMasking:

    def test_run_masking_prepares_and_batch_id_aligned(
        self, monkeypatch, tmp_path,
    ):
        """prepare 被调用且编号让位（自然 [PERSON_001] -> 张三 用 _002）；
        mapping metadata.batch_id 与批次目录名一致。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()

        from mask_tool.core.pipeline import Pipeline
        calls = []
        original_prepare = Pipeline.prepare

        def spy(self, files):
            calls.append([Path(f).name for f in files])
            return original_prepare(self, files)

        monkeypatch.setattr(Pipeline, "prepare", spy)
        monkeypatch.setattr(web, "BATCHES_DIR", tmp_path / "batches")
        monkeypatch.setattr(web, "_add_history", lambda rec: None)
        monkeypatch.setattr(web.st, "rerun", lambda *a, **k: None, raising=False)
        monkeypatch.setattr(web.st, "error", lambda *a, **k: None)
        monkeypatch.setattr(web.st, "warning", lambda *a, **k: None)
        for key in list(web.st.session_state.keys()):
            del web.st.session_state[key]

        # 文档自然文本含 token 样式串（撞号源）+ 敏感词
        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("参考 [PERSON_001] 的编号说明")
        doc.add_paragraph("负责人张三")
        doc.save(buf)
        upload = SimpleNamespace(name="a.docx", read=lambda: buf.getvalue())

        all_results = [
            DetectionResult(
                text="张三", text_type=DetectionType.PERSON,
                source="dictionary", confidence=0.95,
                location=Location(file="a.docx"),
            ),
        ]
        web._run_masking(
            [upload], [0], all_results,
            mode="smart", ner_enabled=False, irreversible=False,
            learn_words=False, batch_id="MSK-TEST", batch_name="t",
            mask_filenames=False,
        )

        # A3：prepare 以已保存的上传文件被调用
        assert calls == [["a.docx"]]
        # A3：编号让位——敏感词拿到的 token 不是输入中已有的 [PERSON_001]
        result = web.st.session_state["mask_result"]
        tokens = json.loads(result["mapping_data"])["tokens"]
        assert "[PERSON_002]" in tokens
        assert "[PERSON_001]" not in tokens
        assert tokens["[PERSON_002]"]["original"] == "张三"
        # B2：批次目录名与 mapping metadata.batch_id 对齐
        assert (tmp_path / "batches" / "MSK-TEST").is_dir()
        assert json.loads(
            (tmp_path / "batches" / "MSK-TEST" / "mapping.json").read_text(
                encoding="utf-8")
        )["metadata"]["batch_id"] == "MSK-TEST"


# ===========================================================================
# A4：CLI --irreversible 单文件模式文件名
# ===========================================================================

class TestA4IrreversibleFileName:

    def test_single_file_irreversible_name_underscores(self, runner):
        """--irreversible 单文件：内容 *** 且文件名 ___ 样式（无 token），
        与目录模式/Web 行为一致（R3-A4：此前文件名误用可逆 token）。"""
        _write_config_dir()
        src = _make_docx(Path("合同-张三.docx"), ["负责人张三"])
        result = runner.invoke(
            app, ["mask", str(src), "--output", "out", "--irreversible"],
        )
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        names = [p.name for p in batch.glob("*.docx")]
        assert len(names) == 1
        name = names[0]
        assert name.endswith("_masked.docx")
        assert "___" in name                     # *** 经 sanitize 落盘为 ___
        assert not TOKEN_RE.search(name)
        assert "张三" not in name
        assert _docx_paragraphs(batch / name) == ["负责人***"]


# ===========================================================================
# B1：16-19 位卡号还原为 int
# ===========================================================================

class TestB1CardNumberRestore:

    def test_restore_value_16_to_19_digit_as_int(self):
        """kind=number 的整值 token：16/19 位 Luhn 卡号还原为 int（下游公式
        引用不断裂）；20 位超界仍降级文本。"""
        from mask_tool.adapters.xlsx_adapter import XlsxAdapter

        for original in ("4111111111111111", "4111111111111111110"):
            token = "[CUSTOM_001]"
            value, hits = XlsxAdapter.restore_value(
                token, {token: {"original": original, "kind": "number"}},
            )
            assert hits == [token]
            assert isinstance(value, int), (original, type(value))
            assert str(value) == original

        # 20 位：超出 ≤19 位从严形态，降级文本（值不丢）
        token20 = "[CUSTOM_002]"
        v20, _ = XlsxAdapter.restore_value(
            token20, {token20: {"original": "41111111111111111100",
                                "kind": "number"}},
        )
        assert isinstance(v20, str) and v20 == "41111111111111111100"

    def test_masked_16_digit_card_roundtrip_int(self, tmp_path):
        """端到端：16 位 Luhn 卡号单元格 mask -> 产品还原为 int。"""
        from mask_tool.adapters.restore import restore_file_content

        wb = Workbook()
        wb.active["A1"] = 4111111111111111
        src = tmp_path / "card.xlsx"
        wb.save(src)

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")
        assert load_workbook(out).active["A1"].value == "[CUSTOM_001]"

        restored = tmp_path / "restored" / "card.xlsx"
        restore_file_content(out, restored, _token_map_of(adapter))
        value = load_workbook(restored).active["A1"].value
        assert isinstance(value, int) and value == 4111111111111111


# ===========================================================================
# B3：CLI --confirm 文件名 SUGGEST 勾选后改名
# ===========================================================================

class TestB3ConfirmFileNameSuggest:

    def test_confirm_checked_suggest_renames_filename(self, runner, monkeypatch):
        """--confirm 全量勾选：ner SUGGEST（文件名"张三"）勾选后文件名
        同步改名（与内容侧/Web 语义一致）；默认模式仍不改名。"""
        from mask_tool.core.confirm import ConfirmEngine
        from mask_tool.core.ner.jieba_ner import JiebaNER

        pytest.importorskip("jieba")
        if not JiebaNER().is_available():
            pytest.skip("jieba 不可用")

        # 词库不带 person：正文/文件名的"张三"仅来自 ner（SUGGEST）
        lexicon = {"company": ["某某科技有限公司"]}
        cfgdir = Path("config")
        cfgdir.mkdir(parents=True, exist_ok=True)
        (cfgdir / "default.yaml").write_text(DEFAULT_CONFIG_YAML, encoding="utf-8")
        (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
        (cfgdir / "sample_lexicon.yaml").write_text(
            SAMPLE_LEXICON_YAML, encoding="utf-8")
        (cfgdir / "lexicon.yaml").write_text(
            yaml.dump(lexicon, allow_unicode=True), encoding="utf-8",
        )

        src = _make_docx(Path("张三科技.docx"), ["负责人张三"])
        monkeypatch.setattr(
            cli_module, "_new_confirm_engine",
            lambda: ConfirmEngine(auto_yes=True),
        )
        result = runner.invoke(
            app, ["mask", str(src), "--output", "out", "--confirm"],
        )
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        names = [p.name for p in batch.glob("*.docx")]
        assert any("[PERSON_" in n for n in names), names
        # 内容侧同步替换（张三 勾选放行）
        assert _docx_paragraphs(
            next(batch.glob("*[PERSON_*.docx"))
        ) == ["负责人[PERSON_001]"]

        # 对照：默认（无 --confirm）SUGGEST 不改名、不替换
        src2 = _make_docx(Path("张三科技.docx"), ["负责人张三"])
        result2 = runner.invoke(app, ["mask", str(src2), "--output", "out2"])
        assert result2.exit_code == 0, result2.output
        batch2 = _batch_dirs(Path("out2"))[0]
        assert (batch2 / "张三科技_masked.docx").exists()
        assert _docx_paragraphs(batch2 / "张三科技_masked.docx") == ["负责人张三"]


# ===========================================================================
# B7：xlsx 批注 author 清除
# ===========================================================================

class TestB7CommentAuthor:

    def test_comment_author_cleared_and_not_restored(self, tmp_path):
        """mask 清除批注 author（真实人名元数据）；产品还原不恢复 author
        （单向脱敏，与 docx 侧语义一致）。"""
        from mask_tool.adapters.restore import restore_file_content

        wb = Workbook()
        wb.active["A1"] = "普通"
        wb.active["A1"].comment = Comment("请张三确认", "审计员")
        src = tmp_path / "cmt.xlsx"
        wb.save(src)

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")
        wbm = load_workbook(out)
        assert wbm.active["A1"].comment.text == "请[PERSON_001]确认"
        assert wbm.active["A1"].comment.author == "匿名"

        restored = tmp_path / "restored" / "cmt.xlsx"
        restore_file_content(out, restored, _token_map_of(adapter))
        wbr = load_workbook(restored)
        assert wbr.active["A1"].comment.text == "请张三确认"
        assert wbr.active["A1"].comment.author == "匿名"   # 单向：不还原


# ===========================================================================
# B6/C11：19 位整值（int / float 形态）预转文本保真
# ===========================================================================

class TestBigIntegerTextPreservation:

    def test_unmasked_19digit_int_kept_as_exact_text(self, tmp_path):
        """未脱敏保留的 19 位整数（非 Luhn，SUGGEST 不替换）预转文本，
        值逐位保真，number_format 置文本格式。

        源文件用手写 XML：openpyxl 自身保存大 int 用 %.16g 科学计数，
        保存产物重载即 float（那是 float 形态分支的用例范围）；
        Excel/外部工具写出的纯数字 <v> 重载才是真 int。"""
        sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<worksheet xmlns="http://schemas.openxmlformats.org/'
                 'spreadsheetml/2006/main"><sheetData>'
                 '<row r="1"><c r="A1"><v>9111111111111111110</v></c></row>'
                 '</sheetData></worksheet>')
        src = _write_raw_xlsx(tmp_path / "big.xlsx", sheet)
        wb0 = load_workbook(str(src))
        assert wb0.active["A1"].value == 9111111111111111110   # 加载即 int

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")
        cell = load_workbook(out).active["A1"]
        assert isinstance(cell.value, str)
        assert cell.value == "9111111111111111110"
        assert cell.number_format == "@"

    def test_unmasked_19digit_float_form_kept_as_exact_text(self, tmp_path):
        """float 形态（19 位 int 读入即 float）：is_integer 且 abs≥10^15
        同样预转文本（R3-C11 扩展），按 double 精确整值逐位落盘。"""
        original = 4.1111111111111111e+18      # is_integer() 且 ≥10^15
        assert original.is_integer()
        wb = Workbook()
        wb.active["A1"] = original
        src = tmp_path / "bigf.xlsx"
        wb.save(src)

        adapter = _make_xlsx_adapter()
        out = adapter.process(src, tmp_path / "out")
        cell = load_workbook(out).active["A1"]
        assert isinstance(cell.value, str)
        assert cell.value == str(int(original))    # double 精确整值
        assert cell.number_format == "@"


# ===========================================================================
# C10：unmask_tree 还原名被占用警告
# ===========================================================================

class TestC10UnmaskNameTakenWarning:

    def test_unmask_tree_warns_when_restore_name_taken(self, tmp_path):
        """还原目标名已被占用：加 _1 后缀落盘，result.warnings 记录
        （CLI 侧打印 result.warnings 自然透出）。"""
        from mask_tool.core.path_masker import PathMasker, PathMapping

        root = tmp_path / "tree"
        root.mkdir()
        (root / "[PERSON_001].txt").write_text("masked", encoding="utf-8")
        (root / "张三.txt").write_text("occupied", encoding="utf-8")  # 占位
        m = PathMapping(
            kind="file", old_name="张三.txt", new_name="[PERSON_001].txt",
            rel_old="张三.txt", rel_new="[PERSON_001].txt", depth=1,
            categories=["person"], tokens_used=["[PERSON_001]"],
            fingerprint="", status="renamed",
        )
        pm = PathMasker(None, None, None)
        result = pm.unmask_tree(root, [m])
        assert (root / "张三_1.txt").exists()
        assert any("已被占用" in w for w in result.warnings), result.warnings


# ===========================================================================
# README：已知限制表述更新（C12）
# ===========================================================================

class TestReadmeLimits:

    def test_readme_r3_wording_updated(self):
        readme = Path(__file__).parent.parent / "README.md"
        text = readme.read_text(encoding="utf-8")
        # 跨 tab 表述：不再声称"检测不到"，改为位置漂移
        assert "跨 tab/换行实体" in text
        assert "位置漂移" in text
        assert "检测不到" not in text
        # 检测临时副本残留提示（2026-09-18 桌面化：原"Web 检测临时副本"）
        assert "检测临时副本" in text
        # ≥16 位整值统一转文本保真声明
        assert "≥16 位" in text
