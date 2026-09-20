"""tests/test_cli.py — CLI 波次2 测试（D1 §8 CLI 1-7 + 目录模式端到端 + D 项专项）

覆盖：H5 配置回退链、N2 词库自动复制、M4 批次目录、H6 --all、PPT/PDF 屏蔽、
M3 reserved/对账/指纹、N1 inspect 盲区消除、--confirm 重构、目录模式
mask->unmask 端到端结构比对。
"""

import json
import re
import shutil
from pathlib import Path

import pytest
import yaml
from docx import Document
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

import mask_tool.cli as cli_module
from mask_tool.cli import app
from mask_tool.core.templates import (
    DEFAULT_CONFIG_YAML, SAMPLE_LEXICON_YAML, WHITELIST_YAML,
)

TOKEN_RE = re.compile(r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]")

LEXICON = {
    "company": ["某某科技有限公司"],
    "person": ["张三", "李四"],
}


@pytest.fixture()
def runner():
    # 大宽度避免 rich 换行截断断言片段
    return CliRunner(env={"COLUMNS": "400"})


@pytest.fixture(autouse=True)
def _cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _flat(output: str) -> str:
    """压平输出（去空白），规避 rich 换行对断言的影响。"""
    return "".join(output.split())


def _write_config_dir(lexicon=None, with_lexicon=True, with_sample=True):
    cfgdir = Path("config")
    cfgdir.mkdir(parents=True, exist_ok=True)
    # NER 关闭保证检测确定性（词典/正则可精确断言）
    default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")
    (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
    (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
    if with_sample:
        (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
    if with_lexicon:
        (cfgdir / "lexicon.yaml").write_text(
            yaml.dump(lexicon or LEXICON, allow_unicode=True), encoding="utf-8",
        )


def _make_docx(path: Path, paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _make_xlsx(path: Path, cells):
    wb = Workbook()
    ws = wb.active
    for coord, value in cells.items():
        ws[coord] = value
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def _docx_texts(path: Path):
    return [p.text for p in Document(str(path)).paragraphs]


def _batch_dirs(output: Path = Path("output")):
    if not output.exists():
        return []
    return sorted(p for p in output.iterdir() if p.is_dir())


def _load_mapping(batch: Path) -> dict:
    return json.loads((batch / "mapping.json").read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# CLI 1：H5 配置回退链
# ---------------------------------------------------------------------------

class TestConfigLoading:

    def test_explicit_config_missing_errors(self, runner):
        src = _make_docx(Path("a.docx"), ["无敏感内容"])
        result = runner.invoke(app, ["mask", str(src), "--config", "nope.yaml"])
        assert result.exit_code == 1
        assert "不存在" in _flat(result.output)

    def test_bare_cwd_fallback_warns_not_silent(self, runner, monkeypatch):
        """任意 CWD：回退到内嵌模板并明示词库为空，不静默（H5）。

        部署锚点同步隔离到 CWD：config_loader 的 runtime_anchor_dirs
        会兜底到源码树根，开发机上 tracked 的 config/sample_lexicon.yaml
        将使"词库为空"分支不可达，故模拟真正裸部署环境。
        """
        import mask_tool.core.config_loader as cl
        monkeypatch.setattr(cl, "runtime_anchor_dirs", lambda: [Path.cwd()])
        monkeypatch.setattr(cl, "writable_anchor_dir", lambda: Path.cwd())
        src = _make_docx(Path("a.docx"), ["联系电话13812345678"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0
        flat = _flat(result.output)
        assert "回退到内嵌默认配置模板" in flat
        assert "词库文件不存在" in flat
        # 词库为空时正则仍生效（手机号 AUTO）
        (out_docx,) = _batch_dirs(Path("out"))[0].glob("*.docx")
        assert _docx_texts(out_docx) == ["联系电话[CUSTOM_001]"]

    def test_config_command_generates_three_files(self, runner):
        result = runner.invoke(app, ["config"])
        assert result.exit_code == 0
        cfgdir = Path("config")
        for name, content in (
            ("default.yaml", DEFAULT_CONFIG_YAML),
            ("sample_lexicon.yaml", SAMPLE_LEXICON_YAML),
            ("whitelist.yaml", WHITELIST_YAML),
        ):
            assert (cfgdir / name).read_text(encoding="utf-8") == content
        # 幂等：再跑一次跳过已存在
        result2 = runner.invoke(app, ["config"])
        assert "跳过(已存在)" in _flat(result2.output)


# ---------------------------------------------------------------------------
# CLI 2：N2 词库自动复制
# ---------------------------------------------------------------------------

class TestLexiconAutocopy:

    def test_lexicon_autocopied_from_sample(self, runner):
        _write_config_dir(with_lexicon=False, with_sample=True)
        src = _make_docx(Path("a.docx"), ["提及某某建设集团有限公司"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0
        assert Path("config/lexicon.yaml").exists()
        assert "自动复制" in _flat(result.output)
        # 复制出的示例词库生效
        (out_docx,) = _batch_dirs(Path("out"))[0].glob("*.docx")
        assert "[COMPANY_001]" in _docx_texts(out_docx)[0]


# ---------------------------------------------------------------------------
# CLI 3：M4 批次目录
# ---------------------------------------------------------------------------

class TestBatchDirs:

    def test_two_masks_two_batch_dirs(self, runner):
        _write_config_dir()
        src = _make_docx(Path("a.docx"), ["张三"])
        for _ in range(2):
            result = runner.invoke(app, ["mask", str(src), "--output", "out"])
            assert result.exit_code == 0
        batches = _batch_dirs(Path("out"))
        assert len(batches) == 2
        assert batches[0].name != batches[1].name
        for b in batches:
            assert (b / "mapping.json").exists()
            assert (b / "report.json").exists()
        # 不再写平铺 output/mapping.json
        assert not (Path("out") / "mapping.json").exists()

    def test_batch_metadata_present(self, runner):
        _write_config_dir()
        src = _make_docx(Path("a.docx"), ["张三"])
        assert runner.invoke(app, ["mask", str(src), "--output", "out"]).exit_code == 0
        data = _load_mapping(_batch_dirs(Path("out"))[0])
        meta = data["metadata"]
        assert meta["batch_id"] == _batch_dirs(Path("out"))[0].name
        assert meta["created_at"] and meta["tool_version"]

    def test_single_file_keeps_masked_suffix(self, runner):
        _write_config_dir()
        src = _make_docx(Path("合同.docx"), ["张三"])
        assert runner.invoke(app, ["mask", str(src), "--output", "out"]).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        assert (batch / "合同_masked.docx").exists()


# ---------------------------------------------------------------------------
# CLI 4：H6 建议项默认不替换 / --all
# ---------------------------------------------------------------------------

class TestH6All:

    def test_suggest_default_vs_all(self, runner):
        _write_config_dir()
        # 正则金额 0.80 -> smart SUGGEST
        src = _make_docx(Path("b.docx"), ["预算500万元整"])
        assert runner.invoke(app, ["mask", str(src), "--output", "out1"]).exit_code == 0
        (out_docx,) = _batch_dirs(Path("out1"))[0].glob("*.docx")
        assert _docx_texts(out_docx) == ["预算500万元整"]

        assert runner.invoke(
            app, ["mask", str(src), "--output", "out2", "--all"],
        ).exit_code == 0
        (out_docx2,) = _batch_dirs(Path("out2"))[0].glob("*.docx")
        assert _docx_texts(out_docx2) == ["预算[AMOUNT_001]整"]


# ---------------------------------------------------------------------------
# CLI 5：PPT/PDF 屏蔽
# ---------------------------------------------------------------------------

class TestBlockedFormats:

    def test_single_blocked_file_exits_with_warning(self, runner):
        _write_config_dir()
        Path("x.pptx").write_bytes(b"PK-fake")
        result = runner.invoke(app, ["mask", "x.pptx", "--output", "out"])
        assert result.exit_code == 1
        flat = _flat(result.output)
        assert "PPT" in flat and "跳过" in flat

    def test_dir_mode_blocked_not_copied(self, runner):
        _write_config_dir()
        src_dir = Path("资料")
        src_dir.mkdir()
        _make_docx(src_dir / "a.docx", ["张三"])
        (src_dir / "b.pdf").write_bytes(b"%PDF-fake")
        (src_dir / "c.pptx").write_bytes(b"PK-fake")
        result = runner.invoke(app, ["mask", str(src_dir), "--output", "out"])
        assert result.exit_code == 0
        tree = _batch_dirs(Path("out"))[0] / src_dir.name
        assert (tree / "a.docx").exists()
        assert not (tree / "b.pdf").exists()
        assert not (tree / "c.pptx").exists()
        flat = _flat(result.output)
        assert "PDF脱敏暂未开放" in flat or "PDF 脱敏暂未开放" in flat


# ---------------------------------------------------------------------------
# CLI 6：M3 reserved / 对账扫描 / 指纹
# ---------------------------------------------------------------------------

class TestM3Protection:

    def _mask(self, runner, paragraphs):
        _write_config_dir()
        src = _make_docx(Path("m.docx"), paragraphs)
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        masked = batch / "m_masked.docx"
        return src, batch, masked

    def test_reserved_natural_token_not_confused(self, runner):
        """文档自然含 [PERSON_001] + mask 张三 -> 新 token 编号让位，unmask 不误还原。"""
        src, batch, masked = self._mask(runner, ["已勾选[PERSON_001]待审", "负责人张三"])
        texts = _docx_texts(masked)
        assert texts == ["已勾选[PERSON_001]待审", "负责人[PERSON_002]"]

        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        restored = Path("restored") / "m_masked.docx"
        # 自然 token 原样保留，张三正确还原
        assert _docx_texts(restored) == ["已勾选[PERSON_001]待审", "负责人张三"]

    def test_residual_in_mapping_errors_unless_force(self, runner, monkeypatch):
        """还原失败残留 ∈ mapping -> 报错退出；--force 忽略。"""
        src, batch, masked = self._mask(runner, ["负责人张三"])

        def broken_restore(f, out_path, token_map):
            # 模拟还原失败：文件照常落盘但内容未还原
            shutil.copy2(f, out_path)

        monkeypatch.setattr(cli_module, "_restore_file_content", broken_restore)
        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 1
        assert "还原后仍残留" in _flat(result.output)

        result2 = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored2", "--force",
        ])
        assert result2.exit_code == 0

    def test_foreign_token_warns_but_exits_zero(self, runner):
        """文档自然含 token 样式文本且不在映射表 -> 警告不中断。"""
        src, batch, masked = self._mask(runner, ["备注[COMPANY_009]待查", "负责人张三"])
        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        assert "不在本次映射表中" in _flat(result.output)

    def test_verify_fingerprint_tampering(self, runner):
        """--verify：篡改 original 触发指纹不匹配。"""
        src, batch, masked = self._mask(runner, ["负责人张三"])
        mapping_path = batch / "mapping.json"
        data = _load_mapping(batch)
        token = next(iter(data["tokens"]))
        data["tokens"][token]["original"] = "被篡改的原文"
        mapping_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(mapping_path),
            "--output", "r1", "--verify",
        ])
        assert result.exit_code == 1
        assert "指纹不匹配" in _flat(result.output)

        result2 = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(mapping_path),
            "--output", "r2", "--verify", "--force",
        ])
        assert result2.exit_code == 0

    def test_verify_ok_when_untampered(self, runner):
        src, batch, masked = self._mask(runner, ["负责人张三"])
        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "r", "--verify",
        ])
        assert result.exit_code == 0


# ---------------------------------------------------------------------------
# CLI 7：N1 inspect 盲区消除 + path 来源项
# ---------------------------------------------------------------------------

class TestInspect:

    def test_inspect_detects_table_content(self, runner):
        _write_config_dir()
        doc = Document()
        doc.add_paragraph("无敏感正文")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "联系人李四"
        doc.save("t.docx")
        result = runner.invoke(app, ["inspect", "t.docx"])
        assert result.exit_code == 0
        assert "李四" in _flat(result.output)

    def test_inspect_detects_filename_as_path_source(self, runner):
        _write_config_dir()
        _make_docx(Path("合同-张三.docx"), ["普通内容"])
        result = runner.invoke(app, ["inspect", "合同-张三.docx"])
        assert result.exit_code == 0
        flat = _flat(result.output)
        assert "张三" in flat
        assert "path" in flat

    def test_inspect_directory_recursive(self, runner):
        _write_config_dir()
        sub = Path("tree/sub")
        _make_docx(sub / "a.docx", ["张三"])
        result = runner.invoke(app, ["inspect", "tree"])
        assert result.exit_code == 0
        assert "张三" in _flat(result.output)


# ---------------------------------------------------------------------------
# --confirm 重构（H6/N3）
# ---------------------------------------------------------------------------

class TestConfirm:

    def test_write_masked_file_removed(self):
        assert not hasattr(cli_module, "_write_masked_file")

    def test_confirm_auto_yes_masks_checked(self, runner, monkeypatch):
        from mask_tool.core.confirm import ConfirmEngine

        monkeypatch.setattr(
            cli_module, "_new_confirm_engine", lambda: ConfirmEngine(auto_yes=True),
        )
        _write_config_dir()
        src = _make_docx(Path("c.docx"), ["负责人张三，预算500万元整"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out", "--confirm"])
        assert result.exit_code == 0
        (out_docx,) = _batch_dirs(Path("out"))[0].glob("*.docx")
        # 勾选项经 allowed_originals + statuses={AUTO,SUGGEST} 放行：
        # 词典 AUTO 的张三 + 正则 SUGGEST 的金额都替换
        texts = _docx_texts(out_docx)
        assert "[PERSON_001]" in texts[0]
        assert "[AMOUNT_001]" in texts[0]

    def test_confirm_reject_all_skips_file(self, runner, monkeypatch):
        class RejectAll:
            learned = []
            auto_yes = False

            def confirm_batch(self, results, file_name=""):
                return []

            def get_learned_words(self):
                return {}

        monkeypatch.setattr(cli_module, "_new_confirm_engine", RejectAll)
        _write_config_dir()
        src = _make_docx(Path("c.docx"), ["负责人张三"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out", "--confirm"])
        assert result.exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        assert list(batch.glob("*.docx")) == []


# ---------------------------------------------------------------------------
# 目录模式端到端：mask -> 断言 -> unmask -> 结构/内容比对
# ---------------------------------------------------------------------------

class TestDirectoryModeEndToEnd:

    def _build_tree(self):
        root = Path("项目资料-某某科技有限公司")
        sub = root / "张三的文件夹"
        sub.mkdir(parents=True)
        self.docx_paras = ["甲方：某某科技有限公司；负责人张三。"]
        _make_docx(root / "合同-张三.docx", self.docx_paras)
        _make_docx(sub / "内部说明.docx", ["抄送李四。"])
        _make_xlsx(root / "名单.xlsx", {"A1": "联系人李四"})
        (root / "说明.txt").write_text("普通文本文件", encoding="utf-8")
        (root / "跳过.pdf").write_bytes(b"%PDF-fake")
        return root

    def test_mask_unmask_roundtrip(self, runner):
        _write_config_dir()
        root = self._build_tree()

        result = runner.invoke(app, ["mask", str(root), "--output", "out"])
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]

        # 树根被改名（目录名含公司名）
        tree = batch / "项目资料-[COMPANY_001]"
        assert tree.is_dir()
        # 文件名脱敏：不加 _masked 后缀，保持原文件名（脱敏后）
        assert (tree / "合同-[PERSON_001].docx").exists()
        assert (tree / "[PERSON_001]的文件夹").is_dir()
        assert (tree / "[PERSON_001]的文件夹" / "内部说明.docx").exists()
        # 非文档原样拷贝；pptx/pdf 未拷贝
        assert (tree / "说明.txt").read_text(encoding="utf-8") == "普通文本文件"
        assert not (tree / "跳过.pdf").exists()

        # 内容脱敏
        masked_paras = _docx_texts(tree / "合同-[PERSON_001].docx")
        assert masked_paras == ["甲方：[COMPANY_001]；负责人[PERSON_001]。"]
        wb = load_workbook(str(tree / "名单.xlsx"))
        assert wb.active["A1"].value == "联系人[PERSON_002]"  # 李四

        # mapping.json 合并 paths 段
        data = _load_mapping(batch)
        assert len(data["paths"]) >= 3
        assert data["metadata"]["total_path_mappings"] == len(data["paths"])
        rels = {pm["rel_new"]: pm["rel_old"] for pm in data["paths"]}
        assert rels["."] == "项目资料-某某科技有限公司"
        assert rels["合同-[PERSON_001].docx"] == "合同-张三.docx"

        # unmask 还原
        result2 = runner.invoke(app, [
            "unmask", str(tree), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result2.exit_code == 0, result2.output
        restored = Path("restored") / "项目资料-某某科技有限公司"

        # 结构比对
        orig_files = sorted(
            p.relative_to(root).as_posix() for p in root.rglob("*")
        )
        restored_files = sorted(
            p.relative_to(restored).as_posix() for p in restored.rglob("*")
        )
        assert restored_files == [f for f in orig_files if f != "跳过.pdf"]

        # 内容比对（D 项逐字断言 + xlsx 数值/文本）
        assert _docx_texts(restored / "合同-张三.docx") == self.docx_paras
        assert _docx_texts(restored / "张三的文件夹" / "内部说明.docx") == ["抄送李四。"]
        wb2 = load_workbook(str(restored / "名单.xlsx"))
        assert wb2.active["A1"].value == "联系人李四"
        assert (restored / "说明.txt").read_text(encoding="utf-8") == "普通文本文件"

    def test_no_mask_names_option(self, runner):
        _write_config_dir()
        root = self._build_tree()
        result = runner.invoke(
            app, ["mask", str(root), "--output", "out", "--no-mask-names"],
        )
        assert result.exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / root.name
        assert tree.is_dir()
        # 名字未脱敏，内容仍脱敏
        assert (tree / "合同-张三.docx").exists()
        assert "[COMPANY_001]" in _docx_texts(tree / "合同-张三.docx")[0]
        data = _load_mapping(batch)
        assert data["paths"] == []


# ---------------------------------------------------------------------------
# unmask 单文件基础路径
# ---------------------------------------------------------------------------

class TestUnmaskSingleFile:

    def test_unmask_restores_content_and_name(self, runner):
        _write_config_dir()
        root = Path("d")
        root.mkdir()
        _make_docx(root / "合同-张三.docx", ["负责人张三"])
        assert runner.invoke(app, ["mask", str(root), "--output", "out"]).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / "d"
        masked_file = tree / "合同-[PERSON_001].docx"

        result = runner.invoke(app, [
            "unmask", str(masked_file), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        # paths 段唯一定位 -> 文件名一并还原
        assert (Path("restored") / "合同-张三.docx").exists()
        assert _docx_texts(Path("restored") / "合同-张三.docx") == ["负责人张三"]

    def test_unmask_xlsx_number_restore(self, runner):
        """xlsx 数值单元格（kind=number）还原为数值而非文本。"""
        _write_config_dir()
        wb = Workbook()
        wb.active["A1"] = 12000000
        wb.active["A1"].number_format = "¥#,##0.00"
        src = Path("n.xlsx")
        wb.save(src)
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        masked = batch / "n_masked.xlsx"
        assert load_workbook(str(masked)).active["A1"].value == "[AMOUNT_001]"

        result2 = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result2.exit_code == 0
        value = load_workbook(str(Path("restored") / "n_masked.xlsx")).active["A1"].value
        assert value == 12000000
        assert isinstance(value, int)

    def test_unmask_xlsx_16digit_card_restores_as_int(self, runner):
        """R3-B1 修订：16 位 Luhn 卡号 kind=number 还原为 int（下游公式
        引用不断裂；旧版降级文本的已知限制已随从严形态放宽到 19 位解除）。"""
        _write_config_dir()
        src = _make_xlsx(Path("card.xlsx"), {"A1": 4111111111111111})
        assert runner.invoke(app, ["mask", str(src), "--output", "out"]).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        masked = batch / "card_masked.xlsx"
        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        value = load_workbook(
            str(Path("restored") / "card_masked.xlsx")
        ).active["A1"].value
        assert isinstance(value, int)
        assert value == 4111111111111111
