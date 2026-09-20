# -*- coding: utf-8 -*-
"""tests/test_web_custom_words.py — I6：Web 临时自定义敏感词

覆盖（每项至少 1 例）：
- 解析函数 _parse_custom_words：换行/中英文逗号混合、空项、去重（保序）、空输入
- 合并模式：手动词注入后检测含该词命中项且状态 AUTO（0.95 档）；
  手动词与词库命中同一文本时去重为一条（source=manual）并置顶
- 仅手动模式：NER/正则/词库通道均不产出，检测结果只含手动词项
- 脱敏链路：_run_detection → _run_masking（同一套手动词注入）产物含
  [CUSTOM_xxx]，restore 还原原文（Web 函数级测试，参考 test_r3_fixes 写法）
- Detector 默认行为不变：manual_words=None / regex_enabled=True 时与旧签名等价
"""

import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from docx import Document

from mask_tool.core.templates import (
    DEFAULT_CONFIG_YAML, SAMPLE_LEXICON_YAML, WHITELIST_YAML,
)

TOKEN_RE = re.compile(
    r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]"
)

LEXICON = {
    "company": ["某某科技有限公司"],
    "person": ["张三", "李四"],
}


@pytest.fixture(autouse=True)
def _cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _write_config_dir(lexicon=None):
    cfgdir = Path("config")
    cfgdir.mkdir(parents=True, exist_ok=True)
    default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")  # NER 关
    (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
    (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
    (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
    (cfgdir / "lexicon.yaml").write_text(
        yaml.dump(lexicon or LEXICON, allow_unicode=True), encoding="utf-8",
    )


def _make_docx_bytes(paragraphs) -> bytes:
    buf = io.BytesIO()
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def _docx_paragraphs(path: Path):
    return [p.text for p in Document(str(path)).paragraphs]


def _upload(name: str, payload: bytes) -> SimpleNamespace:
    return SimpleNamespace(name=name, read=lambda: payload)


def _patch_streamlit(monkeypatch, web):
    """Web 函数级测试：把 st 展示/rerun 类调用全部静音（test_r3_fixes 同法）。"""
    # 拆分后（2026-09-20）：BATCHES_DIR/_add_history 的真实读写方在 ui.service
    from mask_tool.web.ui import service as _svc

    monkeypatch.setattr(_svc, "BATCHES_DIR", Path("batches"))
    monkeypatch.setattr(_svc, "_add_history", lambda rec: None)
    for fn in ("rerun", "error", "warning", "info", "success", "caption",
               "spinner"):
        monkeypatch.setattr(
            web.st, fn,
            (lambda *a, **k: None) if fn != "spinner" else
            (lambda *a, **k: _NullCtx()),
            raising=False,
        )
    for key in list(web.st.session_state.keys()):
        del web.st.session_state[key]


class _NullCtx:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


# ===========================================================================
# 解析函数
# ===========================================================================

class TestParseCustomWords:

    def test_mixed_separators_and_dedup(self):
        """换行/中文逗号/英文逗号混合 + 空项剔除 + 去重保序。"""
        web = pytest.importorskip("mask_tool.web.app")
        text = "绝密项目代号，张三丰\n \n李四光, ，绝密项目代号\n  王五  \n,,"
        assert web._parse_custom_words(text) == [
            "绝密项目代号", "张三丰", "李四光", "王五",
        ]

    def test_empty_and_blank(self):
        """空输入 / 纯分隔符输入 → []（现状不变语义）。"""
        web = pytest.importorskip("mask_tool.web.app")
        assert web._parse_custom_words("") == []
        assert web._parse_custom_words(None) == []
        assert web._parse_custom_words(" \n ,, \n") == []


# ===========================================================================
# 检测语义：合并模式 / 仅手动模式
# ===========================================================================

class TestDetectionSemantics:

    def _detect(self, monkeypatch, paragraphs, manual_words, manual_only=False):
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)
        web._run_detection(
            [_upload("a.docx", _make_docx_bytes(paragraphs))],
            "smart", ner_enabled=False,
            manual_words=manual_words, manual_only=manual_only,
        )
        return web.st.session_state["detection_results"]

    def test_merge_mode_manual_word_auto_status(self, monkeypatch):
        """合并模式：手动词命中项 source=manual、置信度 0.95、状态 AUTO；
        词库项照常存在；手动词条目置顶。"""
        results = self._detect(
            monkeypatch,
            ["负责人李四对接某某科技有限公司", "内部代号绝密项目代号待处理"],
            manual_words=["绝密项目代号"],
        )
        by_source = {}
        for r in results:
            by_source.setdefault(r.source, []).append(r)
        manual = by_source.get("manual", [])
        assert len(manual) == 1
        m = manual[0]
        assert m.text == "绝密项目代号"
        assert m.confidence == 0.95
        assert m.status.value == "auto_mask"
        assert m.text_type.value == "custom"
        # 词库项照常（合并语义）
        assert any(r.text == "李四" and r.source == "dictionary" for r in results)
        assert any(
            r.text == "某某科技有限公司" and r.source == "dictionary"
            for r in results
        )
        # 手动词置顶（稳定排序后 manual 在首位）
        assert results[0].source == "manual"

    def test_merge_mode_dedup_against_lexicon(self, monkeypatch):
        """手动词与词库命中同一文本：仅一条，标 manual（一条为准）。"""
        # "张三" 同时在词库 person 与手动词中
        results = self._detect(
            monkeypatch, ["负责人张三今天离职"],
            manual_words=["张三"],
        )
        hits = [r for r in results if r.text == "张三"]
        assert len(hits) == 1
        assert hits[0].source == "manual"
        assert hits[0].text_type.value == "custom"

    def test_manual_only_mode_skips_all_auto_channels(self, monkeypatch):
        """仅手动模式：词库/正则通道均不产出（NER 亦不跑），只含手动词项。

        文本含词库词（李四/某某科技有限公司）、正则可命中项（手机号），
        manual_only 下全部不得出现。
        """
        results = self._detect(
            monkeypatch,
            ["负责人李四联系13812345678，某某科技有限公司，代号绝密项目代号"],
            manual_words=["绝密项目代号"],
            manual_only=True,
        )
        assert [r.text for r in results] == ["绝密项目代号"]
        assert all(r.source == "manual" for r in results)

    def test_manual_only_mode_pipeline_channels_closed(self, monkeypatch):
        """仅手动模式下 Detector 三通道全关：lexicon 空、正则空、NER 不初始化。"""
        from mask_tool.core.pipeline import Pipeline
        from mask_tool.models.config import MaskConfig

        _write_config_dir()
        cfg = MaskConfig(mode="smart")
        cfg.ner.enabled = True  # 即使侧边栏开着 NER
        cfg.lexicon_path = str(Path("config/lexicon.yaml").resolve())
        cfg.whitelist_path = str(Path("config/whitelist.yaml").resolve())
        p = Pipeline(cfg, manual_words=["甲"], auto_detect_enabled=False)
        assert p.detector.ner_engine is None
        assert p.detector.lexicon == {}
        assert p.detector._regex_rules == []
        assert p.detector.manual_words == ["甲"]

    def test_detector_default_behavior_unchanged(self):
        """旧签名等价：manual_words=None / regex_enabled=True 行为不变。"""
        from mask_tool.core.detector import Detector

        d_old = Detector(dict(LEXICON), set())
        d_new = Detector(dict(LEXICON), set(), manual_words=None,
                         regex_enabled=True)
        text = "李四在某某科技有限公司，电话13812345678"
        assert (
            [(r.text, r.source, r.confidence) for r in d_old.detect(text)]
            == [(r.text, r.source, r.confidence) for r in d_new.detect(text)]
        )
        # 手动词不受白名单过滤（用户显式指定优先）
        d_wl = Detector({}, {"绝密项目代号"}, manual_words=["绝密项目代号"])
        assert [r.text for r in d_wl.detect("代号绝密项目代号")] == ["绝密项目代号"]


# ===========================================================================
# 脱敏链路：检测页勾选手动词 → 产物 [CUSTOM_xxx] → restore 还原
# ===========================================================================

class TestMaskingChainWithManualWords:

    def test_detect_then_mask_then_restore(self, monkeypatch):
        """全链路：_run_detection（手动词+词库同文本去重）→ 勾选手动词项
        → _run_masking（同一套手动词注入）→ 产物含 [CUSTOM_001] →
        restore_file_content 还原。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)

        manual = ["绝密项目代号"]
        upload = _upload(
            "a.docx",
            _make_docx_bytes(["内部代号绝密项目代号，负责人李四对接"]),
        )

        # 检测（合并模式）
        web._run_detection(
            [upload], "smart", ner_enabled=False,
            manual_words=manual, manual_only=False,
        )
        all_results = web.st.session_state["detection_results"]
        idx = next(i for i, r in enumerate(all_results) if r.source == "manual")
        # 同文本去重：词库无重复项（"绝密项目代号"只出现一次）
        assert sum(1 for r in all_results if r.text == "绝密项目代号") == 1

        # 脱敏（与检测同一套手动词——Web 端两处均取自同一控件值）
        web._run_masking(
            [upload], [idx], all_results,
            mode="smart", ner_enabled=False, irreversible=False,
            learn_words=False, batch_id="MSK-I6", batch_name="t",
            mask_filenames=False, manual_words=manual, manual_only=False,
        )

        result = web.st.session_state["mask_result"]
        tokens = json.loads(result["mapping_data"])["tokens"]
        custom_tokens = [t for t in tokens if t.startswith("[CUSTOM_")]
        assert len(custom_tokens) == 1
        assert tokens[custom_tokens[0]]["original"] == "绝密项目代号"
        assert re.fullmatch(r"\[CUSTOM_\d{3,}\]", custom_tokens[0])

        # 产物正文：手动词已替换，未勾选的词库词保留
        out_docx = next(
            Path("batches", "MSK-I6").glob("*.docx")
        )
        paras = _docx_paragraphs(out_docx)
        assert paras == ["内部代号[CUSTOM_001]，负责人李四对接"]

        # restore 还原（产品链路）
        from mask_tool.adapters.restore import restore_file_content

        restored = Path("restored") / "a.docx"
        restore_file_content(out_docx, restored, tokens)
        assert _docx_paragraphs(restored) == ["内部代号绝密项目代号，负责人李四对接"]

    def test_manual_only_masking_no_auto_items(self, monkeypatch):
        """仅手动模式脱敏：勾选手动词后产物只含手动词 token，
        词库词（李四）不替换。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)

        manual = ["内部代号X7"]
        upload = _upload(
            "a.docx",
            _make_docx_bytes(["内部代号X7与李四无关"]),
        )

        web._run_detection(
            [upload], "smart", ner_enabled=True,
            manual_words=manual, manual_only=True,
        )
        all_results = web.st.session_state["detection_results"]
        # 检测面只有手动词项
        assert [(r.text, r.source) for r in all_results] == [
            ("内部代号X7", "manual"),
        ]
        web._run_masking(
            [upload], [0], all_results,
            mode="smart", ner_enabled=True, irreversible=False,
            learn_words=False, batch_id="MSK-I6B", batch_name="t",
            mask_filenames=False, manual_words=manual, manual_only=True,
        )
        out_docx = next(
            Path("batches", "MSK-I6B").glob("*.docx")
        )
        assert _docx_paragraphs(out_docx) == ["[CUSTOM_001]与李四无关"]
        # mapping 只含手动词条目
        tokens = json.loads(
            web.st.session_state["mask_result"]["mapping_data"]
        )["tokens"]
        assert list(tokens) == ["[CUSTOM_001]"]

    def test_empty_manual_words_keeps_current_behavior(self, monkeypatch):
        """空手动词 + 开关全关 = 现状不变：检测/脱敏与旧流程等价。"""
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        _patch_streamlit(monkeypatch, web)

        upload = _upload(
            "a.docx",
            _make_docx_bytes(["负责人李四对接某某科技有限公司"]),
        )
        web._run_detection(
            [upload], "smart", ner_enabled=False,
            manual_words=[], manual_only=False,
        )
        all_results = web.st.session_state["detection_results"]
        assert all(r.source != "manual" for r in all_results)
        assert any(r.text == "李四" for r in all_results)
