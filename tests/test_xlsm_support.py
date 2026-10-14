# -*- coding: utf-8 -*-
"""tests/test_xlsm_support.py — xlsm 格式接入测试（方案 A，2026-10-14）

覆盖：
- 样本构造：xlsx 重打包为 xlsm（macroEnabled content type + vbaProject.bin），
  openpyxl keep_vba 不解析宏部件、保存时原样复制（字节级不变的基础）
- 检测面：extract_texts / detect_file_results 覆盖 .xlsm（检测面=处理面）
- 脱敏：XlsxAdapter 输出保持 .xlsm 后缀、宏部件字节级不变、敏感词已替换
- pipeline：process_file 对 .xlsm 走 xlsx 路径出产物
- 还原：restore_file_content .xlsm 分支内容还原 + 宏部件保留
- 预检：zip 含 vbaProject.bin 时输出宏边界告警
"""

import zipfile
from io import BytesIO
from logging import WARNING
from pathlib import Path

from openpyxl import Workbook, load_workbook

from mask_tool.adapters.extract import detect_file_results, extract_texts
from mask_tool.adapters.restore import restore_file_content
from mask_tool.adapters.xlsx_adapter import XlsxAdapter
from mask_tool.core.detector import Detector
from mask_tool.core.masker import Masker
from mask_tool.core.policy import PolicyEngine
from mask_tool.core.tokenizer import TokenGenerator
from mask_tool.models.config import MaskConfig

VBA_BYTES = b"FAKE_VBA_PROJECT_BIN_\x00\x01\x02"

DEFAULT_LEXICON = {
    "person": ["李四"],
    "company": ["某某科技有限公司"],
}


# ---------------------------------------------------------------------------
# 样本构造：xlsx → xlsm（zip 层植入宏部件）
# ---------------------------------------------------------------------------

def _make_xlsx_bytes(cells: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    for ref, value in cells.items():
        ws[ref] = value
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _make_xlsm_bytes(cells: dict, vba: bytes = VBA_BYTES) -> bytes:
    """xlsx 重打包为 xlsm：workbook content type 改 macroEnabled +
    bin Default 声明 + 植入 vbaProject.bin（openpyxl keep_vba 原样复制）。"""
    src = zipfile.ZipFile(BytesIO(_make_xlsx_bytes(cells)))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = text.replace(
                    "application/vnd.ms-excel.sheet.main+xml",
                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
                text = text.replace(
                    '<Default Extension="xml"',
                    '<Default Extension="bin" '
                    'ContentType="application/vnd.ms-office.vbaProject"/>'
                    '<Default Extension="xml"',
                )
                data = text.encode("utf-8")
            dst.writestr(name, data)
        dst.writestr("xl/vbaProject.bin", vba)
    src.close()
    return out.getvalue()


def _read_vba(path: Path):
    with zipfile.ZipFile(path) as z:
        if "xl/vbaProject.bin" in z.namelist():
            return z.read("xl/vbaProject.bin")
    return None


def _make_adapter(lexicon=None):
    detector = Detector(
        {k: list(v) for k, v in (lexicon or DEFAULT_LEXICON).items()}, set(),
    )
    policy = PolicyEngine(MaskConfig(mode="smart"))
    masker = Masker(TokenGenerator())
    return XlsxAdapter(detector, policy, masker), masker


SAMPLE_CELLS = {
    "A1": "项目负责人：李四",
    "B2": "供应商：某某科技有限公司",
    "C3": "普通文本，不应被改动",
}


# ---------------------------------------------------------------------------
# 检测面
# ---------------------------------------------------------------------------

class TestDetectSurface:
    def test_extract_texts_covers_xlsm(self, tmp_path):
        f = tmp_path / "带宏.xlsm"
        f.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        text = extract_texts(f)
        assert "李四" in text and "某某科技有限公司" in text
        assert "普通文本" in text

    def test_detect_file_results_covers_xlsm(self, tmp_path):
        f = tmp_path / "带宏.xlsm"
        f.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        adapter, _ = _make_adapter()
        results = detect_file_results(f, adapter.detector, adapter.policy)
        texts = {r.text for r in results}
        assert "李四" in texts and "某某科技有限公司" in texts


# ---------------------------------------------------------------------------
# 脱敏：宏部件字节级不变 + 敏感词替换
# ---------------------------------------------------------------------------

class TestMaskKeepsVba:
    def test_mask_output_and_vba_integrity(self, tmp_path):
        src = tmp_path / "合同.xlsm"
        src.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        adapter, masker = _make_adapter()

        out = adapter.process(src, tmp_path / "out")

        assert out.suffix == ".xlsm", "输出必须保持 xlsm 后缀（否则宏丢失）"
        assert out.name == "合同_masked.xlsm"
        assert _read_vba(out) == VBA_BYTES, "宏部件被改写"

        wb = load_workbook(out)
        ws = wb.active
        assert "李四" not in str(ws["A1"].value)
        assert "某某科技有限公司" not in str(ws["B2"].value)
        assert ws["C3"].value == "普通文本，不应被改动"
        assert masker.mappings, "应产生映射记录"

    def test_output_name_passthrough_keeps_suffix(self, tmp_path):
        src = tmp_path / "合同.xlsm"
        src.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out",
                              output_name="改名后.xlsm")
        assert out.name == "改名后.xlsm"

    def test_precheck_warns_vba(self, tmp_path, caplog):
        """zip 预检对宏部件输出边界告警（宏内容不参与脱敏）。"""
        import logging

        src = tmp_path / "合同.xlsm"
        src.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        adapter, _ = _make_adapter()
        with caplog.at_level(logging.WARNING, logger="mask_tool"):
            adapter._precheck_zip(src)
        assert any("VBA" in r.message or "宏" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# pipeline 分发
# ---------------------------------------------------------------------------

class TestPipelineDispatch:
    def test_process_file_routes_xlsm(self, tmp_path):
        from mask_tool.core.pipeline import Pipeline

        lex = tmp_path / "lex.yaml"
        lex.write_text(
            "person:\n  - 李四\n", encoding="utf-8",
        )
        wl = tmp_path / "wl.yaml"
        wl.write_text("whitelist: []\n", encoding="utf-8")
        cfg = MaskConfig(mode="smart")
        pipeline = Pipeline(cfg, lexicon_path=str(lex), whitelist_path=str(wl))

        src = tmp_path / "合同.xlsm"
        src.write_bytes(_make_xlsm_bytes({"A1": "负责人李四"}))
        out = pipeline.process_file(src, tmp_path / "batch")

        assert out is not None and out.suffix == ".xlsm"
        assert _read_vba(out) == VBA_BYTES
        wb = load_workbook(out)
        assert "李四" not in str(wb.active["A1"].value)


# ---------------------------------------------------------------------------
# 还原：内容还原 + 宏部件保留
# ---------------------------------------------------------------------------

class TestRestore:
    def test_restore_roundtrip_keeps_vba(self, tmp_path):
        src = tmp_path / "合同.xlsm"
        src.write_bytes(_make_xlsm_bytes(SAMPLE_CELLS))
        adapter, masker = _make_adapter()
        masked = adapter.process(src, tmp_path / "out")

        token_map = {
            m.token: {"original": m.original} for m in masker.mappings
        }
        assert token_map, "无映射可还原"

        restored = tmp_path / "合同_restored.xlsm"
        restore_file_content(masked, restored, token_map)

        assert _read_vba(restored) == VBA_BYTES, "还原产物宏部件丢失"
        wb = load_workbook(restored)
        ws = wb.active
        assert "李四" in str(ws["A1"].value)
        assert "某某科技有限公司" in str(ws["B2"].value)
        assert ws["C3"].value == "普通文本，不应被改动"

    def test_restore_rejects_unknown_suffix(self, tmp_path):
        f = tmp_path / "a.pptx"
        f.write_bytes(b"x")
        try:
            restore_file_content(f, tmp_path / "b.pptx", {})
        except ValueError as e:
            assert "不支持的还原格式" in str(e)
        else:
            raise AssertionError("应拒绝不支持的还原格式")
