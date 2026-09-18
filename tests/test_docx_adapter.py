"""tests/test_docx_adapter.py — Word 适配器重写版测试（D1 §8 docx 1-5 + D 项专项）

覆盖：多 run 段落完整（H1）、重叠无乱码、全部件覆盖面 fixture（表格/嵌套表格/
页眉/页脚/首页页眉/脚注/尾注/批注/文本框/超链接）、core properties 清空、
mask->unmask roundtrip、相邻实体逐字正确（波次1 发现的丢字符 bug 回归）、
H6 状态过滤、xml:space 保留、output_name 镜像树语义。
"""

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from mask_tool.adapters.docx_adapter import DocxAdapter, clear_core_properties
from mask_tool.core.detector import Detector
from mask_tool.core.masker import Masker
from mask_tool.core.policy import PolicyEngine
from mask_tool.core.tokenizer import TokenGenerator
from mask_tool.models.config import MaskConfig
from mask_tool.models.detection import DetectionStatus

TOKEN_RE = re.compile(r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]")

DEFAULT_LEXICON = {
    "company": ["某某科技有限公司"],
    "person": ["张三", "李四"],
}


def _make_adapter(lexicon=None, mode="smart", irreversible=False):
    detector = Detector({k: list(v) for k, v in (lexicon or DEFAULT_LEXICON).items()}, set())
    policy = PolicyEngine(MaskConfig(mode=mode))
    token_gen = TokenGenerator()
    masker = Masker(token_gen, irreversible=irreversible)
    return DocxAdapter(detector, policy, masker), masker


def _para_texts(path: Path):
    """用与适配器同源的 walker 提取输出文档全部段落文本。"""
    return DocxAdapter.extract_paragraph_texts(Path(path))


def _token_map(masker):
    return {m.token: {"original": m.original} for m in masker.mappings}


# ---------------------------------------------------------------------------
# 覆盖面 fixture：zip 注入 footnotes/endnotes/comments 部件
# ---------------------------------------------------------------------------

_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'

FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<w:footnotes {_W_NS}>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t></w:t></w:r></w:p></w:footnote>'
    '<w:footnote w:id="2"><w:p><w:r><w:t>脚注：张三经手</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
)
ENDNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<w:endnotes {_W_NS}>'
    '<w:endnote w:id="2"><w:p><w:r><w:t>尾注：某某科技有限公司承建</w:t></w:r></w:p></w:endnote>'
    '</w:endnotes>'
)
COMMENTS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<w:comments {_W_NS}>'
    '<w:comment w:id="0" w:author="王五" w:date="2024-01-01T00:00:00Z" w:initials="WW">'
    '<w:p><w:r><w:t>批注：请李四复核</w:t></w:r></w:p></w:comment>'
    '</w:comments>'
)

_PART_SPECS = {
    "footnotes": (
        "/word/footnotes.xml", FOOTNOTES_XML,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        "footnotes",
    ),
    "endnotes": (
        "/word/endnotes.xml", ENDNOTES_XML,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        "endnotes",
    ),
    "comments": (
        "/word/comments.xml", COMMENTS_XML,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        "comments",
    ),
}


def _inject_parts(src: Path, dst: Path, parts=("footnotes", "endnotes", "comments")) -> Path:
    """向已保存的 docx 注入 footnotes/endnotes/comments 部件（含 CT/rels 注册）。"""
    with zipfile.ZipFile(src) as zin:
        items = {n: zin.read(n) for n in zin.namelist()}
    ct_add, rel_add = [], []
    for key in parts:
        partname, xml, ctype, rel_target = _PART_SPECS[key]
        items[partname.lstrip("/")] = xml.encode("utf-8")
        ct_add.append(f'<Override PartName="{partname}" ContentType="{ctype}"/>')
        rel_add.append(
            f'<Relationship Id="rIdX{key}" Type="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships/{rel_target}" Target="{partname.rsplit("/", 1)[-1]}"/>'
        )
    ct = items["[Content_Types].xml"].decode()
    items["[Content_Types].xml"] = ct.replace("</Types>", "".join(ct_add) + "</Types>").encode()
    rels = items["word/_rels/document.xml.rels"].decode()
    items["word/_rels/document.xml.rels"] = rels.replace(
        "</Relationships>", "".join(rel_add) + "</Relationships>",
    ).encode()
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in items.items():
            zout.writestr(n, d)
    return dst


def _build_full_docx(path: Path) -> Path:
    """构造覆盖面 fixture：正文/表格/嵌套表格/页眉/页脚/首页页眉/
    文本框（含外层段落）/超链接，再注入脚注/尾注/批注部件。"""
    doc = Document()
    doc.add_paragraph("正文：某某科技有限公司与张三签约。")

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "联系人"
    table.cell(0, 1).text = "张三"
    inner = table.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "嵌套表格：李四"

    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.header.paragraphs[0].text = "页眉：李四"
    sec.footer.paragraphs[0].text = "页脚：张三"
    fph = sec.first_page_header
    fph.is_linked_to_previous = False
    fph.paragraphs[0].text = "首页页眉：李四"

    # 文本框：外层段落自身文本 + VML 文本框内嵌套段落
    p = doc.add_paragraph()
    p.add_run("外层段落：李四负责")
    pict = (
        f'<w:r {_W_NS} xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:pict><v:shape style=\"width:100pt;height:20pt\"><v:textbox>"
        "<w:txbxContent><w:p><w:r><w:t>文本框：张三</w:t></w:r></w:p></w:txbxContent>"
        "</v:textbox></v:shape></w:pict></w:r>"
    )
    p._p.append(parse_xml(pict))

    # 超链接内文本
    hyper = (
        f'<w:hyperlink {_W_NS} {_R_NS} r:id="rIdHl1">'
        "<w:r><w:t>链接文字：张三</w:t></w:r></w:hyperlink>"
    )
    doc.add_paragraph()._p.append(parse_xml(hyper))

    raw = path.with_suffix(".raw.docx")
    doc.save(raw)
    _inject_parts(raw, path)
    raw.unlink()
    return path


# ---------------------------------------------------------------------------
# D 项专项：相邻实体逐字正确（旧链路丢字符回归）
# ---------------------------------------------------------------------------

class TestAdjacentEntities:
    """波次1 发现的旧链路 bug：单 run 段落「甲方：某某科技有限公司；负责人张三。」
    输出丢字符（`[COMPANY_001[PERSON_001]`）。新链路必须逐字正确，含「；负责」。"""

    def test_single_run_adjacent_exact(self, tmp_path):
        doc = Document()
        doc.add_paragraph("甲方：某某科技有限公司；负责人张三。")
        src = tmp_path / "adj.docx"
        doc.save(src)

        adapter, masker = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        texts = _para_texts(out)
        assert texts == ["甲方：[COMPANY_001]；负责人[PERSON_001]。"]
        assert "；负责" in texts[0]

    def test_multi_run_adjacent_exact(self, tmp_path):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("甲方：")
        p.add_run("某某科技有限公司；负")
        p.add_run("责人张三。")
        src = tmp_path / "adj3.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        texts = _para_texts(out)
        assert texts == ["甲方：[COMPANY_001]；负责人[PERSON_001]。"]


# ---------------------------------------------------------------------------
# docx 1：H1 回归——替换区间横跨 run2-run3，段落全文完整
# ---------------------------------------------------------------------------

class TestMultiRunSpan:

    def test_span_across_runs_full_text_kept(self, tmp_path):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("前置说明：")
        p.add_run("某某科技")
        p.add_run("有限公司承建本项目")
        p.add_run("，尾部说明。")
        src = tmp_path / "span.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")

        d2 = Document(str(out))
        para = d2.paragraphs[0]
        assert para.text == "前置说明：[COMPANY_001]承建本项目，尾部说明。"

        # 首节点承载整段，其余清空但元素保留（rPr 不丢失）
        nodes = DocxAdapter._paragraph_text_nodes(para._p)
        assert [t.text or "" for t in nodes] == [
            "前置说明：[COMPANY_001]承建本项目，尾部说明。", "", "", "",
        ]

    def test_last_run_not_dropped(self, tmp_path):
        """旧 pos_map 缺陷：末 run old_end==len 永不在映射中 -> 整体置空。回归锁定。"""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("开头，")
        p.add_run("某某科技有限公司")
        p.add_run("结尾")
        src = tmp_path / "tail.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        assert _para_texts(out) == ["开头，[COMPANY_001]结尾"]


# ---------------------------------------------------------------------------
# docx 2：重叠检测结果无乱码（无嵌套 token 残片）
# ---------------------------------------------------------------------------

class TestOverlapNoGarbage:

    def test_nested_lexicon_entries_single_token(self, tmp_path):
        lexicon = {"company": ["某某科技有限公司", "科技公司"]}
        doc = Document()
        doc.add_paragraph("乙方某某科技有限公司盖章")
        doc.add_paragraph("简称科技公司")
        src = tmp_path / "nest.docx"
        doc.save(src)

        adapter, _ = _make_adapter(lexicon)
        out = adapter.process(src, tmp_path / "out")
        texts = _para_texts(out)

        # 长区间优先：首段整名一个 token，被吞子串不产生残片；
        # 次段“科技公司”是独立原文，获得自己的 token
        assert texts == ["乙方[COMPANY_001]盖章", "简称[COMPANY_002]"]
        for t in texts:
            # 无嵌套 token 残片（如 "[COMPANY_001[PERSON" / "]]"）
            assert not re.search(r"\[[A-Z]+_\d+\[", t)
            assert "]]" not in t

    def test_regex_overlap_resolved_by_engine(self, tmp_path):
        """19 位数字串旧版产生身份证+银行卡两条重叠；M6+引擎兜底后仅一条。"""
        doc = Document()
        doc.add_paragraph("卡号1234567890123456789结尾")
        src = tmp_path / "digits.docx"
        doc.save(src)

        adapter, masker = _make_adapter()  # 通用正则路径；smart 下银行卡 0.65 为
        # SUGGEST，需放宽 statuses 才进入替换（H6），同时验证重叠仅出一条
        masker.engine.active_statuses = frozenset(
            {DetectionStatus.AUTO_MASK, DetectionStatus.SUGGEST_MASK}
        )
        out = adapter.process(src, tmp_path / "out")
        texts = _para_texts(out)
        assert texts == ["卡号[CUSTOM_001]结尾"]
        assert "]]" not in texts[0]


# ---------------------------------------------------------------------------
# docx 3：覆盖面 fixture
# ---------------------------------------------------------------------------

class TestCoverage:

    @pytest.fixture()
    def fixture_docx(self, tmp_path):
        return _build_full_docx(tmp_path / "full.docx")

    def test_all_parts_masked(self, fixture_docx, tmp_path):
        adapter, masker = _make_adapter()
        out = adapter.process(fixture_docx, tmp_path / "out")
        texts = _para_texts(out)

        joined = "\n".join(texts)
        # 敏感词全部消失
        for word in ("某某科技有限公司", "张三", "李四"):
            assert word not in joined, f"{word} 残留于: {texts}"
        # 各部件都有 token
        assert any(t.startswith("正文：") and "COMPANY" in t for t in texts), texts
        assert any(t.startswith("页眉：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("页脚：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("首页页眉：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("联系人") for t in texts), texts                      # 表格
        assert any(t.startswith("嵌套表格：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("脚注：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("尾注：") and "COMPANY" in t for t in texts), texts
        assert any(t.startswith("批注：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("文本框：") and "PERSON" in t for t in texts), texts
        assert any(t.startswith("链接文字：") and "PERSON" in t for t in texts), texts

    def test_textbox_inner_paragraph_owned_by_itself(self, fixture_docx, tmp_path):
        """外层段落不得吞并文本框内层段落的文本（归属规则）。"""
        adapter, _ = _make_adapter()
        out = adapter.process(fixture_docx, tmp_path / "out")
        d2 = Document(str(out))
        body_ps = [
            "".join(t.text or "" for t in DocxAdapter._paragraph_text_nodes(p._p))
            for p in d2.paragraphs
        ]
        outer = next(t for t in body_ps if t.startswith("外层段落"))
        assert "文本框" not in outer
        assert re.fullmatch(r"外层段落：\[PERSON_\d{3}\]负责", outer)

    def test_comment_author_cleared(self, fixture_docx, tmp_path):
        adapter, _ = _make_adapter()
        out = adapter.process(fixture_docx, tmp_path / "out")
        doc2 = Document(str(out))
        for part in doc2.part.package.iter_parts():
            if str(part.partname) == "/word/comments.xml":
                for c in part.element.iter(qn("w:comment")):
                    assert c.get(qn("w:author")) == ""
                    assert c.get(qn("w:initials")) == ""
                return
        pytest.fail("comments 部件丢失")

    def test_extract_paragraph_texts_covers_all(self, fixture_docx):
        """inspect/confirm 同源提取与 mask 覆盖面一致。"""
        texts = DocxAdapter.extract_paragraph_texts(fixture_docx)
        joined = "\n".join(texts)
        for probe in ("正文：", "页眉：李四", "脚注：张三", "尾注：", "批注：请李四",
                      "文本框：张三", "链接文字：张三", "嵌套表格：李四"):
            assert probe in joined, f"提取盲区: {probe}"


# ---------------------------------------------------------------------------
# docx 4：core properties 清空
# ---------------------------------------------------------------------------

class TestCoreProperties:

    def test_core_properties_cleared(self, tmp_path):
        from docx import Document

        doc = Document()
        doc.add_paragraph("无敏感内容")
        cp = doc.core_properties
        cp.author = "王五"
        cp.last_modified_by = "赵六"
        cp.comments = "内部草稿"
        cp.subject = "合同"
        cp.keywords = "机密"
        cp.category = "合同类"
        cp.title = "某项目合同"
        src = tmp_path / "cp.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")

        cp2 = Document(str(out)).core_properties
        for attr in ("author", "last_modified_by", "comments", "subject",
                     "keywords", "category", "title"):
            assert getattr(cp2, attr) == "", f"{attr} 未清空: {getattr(cp2, attr)!r}"

    def test_clear_core_properties_idempotent(self, tmp_path):
        doc = Document()
        doc.add_paragraph("x")
        clear_core_properties(doc)
        clear_core_properties(doc)
        assert doc.core_properties.author == ""


# ---------------------------------------------------------------------------
# docx 5：mask -> unmask roundtrip
# ---------------------------------------------------------------------------

class TestRoundtrip:

    def test_roundtrip_full_fixture(self, tmp_path):
        src = _build_full_docx(tmp_path / "rt.docx")
        original_texts = DocxAdapter.extract_paragraph_texts(src)

        adapter, masker = _make_adapter()
        out = adapter.process(src, tmp_path / "out")

        token_map = _token_map(masker)
        assert token_map, "无映射产生"
        restored = tmp_path / "restored" / "rt.docx"
        hits = DocxAdapter.restore_document(out, restored, token_map)
        assert hits == set(token_map)

        assert DocxAdapter.extract_paragraph_texts(restored) == original_texts

    def test_roundtrip_plain(self, tmp_path):
        doc = Document()
        doc.add_paragraph("甲方：某某科技有限公司；负责人张三。")
        doc.add_paragraph("无敏感段落")
        src = tmp_path / "p.docx"
        doc.save(src)
        original = DocxAdapter.extract_paragraph_texts(src)

        adapter, masker = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        restored = tmp_path / "r.docx"
        DocxAdapter.restore_document(out, restored, _token_map(masker))
        assert DocxAdapter.extract_paragraph_texts(restored) == original


# ---------------------------------------------------------------------------
# H6 状态过滤 与 xml:space / output_name
# ---------------------------------------------------------------------------

class TestStatusesAndDetails:

    def test_suggest_not_replaced_by_default(self, tmp_path):
        """smart 模式 0.80 置信（正则金额）为 SUGGEST_MASK，默认不替换（H6）。"""
        doc = Document()
        doc.add_paragraph("预算500万元整")
        src = tmp_path / "sug.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        assert _para_texts(out) == ["预算500万元整"]

    def test_suggest_replaced_when_active_statuses_widened(self, tmp_path):
        doc = Document()
        doc.add_paragraph("预算500万元整")
        src = tmp_path / "sug2.docx"
        doc.save(src)

        adapter, masker = _make_adapter()
        masker.engine.active_statuses = frozenset(
            {DetectionStatus.AUTO_MASK, DetectionStatus.SUGGEST_MASK}
        )
        out = adapter.process(src, tmp_path / "out")
        texts = _para_texts(out)
        # 金额正则贪吃至“元”：500万元 整体被替换
        assert texts == ["预算[AMOUNT_001]整"]

    def test_allowed_originals_respected(self, tmp_path):
        """confirm 模式：只替换勾选项（allowed_originals 精确放行）。"""
        doc = Document()
        doc.add_paragraph("某某科技有限公司与张三")
        src = tmp_path / "allow.docx"
        doc.save(src)

        adapter, masker = _make_adapter()
        masker.engine.active_allowed_originals = {"张三"}
        out = adapter.process(src, tmp_path / "out")
        assert _para_texts(out) == ["某某科技有限公司与[PERSON_001]"]

    def test_xml_space_preserved_on_rewrite(self, tmp_path):
        doc = Document()
        doc.add_paragraph("  前后空格 张三  ")
        src = tmp_path / "sp.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        d2 = Document(str(out))
        p = d2.paragraphs[0]
        assert p.text == "  前后空格 [PERSON_001]  "
        first_t = DocxAdapter._paragraph_text_nodes(p._p)[0]
        assert first_t.get(qn("xml:space")) == "preserve"

    def test_output_name_keeps_original_name(self, tmp_path):
        doc = Document()
        doc.add_paragraph("张三")
        src = tmp_path / "orig.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out", output_name="orig.docx")
        assert out.name == "orig.docx"
        # 默认（无 output_name）保持 _masked 后缀
        out2 = adapter.process(src, tmp_path / "out2")
        assert out2.name == "orig_masked.docx"

    def test_no_write_modifies_paragraph_without_entities(self, tmp_path):
        """无命中段落原样保留（w:t 数量与文本不变）。"""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("普通")
        p.add_run("文本")
        src = tmp_path / "plain.docx"
        doc.save(src)

        adapter, _ = _make_adapter()
        out = adapter.process(src, tmp_path / "out")
        d2 = Document(str(out))
        nodes = DocxAdapter._paragraph_text_nodes(d2.paragraphs[0]._p)
        assert [t.text or "" for t in nodes] == ["普通", "文本"]
