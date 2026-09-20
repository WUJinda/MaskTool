# -*- coding: utf-8 -*-
"""tests/test_review_fixes.py — R1 对抗审查 A/B 级修复的回归测试

覆盖：
- A1 确认模式检测面=处理面（docx 页眉独有实体 / xlsx 货币数字合成项 / 富文本 /
  验证列表；engine 确认过滤警告兜底）
- A2 Web unmask 统一公共还原（页眉无 token 残留 + 数字单元格还原为数值）
- B1 目录模式自嵌套排除；B2 路径处置集对齐 + 嵌套 ner 过滤；B3 单文件名脱敏；
  B4 还原示例新根名；B5 Web 检测临时目录清理；B6 Web 配置回退链；
  B7 空目录迁移；B8 跨批次 mapping 错配警告；B9 skipped_unmasked 隔离
"""

import io
import json
import logging
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.datavalidation import DataValidation
from typer.testing import CliRunner

import mask_tool.cli as cli_module
from mask_tool.cli import app
from mask_tool.core.templates import (
    DEFAULT_CONFIG_YAML, SAMPLE_LEXICON_YAML, WHITELIST_YAML,
)
from mask_tool.models.detection import (
    DetectionResult, DetectionStatus, DetectionType, Location,
)

TOKEN_RE = re.compile(
    r"\[(?:COMPANY|GOVERNMENT|PERSON|PROJECT|SUBJECT|LOCATION|AMOUNT|CUSTOM)_\d{3,}\]"
)

LEXICON = {
    "company": ["某某科技有限公司"],
    "person": ["张三", "李四"],
}


@pytest.fixture()
def runner():
    return CliRunner(env={"COLUMNS": "400"})


@pytest.fixture(autouse=True)
def _cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _flat(output: str) -> str:
    return "".join(output.split())


def _write_config_dir(lexicon=None):
    cfgdir = Path("config")
    cfgdir.mkdir(parents=True, exist_ok=True)
    default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")
    (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
    (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
    (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
    (cfgdir / "lexicon.yaml").write_text(
        yaml.dump(lexicon or LEXICON, allow_unicode=True), encoding="utf-8",
    )


def _make_docx(path: Path, paragraphs, header_text=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if header_text is not None:
        doc.sections[0].header.paragraphs[0].text = header_text
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _docx_header_text(path: Path) -> str:
    doc = Document(str(path))
    return doc.sections[0].header.paragraphs[0].text


def _docx_paragraphs(path: Path):
    return [p.text for p in Document(str(path)).paragraphs]


def _batch_dirs(output: Path = Path("output")):
    if not output.exists():
        return []
    return sorted(p for p in output.iterdir() if p.is_dir())


def _load_mapping(batch: Path) -> dict:
    return json.loads((batch / "mapping.json").read_text(encoding="utf-8"))


class _AutoYes:
    """全量勾选的确认引擎替身。"""
    learned = []
    auto_yes = True

    def confirm_batch(self, results, file_name=""):
        return list(results)

    def get_learned_words(self):
        return {}


class _RejectAll:
    learned = []
    auto_yes = False

    def confirm_batch(self, results, file_name=""):
        return []

    def get_learned_words(self):
        return {}


# ===========================================================================
# A1：确认模式检测面 = 处理面
# ===========================================================================

class TestA1ConfirmDetectionSurface:

    def test_confirm_header_only_entity_masked(self, runner, monkeypatch):
        """页眉独有实体：确认表格可见（检测面），勾选后页眉被替换（处理面）。"""
        from mask_tool.core.confirm import ConfirmEngine

        monkeypatch.setattr(
            cli_module, "_new_confirm_engine",
            lambda: ConfirmEngine(auto_yes=True),
        )
        _write_config_dir()
        src = _make_docx(
            Path("h.docx"),
            ["正文无敏感内容"],
            header_text="某某科技有限公司 编号001",
        )
        result = runner.invoke(app, ["mask", str(src), "--output", "out", "--confirm"])
        assert result.exit_code == 0, result.output
        (out_docx,) = _batch_dirs(Path("out"))[0].glob("*.docx")
        header = _docx_header_text(out_docx)
        assert "某某科技有限公司" not in header
        assert "[COMPANY_001]" in header

    def test_confirm_xlsx_currency_synth_visible(self, runner, monkeypatch):
        """货币格式数字单元格（R1 合成项，仅存在于 adapter）：确认模式可见且勾选后替换。"""
        from mask_tool.core.confirm import ConfirmEngine

        monkeypatch.setattr(
            cli_module, "_new_confirm_engine",
            lambda: ConfirmEngine(auto_yes=True),
        )
        _write_config_dir()
        wb = Workbook()
        wb.active["A1"] = 12000000
        wb.active["A1"].number_format = "¥#,##0.00"
        wb.save("n.xlsx")
        result = runner.invoke(app, ["mask", "n.xlsx", "--output", "out", "--confirm"])
        assert result.exit_code == 0, result.output
        (out_xlsx,) = _batch_dirs(Path("out"))[0].glob("*.xlsx")
        assert load_workbook(str(out_xlsx)).active["A1"].value == "[AMOUNT_001]"

    def test_confirm_xlsx_rich_text_visible(self, runner, monkeypatch):
        """富文本单元格（CellRichText）：确认模式可见且块内替换。"""
        from mask_tool.core.confirm import ConfirmEngine

        monkeypatch.setattr(
            cli_module, "_new_confirm_engine",
            lambda: ConfirmEngine(auto_yes=True),
        )
        _write_config_dir()
        wb = Workbook()
        wb.active["A1"] = CellRichText(
            [TextBlock(InlineFont(), "联系人"), TextBlock(InlineFont(b=True), "张三")]
        )
        wb.save("r.xlsx")
        result = runner.invoke(app, ["mask", "r.xlsx", "--output", "out", "--confirm"])
        assert result.exit_code == 0, result.output
        (out_xlsx,) = _batch_dirs(Path("out"))[0].glob("*.xlsx")
        value = load_workbook(str(out_xlsx), rich_text=True).active["A1"].value
        # 块内替换：内联格式保留（仍是 CellRichText），拼接串验证内容
        assert str(value) == "联系人[PERSON_001]"

    def test_detect_file_results_covers_validation_list(self, tmp_path):
        """验证列表（数据验证 list 项）进入检测面（函数级）。"""
        from mask_tool.adapters.extract import detect_file_results
        from mask_tool.core.detector import Detector
        from mask_tool.core.policy import PolicyEngine
        from mask_tool.models.config import MaskConfig

        wb = Workbook()
        ws = wb.active
        dv = DataValidation(type="list", formula1='"张三,普通项"')
        dv.add("B1")  # sqref 必填，否则保存时丢弃
        ws.add_data_validation(dv)
        wb.save(str(tmp_path / "v.xlsx"))

        cfg = MaskConfig(mode="smart")
        cfg.ner.enabled = False
        detector = Detector(LEXICON, whitelist=set(), ner_engine=None)
        policy = PolicyEngine(cfg)
        results = detect_file_results(
            tmp_path / "v.xlsx", detector, policy,
        )
        assert "张三" in {r.text for r in results}

    def test_confirm_filtered_auto_warning(self, tmp_path, caplog):
        """兜底：确认模式下 AUTO 项不在勾选集 -> 警告列出（engine/pipeline 钩子）。"""
        from mask_tool.core.pipeline import Pipeline
        from mask_tool.models.config import MaskConfig

        lex = tmp_path / "lex.yaml"
        lex.write_text(yaml.dump(LEXICON, allow_unicode=True), encoding="utf-8")
        wl = tmp_path / "wl.yaml"
        wl.write_text(WHITELIST_YAML, encoding="utf-8")
        cfg = MaskConfig(mode="smart")
        cfg.ner.enabled = False
        pipeline = Pipeline(cfg, lexicon_path=str(lex), whitelist_path=str(wl))

        src = _make_docx(tmp_path / "w.docx", ["负责人张三", "抄送李四"])
        with caplog.at_level(logging.WARNING, logger="mask_tool"):
            pipeline.process_file(
                src, tmp_path / "out",
                allowed_originals={"张三"},  # 李四未勾选
            )
        assert any(
            "李四" in rec.message and "未做脱敏" in rec.message
            for rec in caplog.records
        ), [r.message for r in caplog.records]


# ===========================================================================
# A2：Web unmask 统一公共还原
# ===========================================================================

class TestA2UnifiedRestore:

    def _mask_docx_with_header(self, runner):
        _write_config_dir()
        src = _make_docx(
            Path("h.docx"), ["正文提到张三"],
            header_text="某某科技有限公司 页眉",
        )
        assert runner.invoke(
            app, ["mask", str(src), "--output", "out_docx"],
        ).exit_code == 0
        batch = _batch_dirs(Path("out_docx"))[0]
        return batch / "h_masked.docx", batch / "mapping.json"

    def _mask_xlsx_with_number(self, runner):
        _write_config_dir()
        wb = Workbook()
        wb.active["A1"] = 12000000
        wb.active["A1"].number_format = "¥#,##0.00"
        wb.save("n.xlsx")
        assert runner.invoke(
            app, ["mask", "n.xlsx", "--output", "out_num"],
        ).exit_code == 0
        batch = _batch_dirs(Path("out_num"))[0]
        return batch / "n_masked.xlsx", batch / "mapping.json"

    def test_restore_module_docx_header_no_token_left(self, runner, tmp_path):
        """公共还原模块：页眉 token 全部还原，无残留。"""
        from mask_tool.adapters.restore import restore_file_content

        masked, mapping_path = self._mask_docx_with_header(runner)
        assert "[COMPANY_001]" in _docx_header_text(masked)
        data = json.loads(mapping_path.read_text(encoding="utf-8"))
        out = tmp_path / "restored" / "h.docx"
        restore_file_content(masked, out, data["tokens"])
        header = _docx_header_text(out)
        assert "[COMPANY" not in header
        assert "某某科技有限公司" in header
        assert _docx_paragraphs(out) == ["正文提到张三"]

    def test_restore_module_xlsx_number_as_int(self, runner, tmp_path):
        """公共还原模块：kind=number 的数字单元格还原为数值。"""
        from mask_tool.adapters.restore import restore_file_content

        masked, mapping_path = self._mask_xlsx_with_number(runner)
        data = json.loads(mapping_path.read_text(encoding="utf-8"))
        out = tmp_path / "restored" / "n.xlsx"
        restore_file_content(masked, out, data["tokens"])
        value = load_workbook(str(out)).active["A1"].value
        assert value == 12000000
        assert isinstance(value, int)

    def test_web_unmask_file_restores_header_and_number(self, runner, tmp_path):
        """Web _unmask_file（str tokens 兼容形态）：页眉还原 + 数值还原为 int。"""
        mask_tool_web = pytest.importorskip("mask_tool.web.app")

        masked_docx, mapping_docx = self._mask_docx_with_header(runner)
        masked_xlsx, mapping_xlsx = self._mask_xlsx_with_number(runner)
        tokens_docx = {
            t: v["original"]
            for t, v in json.loads(mapping_docx.read_text(encoding="utf-8"))["tokens"].items()
        }
        tokens_xlsx = json.loads(
            mapping_xlsx.read_text(encoding="utf-8")
        )["tokens"]  # dict 条目形态（含 kind）

        out_dir = tmp_path / "web_restored"
        out_dir.mkdir()
        r1 = mask_tool_web._unmask_file(
            masked_docx, out_dir / "h.docx", tokens_docx,
        )
        assert r1 is not None
        header = _docx_header_text(r1)
        assert "[COMPANY" not in header and "某某科技有限公司" in header

        r2 = mask_tool_web._unmask_file(
            masked_xlsx, out_dir / "n.xlsx", tokens_xlsx,
        )
        assert r2 is not None
        value = load_workbook(str(r2)).active["A1"].value
        assert value == 12000000 and isinstance(value, int)

    def test_web_legacy_unmask_copies_removed(self):
        """旧版逐 run 替换实现已删除（消除第三份拷贝）。"""
        mask_tool_web = pytest.importorskip("mask_tool.web.app")
        assert not hasattr(mask_tool_web, "_unmask_docx")
        assert not hasattr(mask_tool_web, "_unmask_xlsx")


# ===========================================================================
# B2：路径处置集与内容一致 + 嵌套 ner 过滤
# ===========================================================================

class TestB2PathStatuses:

    def _real_masker(self):
        from mask_tool.core.detector import Detector
        from mask_tool.core.path_masker import PathMasker
        from mask_tool.core.policy import PolicyEngine
        from mask_tool.core.tokenizer import TokenGenerator
        from mask_tool.core.ner.jieba_ner import JiebaNER
        from mask_tool.models.config import MaskConfig

        cfg = MaskConfig(mode="smart")
        ner = JiebaNER()
        detector = Detector(
            {"company": ["某某科技有限公司"]}, whitelist=set(), ner_engine=ner,
        )
        return PathMasker(detector, PolicyEngine(cfg), TokenGenerator())

    def test_detect_name_filters_nested_ner(self):
        """嵌套于词典更长词条内的 ner 项被滤除（词典命中覆盖处不叠加 ner）。"""
        from mask_tool.core.path_masker import PathMasker
        from mask_tool.core.policy import PolicyEngine
        from mask_tool.core.tokenizer import TokenGenerator
        from mask_tool.models.config import MaskConfig

        class StubDetector:
            def detect(self, name, context=""):
                return [
                    DetectionResult(
                        text="张三科技有限公司",
                        text_type=DetectionType.COMPANY,
                        source="dictionary", confidence=0.95,
                        location=Location(file="x"),
                    ),
                    DetectionResult(
                        text="张三", text_type=DetectionType.PERSON,
                        source="ner", confidence=0.65,
                        location=Location(file="x"),
                    ),
                ]

        pm = PathMasker(
            StubDetector(), PolicyEngine(MaskConfig(mode="smart")), TokenGenerator(),
        )
        kept = pm.detect_name("张三科技有限公司")
        assert [r.text for r in kept] == ["张三科技有限公司"]

    def test_build_new_name_default_auto_only(self):
        """SUGGEST（ner 0.65）默认不改名（与内容侧 H6 对齐）。"""
        pm = self._real_masker()
        # 检测面上 SUGGEST 项存在（确认表格可见，可勾选）
        detections = pm.policy.apply(pm.detect_name("张三科技"))
        assert any(r.status == DetectionStatus.SUGGEST_MASK for r in detections)
        # 默认 statuses={AUTO}：不改名
        new_name, _ = pm.build_new_name("张三科技.xlsx", kind="file")
        assert new_name == "张三科技.xlsx"

    def test_build_new_name_all_statuses_replaces(self):
        """statuses 含 SUGGEST（--all / 确认语义）时改名。"""
        pm = self._real_masker()
        new_name, _ = pm.build_new_name(
            "张三科技.xlsx", kind="file",
            statuses=frozenset(
                {DetectionStatus.AUTO_MASK, DetectionStatus.SUGGEST_MASK}
            ),
        )
        assert new_name != "张三科技.xlsx"
        assert "[PERSON_" in new_name

    def test_cli_dir_mode_path_statuses_follow_all(self, runner):
        """端到端：目录模式默认下 SUGGEST-only 命中文件名不改名；--all 后改名。"""
        from mask_tool.core.ner.jieba_ner import JiebaNER
        pytest.importorskip("jieba")
        if not JiebaNER().is_available():
            pytest.skip("jieba 不可用")

        lexicon = {"company": ["某某科技有限公司"]}
        cfgdir = Path("config")
        cfgdir.mkdir(parents=True, exist_ok=True)
        default = DEFAULT_CONFIG_YAML.replace("enabled: true", "enabled: false")
        # NER 显式开启（用模板原文但 jieba 兜底）：
        (cfgdir / "default.yaml").write_text(default, encoding="utf-8")
        (cfgdir / "whitelist.yaml").write_text(WHITELIST_YAML, encoding="utf-8")
        (cfgdir / "sample_lexicon.yaml").write_text(SAMPLE_LEXICON_YAML, encoding="utf-8")
        # 词库不带 person：文件名"张三科技.docx"的"张三"仅来自 ner(SUGGEST)
        (cfgdir / "lexicon.yaml").write_text(
            yaml.dump(lexicon, allow_unicode=True), encoding="utf-8",
        )
        # NER 需要开启：重建 default 开启 ner
        (cfgdir / "default.yaml").write_text(DEFAULT_CONFIG_YAML, encoding="utf-8")

        src_dir = Path("tree")
        src_dir.mkdir()
        _make_docx(src_dir / "张三科技.docx", ["无敏感正文"])

        r1 = runner.invoke(app, ["mask", str(src_dir), "--output", "out1"])
        assert r1.exit_code == 0, r1.output
        tree1 = _batch_dirs(Path("out1"))[0] / "tree"
        # 默认：ner SUGGEST 不替换 -> 文件名保持
        assert (tree1 / "张三科技.docx").exists()
        assert not TOKEN_RE.search(tree1.joinpath("张三科技.docx").name)

        r2 = runner.invoke(app, ["mask", str(src_dir), "--output", "out2", "--all"])
        assert r2.exit_code == 0, r2.output
        tree2 = _batch_dirs(Path("out2"))[0] / "tree"
        masked = [p.name for p in tree2.iterdir() if p.suffix == ".docx"]
        assert any("[PERSON_" in n for n in masked), masked


# ===========================================================================
# B3：CLI 单文件名脱敏
# ===========================================================================

class TestB3SingleFileNameMasking:

    def test_single_file_stem_masked_by_default(self, runner):
        """默认：产物主名脱敏 + _masked 后缀，paths 段记录，unmask 可还原。"""
        _write_config_dir()
        src = _make_docx(Path("合同-张三.docx"), ["负责人张三"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        masked = batch / "合同-[PERSON_001]_masked.docx"
        assert masked.exists()
        assert _docx_paragraphs(masked) == ["负责人[PERSON_001]"]

        data = _load_mapping(batch)
        rels = {pm["rel_new"]: pm["rel_old"] for pm in data["paths"]}
        assert rels["合同-[PERSON_001]_masked.docx"] == "合同-张三_masked.docx"

        result2 = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result2.exit_code == 0, result2.output
        assert _docx_paragraphs(Path("restored") / "合同-张三_masked.docx") == ["负责人张三"]

    def test_single_file_no_mask_names_keeps_original_stem(self, runner):
        """--no-mask-names：保持旧版 {原名}_masked 语义。"""
        _write_config_dir()
        src = _make_docx(Path("合同-张三.docx"), ["负责人张三"])
        result = runner.invoke(
            app, ["mask", str(src), "--output", "out", "--no-mask-names"],
        )
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        assert (batch / "合同-张三_masked.docx").exists()
        assert _load_mapping(batch).get("paths", []) == []

    def test_single_file_clean_name_unchanged(self, runner):
        """无命中主名：行为与旧版一致（原名 + _masked，paths 段空）。"""
        _write_config_dir()
        src = _make_docx(Path("普通文件.docx"), ["负责人张三"])
        result = runner.invoke(app, ["mask", str(src), "--output", "out"])
        assert result.exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        assert (batch / "普通文件_masked.docx").exists()
        assert _load_mapping(batch).get("paths", []) == []


# ===========================================================================
# B4：还原命令示例用新根名
# ===========================================================================

class TestB4RestoreExample:

    def test_restore_example_points_to_renamed_root(self, runner):
        _write_config_dir()
        root = Path("项目资料-某某科技有限公司")
        root.mkdir()
        _make_docx(root / "a.docx", ["甲方某某科技有限公司"])
        result = runner.invoke(app, ["mask", str(root), "--output", "out"])
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]

        # 示例命令中的目标路径真实存在（新根名）
        m = re.search(r'unmask "([^"]+)"', result.output)
        assert m, result.output
        target = Path(m.group(1))
        assert target.exists(), f"示例路径不存在: {target}"
        assert target.name == "项目资料-[COMPANY_001]"


# ===========================================================================
# B5：Web 检测流程临时目录清理
# ===========================================================================

class TestB5WebDetectionTmpCleanup:

    def test_run_detection_cleans_previous_tmp(self, monkeypatch):
        web = pytest.importorskip("mask_tool.web.app")

        monkeypatch.setattr(web.st, "rerun", lambda: None, raising=False)
        for key in list(web.st.session_state.keys()):
            del web.st.session_state[key]

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("联系电话13812345678")
        doc.save(buf)
        upload = SimpleNamespace(name="a.docx", read=lambda: buf.getvalue())

        web._run_detection([upload], "smart", ner_enabled=False)
        tmp1 = web.st.session_state.get("tmp_dir")
        assert tmp1 and Path(tmp1).exists()

        web._run_detection([upload], "smart", ner_enabled=False)
        tmp2 = web.st.session_state.get("tmp_dir")
        assert tmp2 and Path(tmp2).exists() and tmp2 != tmp1
        assert not Path(tmp1).exists()  # 旧一轮的上传副本已清理


# ===========================================================================
# B6：Web 配置回退链对齐 CLI
# ===========================================================================

class TestB6WebConfigFallback:

    def test_bare_cwd_falls_back_with_warning(self, monkeypatch):
        """任意 CWD：内嵌模板回退 + 词库为空警告（不再静默）。

        部署锚点同步隔离到 CWD（与 test_cli 同名测试同理）：否则
        源码树根的 config/sample_lexicon.yaml 会被锚点链找到，
        "词库文件不存在"警告不可达。
        """
        web = pytest.importorskip("mask_tool.web.app")
        import mask_tool.core.config_loader as cl
        monkeypatch.setattr(cl, "runtime_anchor_dirs", lambda: [Path.cwd()])
        monkeypatch.setattr(cl, "writable_anchor_dir", lambda: Path.cwd())
        warns = []
        monkeypatch.setattr(web.st, "warning", lambda m, **kw: warns.append(str(m)))

        cfg = web._load_config("smart")
        assert cfg is not None
        flat = "".join(warns)
        assert "回退到内嵌默认配置模板" in flat
        assert "词库文件不存在" in flat

    def test_cwd_config_loaded_without_template_warning(self, monkeypatch):
        web = pytest.importorskip("mask_tool.web.app")
        _write_config_dir()
        warns = []
        monkeypatch.setattr(web.st, "warning", lambda m, **kw: warns.append(str(m)))
        cfg = web._load_config("smart")
        assert cfg is not None
        flat = "".join(warns)
        assert "回退到内嵌默认配置模板" not in flat


# ===========================================================================
# B7：空目录迁移（含目录名脱敏）
# ===========================================================================

class TestB7EmptyDirs:

    def test_empty_dir_migrated_and_name_masked(self, runner):
        _write_config_dir()
        src_dir = Path("tree")
        (src_dir / "某某科技有限公司档案").mkdir(parents=True)
        _make_docx(src_dir / "a.docx", ["负责人张三"])

        result = runner.invoke(app, ["mask", str(src_dir), "--output", "out"])
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / "tree"
        # 空目录迁移且目录名脱敏
        masked_dir = tree / "[COMPANY_001]档案"
        assert masked_dir.is_dir()
        assert list(masked_dir.iterdir()) == []

        # unmask 后空目录按原名恢复
        result2 = runner.invoke(app, [
            "unmask", str(tree), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result2.exit_code == 0, result2.output
        restored_dir = Path("restored") / "tree" / "某某科技有限公司档案"
        assert restored_dir.is_dir()
        assert list(restored_dir.iterdir()) == []


# ===========================================================================
# B8：跨批次 mapping 错配警告
# ===========================================================================

class TestB8MappingMismatch:

    def test_unmask_mismatched_inputs_warns(self, runner):
        _write_config_dir()
        _make_docx(Path("源文件甲.docx"), ["负责人张三"])
        assert runner.invoke(
            app, ["mask", "源文件甲.docx", "--output", "out"],
        ).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]

        # 完全不相干的文件 + 该批次 mapping
        _make_docx(Path("完全不同的文档.docx"), ["普通内容"])
        result = runner.invoke(app, [
            "unmask", "完全不同的文档.docx",
            "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        flat = _flat(result.output)
        assert "无任何交集" in flat
        assert "可能混用了其他批次的映射表" in flat

    def test_unmask_matching_inputs_no_warning(self, runner):
        _write_config_dir()
        src = _make_docx(Path("源文件甲.docx"), ["负责人张三"])
        assert runner.invoke(
            app, ["mask", str(src), "--output", "out"],
        ).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        masked = batch / "源文件甲_masked.docx"
        result = runner.invoke(app, [
            "unmask", str(masked), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        assert "无任何交集" not in _flat(result.output)


# ===========================================================================
# B9：目录 + --confirm 未勾选文件隔离
# ===========================================================================

class TestB9SkippedUnmasked:

    def test_unchecked_files_isolated(self, runner, monkeypatch):
        monkeypatch.setattr(cli_module, "_new_confirm_engine", _RejectAll)
        _write_config_dir()
        src_dir = Path("tree")
        src_dir.mkdir()
        _make_docx(src_dir / "合同-张三.docx", ["负责人张三"])

        result = runner.invoke(
            app, ["mask", str(src_dir), "--output", "out", "--confirm"],
        )
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / "tree"

        # 未勾选文件进隔离目录，保持原名原内容
        skipped = tree / "skipped_unmasked" / "合同-张三.docx"
        assert skipped.exists()
        assert _docx_paragraphs(skipped) == ["负责人张三"]
        # 产物树中不再出现"名已脱敏、内容未脱敏"的文件
        assert not (tree / "合同-[PERSON_001].docx").exists()
        # 摘要明确提示数量与目录
        flat = _flat(result.output)
        assert "skipped_unmasked" in flat
        assert "未勾选（内容未脱敏）文件:1个" in flat

    def test_skipped_dir_exempt_from_rename(self, runner, monkeypatch):
        monkeypatch.setattr(cli_module, "_new_confirm_engine", _RejectAll)
        _write_config_dir()
        src_dir = Path("tree")
        src_dir.mkdir()
        _make_docx(src_dir / "合同-张三.docx", ["负责人张三"])
        _make_docx(src_dir / "sub" / "内部-李四.docx", ["抄送李四"])

        result = runner.invoke(
            app, ["mask", str(src_dir), "--output", "out", "--confirm"],
        )
        assert result.exit_code == 0, result.output
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / "tree"
        # 隔离区内文件名不做脱敏改名（保持原名与相对结构）
        assert (tree / "skipped_unmasked" / "合同-张三.docx").exists()
        assert (tree / "skipped_unmasked" / "sub" / "内部-李四.docx").exists()
        # mapping 的 paths 段不含隔离区条目
        data = _load_mapping(batch)
        for pm in data.get("paths", []):
            assert "skipped_unmasked" not in pm["rel_old"]
            assert "skipped_unmasked" not in pm["rel_new"]


# ===========================================================================
# B1：输出目录自嵌套排除
# ===========================================================================

class TestB1SelfNesting:

    def test_output_inside_source_not_sucked_in(self, runner):
        _write_config_dir()
        src_dir = Path("docs")
        src_dir.mkdir()
        _make_docx(src_dir / "a.docx", ["负责人张三"])

        # 第一次：output 建在源目录内部
        assert runner.invoke(
            app, ["mask", str(src_dir), "--output", str(src_dir / "out")],
        ).exit_code == 0
        # 第二次：历史批次产物不被吸入
        result = runner.invoke(
            app, ["mask", str(src_dir), "--output", str(src_dir / "out")],
        )
        assert result.exit_code == 0, result.output
        flat = _flat(result.output)
        assert "找到1个文件待处理" in flat
        batches = _batch_dirs(src_dir / "out")
        assert len(batches) == 2
        # 第二批镜像树内无嵌套 out/ 历史产物
        tree = batches[1] / "docs"
        assert not (tree / "out").exists()
        assert list(tree.glob("*.docx")) == [tree / "a.docx"]

    def test_unmask_tree_walk_unaffected(self, runner):
        """unmask 自身的树内遍历（还原树）不受 exclude 参数影响。"""
        _write_config_dir()
        src_dir = Path("tree")
        src_dir.mkdir()
        _make_docx(src_dir / "a.docx", ["负责人张三"])
        assert runner.invoke(
            app, ["mask", str(src_dir), "--output", "out"],
        ).exit_code == 0
        batch = _batch_dirs(Path("out"))[0]
        tree = batch / "tree"
        result = runner.invoke(app, [
            "unmask", str(tree), "--mapping", str(batch / "mapping.json"),
            "--output", "restored",
        ])
        assert result.exit_code == 0
        assert _docx_paragraphs(Path("restored") / "tree" / "a.docx") == ["负责人张三"]


# ===========================================================================
# B10：README 已知限制
# ===========================================================================

class TestB10Readme:

    def test_readme_documents_docx_revision_limitation(self):
        readme = Path(__file__).parent.parent / "README.md"
        text = readme.read_text(encoding="utf-8")
        assert "修订记录" in text
        assert "接受所有修订" in text
        assert "skipped_unmasked" in text
