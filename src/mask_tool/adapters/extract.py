"""检测面提取公共模块（审查 R1-A1：检测面 = 处理面）

CLI ``--confirm`` / ``inspect`` 与 Web 检测流程共用的提取入口：

- docx：委托 ``DocxAdapter.extract_paragraph_texts``（XML walker 全部件：
  正文/表格/页眉脚/脚注/尾注/批注/文本框/超链接），与 mask 处理面同源。
- xlsx：遍历面与 ``XlsxAdapter._process_workbook`` 一致（单元格纯文本/富文本/
  数字 canonical/公式字面量、批注、超链接 tooltip、页眉脚、数据验证 list 项、
  定义名称值/备注）；数字单元格的合成检测规则直接复用
  ``XlsxAdapter._synth_numeric_results`` / ``_numeric_canonical`` /
  ``_merge_numeric_results``——禁止在本模块复制第三份规则。

确认模式的承诺是"全量展示、勾选放行"：确认表格的检测面必须覆盖 adapter 将要
处理的全部检测项，否则 ``allowed_originals`` 会把盲区实体静默漏脱（A1）。
"""

from pathlib import Path
from typing import List

from mask_tool.models.detection import DetectionResult, Location


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------

def extract_texts(file_path) -> str:
    """提取全部件纯文本（换行拼接；通用检测 / unmask 对账扫描用）。

    与 ``mask_tool.cli._extract_text`` 语义一致并扩展了 xlsx 盲区部件；
    pptx/pdf 已屏蔽，遇到时抛 ValueError（与 CLI 现行为一致）。
    """
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        from mask_tool.adapters.docx_adapter import DocxAdapter
        return "\n".join(DocxAdapter.extract_paragraph_texts(file_path))
    if suffix == ".xlsx":
        return "\n".join(extract_xlsx_texts(file_path))
    if suffix in (".pptx", ".pdf"):
        raise ValueError(
            f"{suffix} 脱敏暂未开放，无法检测（见 core/formats.py BLOCKED_EXTS）"
        )
    return ""


def detect_file_results(file_path, detector, policy) -> List[DetectionResult]:
    """与 adapter 处理面同源的检测结果列表（confirm 表格 / Web 检测用）。

    文本部件逐段 ``detector.detect``；xlsx 数字单元格走 synth 合成规则 +
    generic detect + merge（与 ``XlsxAdapter._process_numeric_cell`` 相同口径）；
    全部结果统一过 ``policy.apply``。
    """
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    results: List[DetectionResult] = []
    if suffix == ".docx":
        from mask_tool.adapters.docx_adapter import DocxAdapter
        file_path_str = str(file_path)
        for para in DocxAdapter.extract_paragraph_texts(file_path):
            results.extend(detector.detect(para, file_path_str))
    elif suffix == ".xlsx":
        results.extend(_detect_xlsx_results(file_path, detector))
    return policy.apply(results)


# ---------------------------------------------------------------------------
# xlsx 提取（遍历面 = XlsxAdapter 处理面）
# ---------------------------------------------------------------------------

def extract_xlsx_texts(file_path) -> List[str]:
    """xlsx 全部件文本列表（提取面与 XlsxAdapter 处理面一致）。

    覆盖：单元格纯文本 / 富文本块与整串 / 数字 canonical 渲染 / 公式双引号
    字面量 / 批注 / 超链接 tooltip / 页眉脚 6 部位 x L/C/R / 数据验证 list 项 /
    定义名称值/备注字面量。工作表名 adapter 仅提示不改名，不产生替换盲区，
    不纳入（与处理面语义一致）。
    """
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.cell.rich_text import CellRichText

    from mask_tool.adapters.xlsx_adapter import XlsxAdapter

    wb = load_workbook(str(file_path), rich_text=True)
    texts: List[str] = []
    try:
        for ws in wb.worksheets:
            _collect_header_footer(ws, texts)
            _collect_data_validations(ws, texts)
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    _collect_comment(cell, texts)
                    _collect_hyperlink_tooltip(cell, texts)
                    value = cell.value
                    if value is None or isinstance(value, bool) or cell.is_date:
                        continue
                    if isinstance(value, CellRichText):
                        # 富文本：逐块 + 整串（与 _process_rich_cell 两遍口径一致）
                        for block in value:
                            block_text = (
                                block if isinstance(block, str)
                                else getattr(block, "text", None)
                            )
                            if isinstance(block_text, str) and block_text.strip():
                                texts.append(block_text)
                        whole = str(value)
                        if whole.strip():
                            texts.append(whole)
                    elif isinstance(value, (int, float)):
                        canon = XlsxAdapter._numeric_canonical(value)
                        if canon:
                            texts.append(canon)
                    elif isinstance(value, str) or cell.data_type == "f":
                        # str：纯文本/普通公式串；data_type=='f' 的非 str 值 =
                        # ArrayFormula（R3-A2），token/实体只会在 .text 的
                        # 双引号字面量内——与写侧 _process_formula_cell 对称
                        text = (
                            value if isinstance(value, str)
                            else getattr(value, "text", None)
                        )
                        if isinstance(text, str) and text.strip():
                            if cell.data_type == "f":
                                _collect_formula_literals(text, texts)
                            else:
                                texts.append(text)
            _collect_defined_names(
                getattr(ws, "defined_names", {}), texts,
            )
        _collect_defined_names(getattr(wb, "defined_names", {}), texts)
    finally:
        wb.close()
    return texts


def _detect_xlsx_results(file_path: Path, detector) -> List[DetectionResult]:
    """xlsx 检测结果（数字单元格 synth 规则与 adapter 写路径共用同一实现）。"""
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.cell.rich_text import CellRichText

    from mask_tool.adapters.xlsx_adapter import XlsxAdapter

    file_path_str = str(file_path)
    results: List[DetectionResult] = []
    wb = load_workbook(str(file_path), rich_text=True)
    try:
        for ws in wb.worksheets:
            # 文本部件：与 extract_xlsx_texts 同一收集器，逐段检测
            part_texts: List[str] = []
            _collect_header_footer(ws, part_texts)
            _collect_data_validations(ws, part_texts)
            for text in part_texts:
                results.extend(detector.detect(text, file_path_str))
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    loc = Location(
                        file=file_path_str, sheet=ws.title, cell_ref=cell.coordinate,
                    )
                    comment = getattr(cell, "comment", None)
                    if comment is not None:
                        text = getattr(comment, "text", None)
                        if isinstance(text, str) and text.strip():
                            results.extend(detector.detect(text, file_path_str))
                    tooltip = getattr(getattr(cell, "hyperlink", None), "tooltip", None)
                    if isinstance(tooltip, str) and tooltip.strip():
                        results.extend(detector.detect(tooltip, file_path_str))
                    value = cell.value
                    if value is None or isinstance(value, bool) or cell.is_date:
                        continue
                    if isinstance(value, CellRichText):
                        block_texts = []
                        for block in value:
                            block_text = (
                                block if isinstance(block, str)
                                else getattr(block, "text", None)
                            )
                            if isinstance(block_text, str) and block_text.strip():
                                block_texts.append(block_text)
                        whole = str(value)
                        if whole.strip():
                            block_texts.append(whole)
                        for text in block_texts:
                            results.extend(detector.detect(text, file_path_str))
                    elif isinstance(value, (int, float)):
                        # 数字单元格：synth 合成规则（与写路径同一实现）+ 通用检测
                        canon = XlsxAdapter._numeric_canonical(value)
                        if canon is None:
                            continue
                        synth = XlsxAdapter._synth_numeric_results(
                            canon, cell.number_format or "", loc,
                        )
                        generic = detector.detect(canon, file_path_str)
                        results.extend(
                            XlsxAdapter._merge_numeric_results(synth, generic, loc)
                        )
                    elif isinstance(value, str) or cell.data_type == "f":
                        # ArrayFormula（R3-A2）：data_type=='f' 的非 str 值，
                        # 实体只会在 .text 的双引号字面量内
                        text = (
                            value if isinstance(value, str)
                            else getattr(value, "text", None)
                        )
                        if not isinstance(text, str) or not text.strip():
                            continue
                        if cell.data_type == "f":
                            literals: List[str] = []
                            _collect_formula_literals(text, literals)
                            for lit in literals:
                                results.extend(detector.detect(lit, file_path_str))
                        else:
                            results.extend(detector.detect(text, file_path_str))
            # 定义名称（工作表级）
            results.extend(
                _detect_defined_names(
                    getattr(ws, "defined_names", {}), detector, file_path_str,
                )
            )
        # 定义名称（工作簿级）
        results.extend(
            _detect_defined_names(
                getattr(wb, "defined_names", {}), detector, file_path_str,
            )
        )
    finally:
        wb.close()
    return results


# ---------------------------------------------------------------------------
# 部件收集器（提取与检测两条路径共用）
# ---------------------------------------------------------------------------

# 页眉页脚部件清单复用 xlsx_adapter 私有常量（R2-C：消除双定义漂移风险）


def _collect_header_footer(ws, texts: List[str]) -> None:
    from mask_tool.adapters.xlsx_adapter import _HF_PARTS

    for part_name in _HF_PARTS:
        item = getattr(ws, part_name, None)
        if item is None:
            continue
        for position in ("left", "center", "right"):
            text = getattr(getattr(item, position, None), "text", None)
            if isinstance(text, str) and text.strip():
                texts.append(text)


def _collect_data_validations(ws, texts: List[str]) -> None:
    dv_list = getattr(getattr(ws, "data_validations", None),
                      "dataValidation", None) or []
    for dv in dv_list:
        if getattr(dv, "type", None) != "list":
            continue
        formula1 = getattr(dv, "formula1", None)
        if (not isinstance(formula1, str) or len(formula1) < 2
                or not formula1.startswith('"')):
            continue
        inner = formula1[1:-1] if formula1.endswith('"') else formula1[1:]
        for item in inner.split(","):
            if item.strip():
                texts.append(item)


def _collect_comment(cell, texts: List[str]) -> None:
    comment = getattr(cell, "comment", None)
    if comment is None:
        return
    text = getattr(comment, "text", None)
    if isinstance(text, str) and text.strip():
        texts.append(text)


def _collect_hyperlink_tooltip(cell, texts: List[str]) -> None:
    tooltip = getattr(getattr(cell, "hyperlink", None), "tooltip", None)
    if isinstance(tooltip, str) and tooltip.strip():
        texts.append(tooltip)


def _collect_formula_literals(formula: str, texts: List[str]) -> None:
    from mask_tool.adapters.xlsx_adapter import _FORMULA_LITERAL

    for m in _FORMULA_LITERAL.finditer(formula):
        literal = m.group(1).replace('""', '"')
        if literal.strip():
            texts.append(literal)


def _iter_defined_name_literals(defined_names):
    """产出定义名称的值/备注中的双引号字面量（与 _process_defined_names 同拆法）。"""
    from mask_tool.adapters.xlsx_adapter import _FORMULA_LITERAL

    try:
        items = list(defined_names.items())
    except Exception:
        return
    for _name, dn in items:
        for attr in ("value", "comment"):
            text = getattr(dn, attr, None)
            if isinstance(text, str) and '"' in text:
                for m in _FORMULA_LITERAL.finditer(text):
                    literal = m.group(1).replace('""', '"')
                    if literal.strip():
                        yield literal


def _collect_defined_names(defined_names, texts: List[str]) -> None:
    for literal in _iter_defined_name_literals(defined_names):
        texts.append(literal)


def _detect_defined_names(defined_names, detector, file_path_str: str) -> List[DetectionResult]:
    results: List[DetectionResult] = []
    for literal in _iter_defined_name_literals(defined_names):
        results.extend(detector.detect(literal, file_path_str))
    return results
