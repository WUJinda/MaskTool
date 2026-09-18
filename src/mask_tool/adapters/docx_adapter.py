"""Word 文档(.docx)脱敏适配器（重写版，设计依据 D1 §3）

XML 级 walker 覆盖面（取代 doc.paragraphs/doc.tables 的 API 级遍历）：
- document.xml：正文 + 表格 + 嵌套表格（body.iter(w:p) 全覆盖）
- 各 section 页眉/页脚部件（含 first_page/even_page 变体，按 partname 全量枚举）
- 脚注(footnotes.xml) / 尾注(endnotes.xml) / 批注(comments.xml) 部件
- 文本框（w:txbxContent 内 w:p）与超链接内文本

检测单元 = w:p 整段拼接文本（跨 run 实体天然完整匹配，P1 决策）；替换经
统一引擎区间重建后整段并入首 w:t（保留其 rPr），其余 w:t 清空但元素保留——
从数据结构上不存在旧版 pos_map 切片错位与末 run 整体置空问题（H1 根治）。

已知限制（README 同步声明）：
- w:tab / w:br 是元素而非文本，不参与拼接串 → 跨 tab/软换行的实体
  字符可检测替换与还原，但重建后 tab/软换行可能偏离原位置（位置漂移）；
- 替换发生过的段落内混合格式退化为首 run 格式（P1 拍板接受）。
"""

import re
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List, Optional, Set, Tuple

from docx import Document
from docx.opc.oxml import serialize_part_xml
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from mask_tool.adapters.base import FileAdapter
from mask_tool.core.engine import ReplacementEngine
from mask_tool.models.detection import Location

# ---- 部件枚举（按 partname；与 python-docx 是否注册部件类无关） ----

_DOCUMENT_PARTNAME = "/word/document.xml"
_STORY_PARTNAME_RE = re.compile(r"^/word/(?:header|footer)\d*\.xml$")
_COMMENTS_PARTNAME = "/word/comments.xml"
# python-docx 未注册部件类的 XML 部件：以 blob 形式挂载，需 解析-修改-回写
_BLOB_PARTNAMES = frozenset({"/word/footnotes.xml", "/word/endnotes.xml"})

# 部件中文标签（document 用空串：正文无需位置前缀）
_PART_LABEL_PREFIXES = (
    ("header", "页眉"),
    ("footer", "页脚"),
    ("footnotes", "脚注"),
    ("endnotes", "尾注"),
    ("comments", "批注"),
)


def _part_label(partname: str) -> str:
    """partname -> 中文标签（header2.xml -> 页眉2；document.xml -> ""）。"""
    base = partname.rsplit("/", 1)[-1]
    base = base.rsplit(".", 1)[0]
    for key, label in _PART_LABEL_PREFIXES:
        if base.startswith(key):
            suffix = base[len(key):]
            return f"{label}{suffix}" if suffix else label
    return ""


def clear_core_properties(doc) -> None:
    """清空文档身份元数据（D1 §3.3）。

    created / modified 时间戳非身份信息，保留；个别属性在特殊文档上只读时跳过。
    """
    cp = doc.core_properties
    for attr in (
        "author", "last_modified_by", "comments", "subject", "keywords",
        "category", "content_status", "identifier", "language", "version",
        "title",
    ):
        try:
            setattr(cp, attr, "")
        except Exception:
            pass
    try:
        cp.revision = 0
    except Exception:
        pass


class DocxAdapter(FileAdapter):
    """Word 文档脱敏适配器（XML 级 walker）"""

    def supported_extensions(self) -> list[str]:
        return [".docx"]

    # ------------------------------------------------------------------
    # 对外入口
    # ------------------------------------------------------------------

    def process(
        self,
        input_path: Path,
        output_dir: Path,
        output_name: Optional[str] = None,
    ) -> Path:
        """处理 Word 文档：清空元数据 -> walker 逐段脱敏 -> 另存。

        Args:
            input_path: 输入文件路径
            output_dir: 输出目录
            output_name: 输出文件名；None 时为 {stem}_masked.docx
                （目录模式由 CLI 传原文件名以构建镜像树）
        """
        input_path = Path(input_path)
        output_dir = Path(output_dir)

        doc = Document(str(input_path))
        clear_core_properties(doc)

        file_path = str(input_path)
        for label, root, flush in self._iter_xml_roots(doc):
            self._clear_comment_authors(root)
            for idx, p_elem in enumerate(root.iter(qn("w:p"))):
                self._process_xml_paragraph(
                    p_elem, file_path, part_label=label, para_index=idx,
                )
            if flush is not None:
                flush()

        output_dir.mkdir(parents=True, exist_ok=True)
        name = output_name or f"{input_path.stem}_masked.docx"
        output_path = output_dir / name
        doc.save(str(output_path))
        return output_path

    # ------------------------------------------------------------------
    # unmask（D1 §3.2：复用同一 walker，roundtrip 对称）
    # ------------------------------------------------------------------

    @classmethod
    def restore_document(
        cls,
        input_path: Path,
        output_path: Path,
        token_map: Dict[str, Any],
    ) -> Set[str]:
        """按 mapping 还原 docx：walker 取段落全文 -> engine.unmask_text
        整段替换 -> 整段并入首 w:t。页眉/脚注/批注等覆盖面与 mask 一致。

        Args:
            input_path: 脱敏后的文件
            output_path: 还原输出路径（输入文件不被修改）
            token_map: token -> 映射条目（mapping.json 的 dict 或 TokenMapping
                对象，取 original 字段）

        Returns:
            命中（实际发生还原）的 token 集合
        """
        originals: Dict[str, str] = {}
        for token, entry in token_map.items():
            if isinstance(entry, dict):
                original = entry.get("original", "")
            else:
                original = getattr(entry, "original", "")
            if token and original:
                originals[str(token)] = str(original)

        doc = Document(str(input_path))
        hits: Set[str] = set()
        for _label, root, flush in cls._iter_xml_roots(doc):
            for p_elem in root.iter(qn("w:p")):
                nodes = cls._paragraph_text_nodes(p_elem)
                if not nodes:
                    continue
                full_text = "".join(t.text or "" for t in nodes)
                if "[" not in full_text:
                    continue
                new_text, hit = ReplacementEngine.unmask_text(full_text, originals)
                if not hit:
                    continue
                hits |= hit
                cls._rewrite_text_nodes(nodes, new_text)
            if flush is not None:
                flush()

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return hits

    # ------------------------------------------------------------------
    # 提取（inspect / --confirm 全文检测用，与 mask 同源，消除 N1 盲区）
    # ------------------------------------------------------------------

    @classmethod
    def extract_paragraph_texts(cls, input_path: Path) -> List[str]:
        """walker 同源提取全部非空段落文本（含表格/页眉脚/脚注/尾注/批注/
        文本框/超链接）。"""
        doc = Document(str(input_path))
        texts: List[str] = []
        for _label, root, _flush in cls._iter_xml_roots(doc):
            for p_elem in root.iter(qn("w:p")):
                nodes = cls._paragraph_text_nodes(p_elem)
                text = "".join(t.text or "" for t in nodes)
                if text.strip():
                    texts.append(text)
        return texts

    # ------------------------------------------------------------------
    # XML walker
    # ------------------------------------------------------------------

    @staticmethod
    def _iter_xml_roots(
        doc,
    ) -> Iterator[Tuple[str, Any, Optional[Callable[[], None]]]]:
        """产出 (部件标签, XML 根元素, 回写钩子|None)。

        - document/header/footer/comments：python-docx 已注册部件类，
          element 直接可变（引用落盘，无需钩子）；
        - footnotes/endnotes：通用 Part（blob），解析副本修改后经钩子回写
          ``part._blob``，保存时按 blob 序列化。

        标签规则：document -> ""；header2.xml -> "页眉2"；footnotes -> "脚注"。
        """
        package = doc.part.package
        for part in package.iter_parts():
            partname = str(part.partname)
            flush: Optional[Callable[[], None]] = None
            if partname == _DOCUMENT_PARTNAME:
                root = part.element
            elif _STORY_PARTNAME_RE.match(partname) or partname == _COMMENTS_PARTNAME:
                root = getattr(part, "element", None)
                if root is None:
                    continue
            elif partname in _BLOB_PARTNAMES:
                blob = getattr(part, "blob", b"") or b""
                if not blob:
                    continue
                root = parse_xml(blob)

                def flush(part=part, root=root):  # noqa: F811
                    part._blob = serialize_part_xml(root)
            else:
                continue
            yield _part_label(partname), root, flush

    @staticmethod
    def _paragraph_text_nodes(p_elem) -> List[Any]:
        """段落的 w:t 节点集合 = { t | t 的最近祖先 w:p 是本段 }。

        该归属规则使文本框内嵌套 w:p、超链接内 w:t 归入正确的段落单元，
        同时杜绝外层段落重复吞并 txbxContent 内层段落的文本。
        """
        p_tag = qn("w:p")
        nodes: List[Any] = []
        for t in p_elem.iter(qn("w:t")):
            anc = t.getparent()
            while anc is not None and anc.tag != p_tag:
                anc = anc.getparent()
            if anc is p_elem:
                nodes.append(t)
        return nodes

    @staticmethod
    def _clear_comment_authors(root) -> int:
        """清空批注作者元数据（w:author / w:initials）。返回清空条数。"""
        count = 0
        for c in root.iter(qn("w:comment")):
            c.set(qn("w:author"), "")
            if c.get(qn("w:initials")):
                c.set(qn("w:initials"), "")
            count += 1
        return count

    # ------------------------------------------------------------------
    # 段落处理
    # ------------------------------------------------------------------

    def _process_xml_paragraph(
        self,
        p_elem,
        file_path: str,
        *,
        part_label: str = "",
        para_index: Optional[int] = None,
    ) -> None:
        """检测整段拼接文本 -> 引擎替换 -> 整段并入首 w:t。

        引擎调用经 base._mask_text_block：不显式传 statuses/allowed_originals，
        Pipeline.process_file 设置的文件作用域参数（--all / --confirm）生效。
        """
        nodes = self._paragraph_text_nodes(p_elem)
        if not nodes:
            return
        full_text = "".join(t.text or "" for t in nodes)
        if not full_text.strip():
            return
        location = Location(file=file_path, paragraph=para_index)
        outcome = self._mask_text_block(
            full_text,
            file_path,
            context_tag=part_label,
            location=location,
        )
        if outcome is None:
            return
        self._rewrite_text_nodes(nodes, outcome.text)

    @staticmethod
    def _rewrite_text_nodes(nodes: List[Any], new_text: str) -> None:
        """P1：整段并入首 w:t（保其 rPr），其余置空但元素保留。

        首尾含空白的新文本声明 xml:space="preserve"，防止序列化裁剪。
        """
        new_texts = ReplacementEngine.plan_run_rewrite(
            [t.text or "" for t in nodes], new_text,
        )
        for node, txt in zip(nodes, new_texts):
            node.text = txt
            if txt and txt != txt.strip():
                node.set(qn("xml:space"), "preserve")
