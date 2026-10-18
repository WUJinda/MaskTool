"""文件内容还原公共模块（审查 R1-A2：CLI 与 Web 共用同一实现）

从 CLI 抽取的还原路径，消除 Web/CLI/adapter 三份拷贝：

- docx：``DocxAdapter.restore_document``（XML walker，覆盖面与 mask 一致：
  页眉/页脚/脚注/尾注/批注/文本框全部还原，不再残留 token）
- xlsx：单元格（纯文本/富文本块/公式串/数组公式 ArrayFormula）/批注/页眉脚
  /数据验证 list/定义名称值+备注/超链接 tooltip 走
  ``XlsxAdapter.restore_value``；``load_workbook(rich_text=True)`` 保内联格式；
  kind=number 的整值 token 还原为数值（int/float）而非文本

R3-A1/A2：还原遍历面与 ``XlsxAdapter._process_workbook`` 写入面对称——
写入侧会替换 token 的每个部件（DV/定义名称/tooltip/ArrayFormula），
还原侧逐一覆盖，保证 token 零残留。

token_map 契约：``{token: 条目}``，条目为 mapping.json 的 dict（含 original /
可选 kind）或 TokenMapping 对象——与 ``XlsxAdapter.restore_value`` 一致。
"""

from pathlib import Path
from typing import Dict


def bracket_gate(token_map: Dict[str, dict]) -> bool:
    """是否允许各还原路径启用 ``"[" in text`` 快速短路优化。

    仅当 **全部** 键均为含 ``[`` 的 token 样式串时返回 True；只要存在
    任意伪装词键（alias/混合批次），含伪装词但不含 ``[`` 的文本段落
    就可能被短路误跳导致静默漏还原，必须禁用短路（宁慢勿漏）。
    """
    return bool(token_map) and all("[" in str(k) for k in token_map)


def restore_file_content(f: Path, output_path: Path, token_map: Dict[str, dict]) -> None:
    """单文件内容还原：docx 走 walker，xlsx 走 restore_value（数值还原）。"""
    f = Path(f)
    suffix = f.suffix.lower()
    if suffix == ".docx":
        from mask_tool.adapters.docx_adapter import DocxAdapter

        DocxAdapter.restore_document(f, output_path, token_map)
    elif suffix in (".xlsx", ".xlsm"):
        restore_xlsx_file(f, output_path, token_map)
    elif suffix == ".pdf":
        restore_pdf_file(f, output_path, token_map)
    else:
        raise ValueError(f"不支持的还原格式: {suffix}")


def restore_pdf_file(
    input_path: Path, output_path: Path, token_map: Dict[str, dict],
) -> None:
    """PDF 内容还原：逐页搜索 token → redaction 擦除 → 原位插回原文。

    原文可能含中文，用 PyMuPDF 内置 CJK 字体（china-s）插回；插入区宽度
    只有原 token 的宽度，长原文缩小字号至下限后允许溢出（还原排版是
    近似的，可读性优先）。"""
    import fitz

    from mask_tool.adapters.pdf_adapter import (
        _apply_edits, _block_common_size, _collect_edit, iter_pdf_text_blocks,
    )

    doc = fitz.open(str(input_path))
    try:
        if doc.is_encrypted:
            raise ValueError("PDF 已加密，无法还原（请先解除密码保护）")
        for page in doc:
            edits = []
            originals = []
            for full, chars in iter_pdf_text_blocks(page):
                hits = [
                    (token, entry)
                    for token, entry in token_map.items()
                    if token and token in full
                ]
                if not hits:
                    continue
                # 插入字号用块内众数（正常文本字号），token 缩小字号不传染
                common_size = _block_common_size(chars)
                for token, entry in hits:
                    original = (
                        entry.get("original", "")
                        if isinstance(entry, dict) else str(entry)
                    )
                    pos = full.find(token)
                    while pos != -1:
                        # pad=0：token 字符是本工具插入的、bbox 精确覆盖实际
                        # 渲染字形；不外扩，溢出部分对右邻的相交压到最小
                        # （对抗审查 P1-1：≥1pt 相交即整字连带删除）
                        edit = _collect_edit(chars, pos, pos + len(token),
                                             pad=0.0)
                        if edit is not None:
                            if edit["size"] < common_size:
                                edit["size"] = common_size
                            edits.append(edit)
                            originals.append(original)
                        pos = full.find(token, pos + 1)
            if edits:
                _apply_edits(page, edits, originals, fontname="china-s")
        doc.save(str(output_path), garbage=3, deflate=True)
    finally:
        doc.close()


def restore_xlsx_file(
    input_path: Path, output_path: Path, token_map: Dict[str, dict],
) -> None:
    """xlsx/xlsm 内容还原（D2 §6.1）：单元格/富文本块/公式串/ArrayFormula/批注/
    页眉脚/数据验证 list/定义名称值+备注/超链接 tooltip 走
    XlsxAdapter.restore_value，kind=number 的整值 token 还原为数值。
    xlsm 的 VBA 宏部件经 keep_vba 原样保留（与脱敏侧对称）。"""
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.cell.rich_text import CellRichText
    from mask_tool.adapters.xlsx_adapter import XlsxAdapter

    keep_vba = input_path.suffix.lower() == ".xlsm"
    wb = load_workbook(str(input_path), rich_text=True, keep_vba=keep_vba)
    # alias 伪装词批次禁用 '[' 短路（见 bracket_gate）：伪装词不含 '['
    gate = bracket_gate(token_map)

    def restore_text(text: str):
        value, hits = XlsxAdapter.restore_value(text, token_map)
        return value if hits else None

    # 全局定义名称（与写入侧 _process_workbook 首步对称，R3-A1）
    _restore_defined_names(getattr(wb, "defined_names", {}), restore_text, gate)

    for ws in wb.worksheets:
        _restore_data_validations(ws, restore_text, gate)   # 写入侧 §4.4
        _restore_defined_names(getattr(ws, "defined_names", {}),
                               restore_text, gate)          # 写入侧 §4.5（sheet 级）
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                comment = getattr(cell, "comment", None)
                if comment is not None:
                    text = getattr(comment, "text", None)
                    if isinstance(text, str) and (not gate or "[" in text):
                        value = restore_text(text)
                        if value is not None:
                            comment.text = value
                # 超链接 tooltip（写入侧 §4.7 只写 token 进 tooltip）
                tooltip = getattr(getattr(cell, "hyperlink", None), "tooltip", None)
                if isinstance(tooltip, str) and (not gate or "[" in tooltip):
                    new_tooltip = restore_text(tooltip)
                    if new_tooltip is not None:
                        cell.hyperlink.tooltip = str(new_tooltip)
                value = cell.value
                if value is None:
                    continue
                if not isinstance(value, (str, CellRichText)):
                    # 数组公式（ArrayFormula，R3-A2）：token 只写入 .text 的
                    # 双引号字面量内，整串替换等价；.text 为可写属性
                    formula = getattr(value, "text", None)
                    if isinstance(formula, str) and (not gate or "[" in formula):
                        new_formula = restore_text(formula)
                        if new_formula is not None:
                            value.text = str(new_formula)
                    continue
                text = str(value)
                if gate and "[" not in text:
                    continue
                if isinstance(value, CellRichText):
                    for i, block in enumerate(value):
                        block_text = block if isinstance(block, str) else getattr(block, "text", None)
                        if isinstance(block_text, str) and (not gate or "[" in block_text):
                            new_block, hits = XlsxAdapter.restore_value(block_text, token_map)
                            if hits:
                                if isinstance(block, str):
                                    value[i] = new_block
                                else:
                                    block.text = new_block
                else:
                    new_value = restore_text(text)
                    if new_value is not None:
                        cell.value = new_value
        for part_name in ("oddHeader", "evenHeader", "firstHeader",
                          "oddFooter", "evenFooter", "firstFooter"):
            item = getattr(ws, part_name, None)
            if item is None:
                continue
            for position in ("left", "center", "right"):
                part = getattr(item, position, None)
                text = getattr(part, "text", None)
                if isinstance(text, str) and (not gate or "[" in text):
                    value = restore_text(text)
                    if value is not None:
                        part.text = value

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def _restore_data_validations(ws, restore_text, gate: bool = True) -> None:
    """数据验证 list 还原（R3-A1）：与写入侧 ``_process_data_validations``
    对称——token 写在引号包裹的逗号串内（``"item1,item2"`` 格式），
    剥引号整串替换后按原格式回写。"""
    dv_list = getattr(getattr(ws, "data_validations", None),
                      "dataValidation", None) or []
    for dv in dv_list:
        if getattr(dv, "type", None) != "list":
            continue
        formula1 = getattr(dv, "formula1", None)
        if (not isinstance(formula1, str) or len(formula1) < 2
                or not formula1.startswith('"')):
            continue
        inner = formula1[1:-1] if formula1.endswith('"') else formula1[1:]
        if not inner or (gate and "[" not in inner):
            continue
        new_inner = restore_text(inner)
        if new_inner is not None:
            dv.formula1 = '"' + str(new_inner) + '"'


def _restore_defined_names(defined_names, restore_text, gate: bool = True) -> None:
    """定义名称还原（R3-A1）：值/备注整串 token 替换，与写入侧
    ``_process_defined_names`` 对称（token 只写入双引号字面量内，
    整串替换与逐字面量替换等价）。"""
    try:
        items = list(defined_names.items())
    except Exception:
        return
    for _name, dn in items:
        for attr in ("value", "comment"):
            text = getattr(dn, attr, None)
            if isinstance(text, str) and (not gate or "[" in text):
                new_text = restore_text(text)
                if new_text is not None:
                    setattr(dn, attr, str(new_text))
