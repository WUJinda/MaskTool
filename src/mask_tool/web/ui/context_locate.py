# -*- coding: utf-8 -*-
"""上下文定位弹窗：检测结果表「上下文」列点击 → 查看敏感词原文位置。

交互链路（R10）：
  AgGrid「上下文」列 cellRenderer（可点击 span）→ 点击时对隐藏 ``loc`` 列
  ``setDataValue(递增序号)`` → 复用既有 VALUE_CHANGED 全表回传 →
  ``_consume_grid_locate_click`` 消费点击（序号防重复消费：组件 return_value
  在下一次事件前保持不变，靠递增序号区分新旧点击）→
  ``session_state["ctx_locate_idx"]`` → ``@st.dialog`` 定位弹窗
  （``on_dismiss`` 回调清标记，右上角 ✕ 即原生关闭按钮）。

视图分型（按检测结果定位能力）：
  docx  结构化块视图——正文按 body 块级顺序提取（标题分级/列表/表格行，
        Markdown 风格渲染），页眉/脚注等部件带标签 chip；
  xlsx  工作表行视图——表格化格子渲染（列字母 + 值），公式显示缓存值、
        数字按 number_format 美化、日期格式化；
  path  文件主名卡片（source="path"，词来自文件名而非正文）；
  miss  定位失败降级卡片（无法逐字匹配时展示检测时的上下文快照）。

文本保真（R11）：
  - 每个块保留两个口径：``text``（检测口径，w:t/canonical 拼接，用于定位
    匹配，保证 context 精确命中）与 ``disp``（显示口径，含 tab/软回车、
    公式缓存值、格式化数字），乱码/黏连/碎片问题由此消除；
  - 段落显示文本把 tab/软回车归一为空格，保证单行内连贯。

展示约束：命中行 ±10 行清晰，再向外各 3 行模糊提示尚有内容；行区
``overflow: hidden`` 不可滚动，超长行按显示宽度截断（命中行以敏感词为中心
截取，保证关键词可见）。弹窗只读，不改动文档与勾选状态。
"""
import html as _html
import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import streamlit as st

from mask_tool.models.detection import DetectionResult

# ±10 行清晰 + 两侧各 3 行模糊
CONTEXT_RANGE = 10
BLUR_PAD = 3
# 行显示宽度预算（显示宽度单位：CJK/全角=2，其余=1；≈ 大号弹窗正文可用宽度）
_LINE_BUDGET = 104
# 表格单元格：单值显示宽度上限与最多展示格数
_CELL_UNITS = 20
_MAX_CELLS = 8
# docx 标题样式名（本地化 Office 会用「标题 N」）
_HEADING_RE = re.compile(r"^(?:Heading|标题)\s*([1-9])")


# ──────────────────────────────────────────────
# AgGrid 单元格渲染器与点击消费
# ──────────────────────────────────────────────

def _ctx_cell_renderer():
    """「上下文」列 cellRenderer：可点击 span，点击写隐藏 loc 列触发回传。

    复用 AgGrid 既有 ``VALUE_CHANGED`` 通道（勾选列同链路），无需新增组件：
    点击 → ``node.setDataValue('loc', 递增序号)`` → 全表数据回传。
    序号必须递增（同列重复点击同一行也产生新事件）；300ms 防抖吸收双击。
    类式渲染器（init/getGui）绕过 React 子元素限制（st-aggrid 1.2.1 内置
    React 版 ag-grid，函数式渲染器返回 DOM 元素会触发 React error #31）。
    refresh 必须同步缓存 params：行组件复用且值未变时 ag-grid 不重挂渲染
    器，陈旧 node 会导致点击写错行。样式 ``mt-ctx-link`` 由 assets.py 注入。
    """
    from st_aggrid import JsCode

    return JsCode(
        """
        class MtCtxRenderer {
            init(params) {
                this.params = params;
                var el = document.createElement('span');
                el.className = 'mt-ctx-link';
                el.title = '点击定位该敏感词在文档中的位置';
                var zi = document.createElement('span');
                zi.className = 'zi';
                zi.textContent = '🔍';
                this.tx = document.createElement('span');
                this.tx.className = 'tx';
                this.tx.textContent = params.value || '';
                el.appendChild(zi);
                el.appendChild(this.tx);
                var self = this;
                el.addEventListener('click', function (ev) {
                    ev.stopPropagation();
                    var w = window;
                    var now = Date.now();
                    if (w.__mtLocAt && now - w.__mtLocAt < 300) { return; }
                    w.__mtLocAt = now;
                    w.__mtLocSeq = (w.__mtLocSeq || 0) + 1;
                    try { self.params.node.setDataValue('loc', w.__mtLocSeq); }
                    catch (e) { /* node 不可用时静默 */ }
                });
                this.eGui = el;
            }
            getGui() { return this.eGui; }
            refresh(params) {
                // 行组件复用时同步引用：否则点击会用旧行的 node（行错位）
                this.params = params;
                this.tx.textContent = params.value || '';
                return true;
            }
            destroy() { this.eGui = null; }
        }
        """
    )


def _consume_grid_locate_click(grid_response) -> Optional[int]:
    """从 AgGrid 回传数据消费「上下文」点击，返回检测结果全局索引。

    组件 return_value 在下一次事件前保持不变：同一序号只会消费一次
    （``last_locate_seq`` 会话守卫），避免弹窗被旧事件反复重开。
    注：st-aggrid 1.x 返回 AgGridReturn（Mapping，非 dict），
    取值统一走 ``.get()``，data 为 DataFrame 或 list[dict]。
    """
    try:
        data = grid_response.get("data")
    except Exception:
        return None
    if data is None:
        return None
    if hasattr(data, "to_dict"):  # pandas DataFrame
        rows = data.to_dict("records")
    elif isinstance(data, list):
        rows = [r for r in data if isinstance(r, dict)]
    else:
        return None
    best_idx, best_seq = None, 0
    for row in rows:
        try:
            seq = int(row.get("loc") or 0)
        except (TypeError, ValueError):
            continue
        if seq > best_seq:
            best_seq = seq
            try:
                idx = int(row.get("index", -1))
            except (TypeError, ValueError):
                idx = -1
            best_idx = idx if idx >= 0 else None
    last_seq = int(st.session_state.get("last_locate_seq") or 0)
    if best_idx is None or best_seq <= last_seq:
        return None
    st.session_state["last_locate_seq"] = best_seq
    import logging
    logging.getLogger("mask_tool").info(
        "R10 上下文定位：seq=%d → 检测索引=%d", best_seq, best_idx)
    return best_idx


# ──────────────────────────────────────────────
# 文档块提取（R11：结构 + 双口径文本；会话级缓存）
# ──────────────────────────────────────────────

def _cache():
    return st.session_state.setdefault("_locate_doc_cache", {})


def _norm_disp(text: str) -> str:
    """显示口径归一：tab/软回车 → 空格，连续空白压缩，文本连贯单行。"""
    text = text.replace("\r", "\n").replace("\t", " ")
    parts = [seg.strip() for seg in text.split("\n")]
    return " ".join(seg for seg in parts if seg)


def _wt_match_text(p_elem) -> str:
    """检测口径段落文本：w:t 直接拼接（与 DocxAdapter walker 完全一致，
    保证 context 切片能在此口径精确命中）。"""
    from mask_tool.adapters.docx_adapter import DocxAdapter

    return "".join(t.text or "" for t in DocxAdapter._paragraph_text_nodes(p_elem))


def _wt_display_text(p_elem) -> str:
    """显示口径段落文本：按文档顺序补 w:tab（空格）与 w:br/w:cr（软回车，
    归一为空格），消除检测口径丢失分隔符造成的词黏连。"""
    from docx.oxml.ns import qn

    out: List[str] = []
    for node in p_elem.iter():
        if node.tag == qn("w:t"):
            out.append(node.text or "")
        elif node.tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
            out.append(" ")
    return _norm_disp("".join(out))


def _heading_level(para) -> int:
    """段落标题级别（样式名 Heading N / 标题 N）；非标题返回 0。"""
    try:
        name = (para.style.name if para.style is not None else "") or ""
    except Exception:
        return 0
    m = _HEADING_RE.match(name.strip())
    return min(int(m.group(1)), 4) if m else 0


def _is_list_item(p_elem) -> bool:
    """列表段落判定：直接格式 w:numPr（项目符号/编号列表均归为列表）。"""
    from docx.oxml.ns import qn

    pPr = p_elem.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:numPr")) is not None


def _extract_docx_blocks(path: Path) -> List[Dict]:
    """docx 结构化块：正文 body 顺序（段落/表格行），其他部件平铺带标签。

    块结构（双口径）：
      {"kind": "h"|"li"|"para", "level", "text": 检测口径, "disp": 显示口径, "part"}
      {"kind": "thead"|"trow", "cells": [{"m": 检测口径, "d": 显示口径}],
       "text": cells.m 拼接（匹配用）, "part": "表格"}
    """
    cached = _cache().get(("docx2", str(path)))
    if cached is not None:
        return cached

    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: List[Dict] = []

    # ── 正文：body 子元素顺序遍历，表格保持结构 ──
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            m_text = _wt_match_text(child)
            if not m_text.strip():
                continue
            para = Paragraph(child, doc)
            level = _heading_level(para)
            kind = "h" if level else ("li" if _is_list_item(child) else "para")
            blocks.append({
                "kind": kind, "level": level,
                "text": m_text, "disp": _norm_disp(para.text) or m_text,
                "part": "",
            })
        elif child.tag == qn("w:tbl"):
            try:
                table = Table(child, doc)
            except Exception:
                continue
            first = True
            for row in table.rows:
                cells = []
                for c in row.cells:
                    m = _wt_cell_match_text(c)
                    if not m:
                        continue
                    cells.append({"m": m, "d": _norm_disp(c.text) or m})
                if not cells:
                    continue
                blocks.append({
                    "kind": "thead" if first else "trow",
                    "cells": cells,
                    "text": " ".join(c["m"] for c in cells),
                    "part": "表格",
                })
                first = False

    # ── 其他部件（页眉/页脚/脚注/批注/文本框）：walker 同源平铺带标签 ──
    from mask_tool.adapters.docx_adapter import DocxAdapter

    for label, root, _flush in DocxAdapter._iter_xml_roots(doc):
        if not label:  # document 已按结构化处理
            continue
        for p_elem in root.iter(qn("w:p")):
            m_text = _wt_match_text(p_elem)
            if not m_text.strip():
                continue
            blocks.append({
                "kind": "para", "level": 0,
                "text": m_text, "disp": _wt_display_text(p_elem) or m_text,
                "part": label,
            })

    _cache()[("docx2", str(path))] = blocks
    return blocks


def _wt_cell_match_text(cell) -> str:
    """表格单元格的检测口径文本：单元格内全部 w:t 拼接（与 walker 把
    单元格段落逐段送检的口径一致到“包含”关系——定位只需子串命中）。"""
    from docx.oxml.ns import qn

    tc = getattr(cell, "_tc", None)
    if tc is None:
        return ""
    return "".join(t.text or "" for t in tc.iter(qn("w:t")))


def _pretty_number(value, number_format: str) -> str:
    """数字按 number_format 常见形态美化（千分位/小数位/百分比/货币）。

    只处理常见形态；识别不了的回落 canonical（检测口径），保证词可高亮。
    """
    from mask_tool.adapters.xlsx_adapter import XlsxAdapter

    nf = number_format or ""
    try:
        body = nf.split(";")[0]
        if "%" in body:
            frac = body.split(".", 1)[1] if "." in body else ""
            dec = min(frac.count("0"), 4)
            return f"{value * 100:.{dec}f}%"
        frac = body.split(".", 1)[1] if "." in body else ""
        dec = min(frac.count("0"), 6)  # 0.00 → 2 位；无小数部分 → 0 位
        thousands = "#,##" in body or ",0" in body
        if dec:
            s = f"{value:,.{dec}f}" if thousands else f"{value:.{dec}f}"
        elif thousands:
            s = f"{value:,.0f}" if float(value).is_integer() else f"{value:,g}"
        else:
            return XlsxAdapter._numeric_canonical(value) or str(value)
        cur = ""
        for sym in ("¥", "￥", "$", "€"):
            if sym in body:
                cur = sym
                break
        return cur + s
    except Exception:
        return XlsxAdapter._numeric_canonical(value) or str(value)


def _extract_xlsx_sheets(path: Path) -> List[Dict]:
    """xlsx 各工作表非空行（结构化格子，双读显示缓存值）。

    每行 {"no": Excel 行号, "kind": "thead"|"trow",
          "cells": [{"coord", "m": 匹配口径, "d": 显示口径}], "hays": [(coord, text)]}。
    匹配口径 = 原始文本/富文本整串/数字 canonical（与检测面一致）；
    显示口径 = 公式缓存值（data_only 双读）、number_format 美化数字、
    格式化日期——两者都进 hays，保证 context 与显示形态均可命中。
    """
    cached = _cache().get(("xlsx2", str(path)))
    if cached is not None:
        return cached

    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.cell.rich_text import CellRichText

    from mask_tool.adapters.xlsx_adapter import XlsxAdapter

    # 第二遍：data_only 读公式缓存值（与 rich_text 互斥，单独加载）
    cached_vals: Dict[str, str] = {}
    try:
        wb2 = load_workbook(str(path), data_only=True, read_only=True)
        for ws2 in wb2.worksheets:
            for row in ws2.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is not None and not isinstance(v, bool):
                        cached_vals[f"{ws2.title}!{cell.coordinate}"] = v
        wb2.close()
    except Exception:
        cached_vals = {}

    wb = load_workbook(str(path), rich_text=True)
    sheets: List[Dict] = []
    try:
        for ws in wb.worksheets:
            rows: List[Dict] = []
            first_data_row = True
            for row in ws.iter_rows():
                cells: List[Dict] = []
                hays: List[Tuple[str, str]] = []
                row_no = None
                for cell in row:
                    if row_no is None:
                        row_no = cell.row
                    if isinstance(cell, MergedCell):
                        continue
                    v = cell.value
                    key = f"{ws.title}!{cell.coordinate}"
                    m = d = None
                    if isinstance(v, CellRichText):
                        s = str(v)
                        if s.strip():
                            m = d = s
                    elif isinstance(v, bool):
                        continue
                    elif cell.is_date:
                        try:
                            has_time = getattr(v, "hour", 0) or getattr(v, "minute", 0)
                            d = v.strftime("%Y-%m-%d %H:%M" if has_time else "%Y-%m-%d")
                            m = d
                        except Exception:
                            m = d = str(v)
                    elif isinstance(v, (int, float)):
                        m = XlsxAdapter._numeric_canonical(v)
                        d = _pretty_number(v, cell.number_format or "")
                    elif isinstance(v, str):
                        if not v.strip():
                            continue
                        m = v
                        if cell.data_type == "f" and key in cached_vals:
                            cv = cached_vals[key]
                            d = _pretty_number(cv, cell.number_format or "") \
                                if isinstance(cv, (int, float)) and not isinstance(cv, bool) \
                                else (str(cv) if cv is not None else v)
                        else:
                            d = v
                    if m is None:
                        continue
                    cells.append({"coord": cell.coordinate, "m": m, "d": d or m})
                    hays.append((cell.coordinate, m))
                    if d and d != m:
                        hays.append((cell.coordinate, d))
                if cells:
                    rows.append({
                        "no": row_no,
                        "kind": "thead" if first_data_row else "trow",
                        "cells": cells, "hays": hays,
                    })
                    first_data_row = False
            sheets.append({"sheet": ws.title, "rows": rows})
    finally:
        wb.close()
    _cache()[("xlsx2", str(path))] = sheets
    return sheets


# ──────────────────────────────────────────────
# 定位（先按上下文精确匹配，回退按敏感词匹配）
# ──────────────────────────────────────────────

def _find_target(blocks: List[Dict], r: DetectionResult) -> Optional[int]:
    """docx：返回命中块下标。context 是检测口径文本的连续子串，优先精确
    匹配；不足 8 字（极短文本部件）或未命中时回退按敏感词匹配。"""
    ctx = (r.context or "").strip()
    if len(ctx) >= 8:
        for i, b in enumerate(blocks):
            if ctx in b["text"]:
                return i
    for i, b in enumerate(blocks):
        if r.text in b["text"]:
            return i
    return None


def _find_xlsx_target(sheets: List[Dict], r: DetectionResult) -> Optional[Tuple[int, int, str]]:
    """xlsx：返回 (工作表下标, 行下标, 命中单元格坐标)。

    两遍扫描：先全簿按 context 精确匹配，再回退按敏感词匹配；
    hay 覆盖检测口径（canonical/原文）与显示口径（格式化形态）。
    """
    ctx = (r.context or "").strip()

    def _scan(needle: Optional[str]) -> Optional[Tuple[int, int, str]]:
        if not needle:
            return None
        for si, sh in enumerate(sheets):
            for ri, row in enumerate(sh["rows"]):
                for coord, hay in row["hays"]:
                    if needle in hay:
                        return si, ri, coord
        return None

    return _scan(ctx if len(ctx) >= 8 else None) or _scan(r.text)


# ──────────────────────────────────────────────
# 文本整形：显示宽度截断 + 高亮（全部 HTML 转义后再注入 mark）
# ──────────────────────────────────────────────

def _wch(ch: str) -> int:
    return 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1


def _units(s: str) -> int:
    return sum(_wch(c) for c in s)


def _clip_fwd(s: str, start: int, budget: int) -> str:
    """从 start 起截取至多 budget 显示宽度单位的子串。"""
    out: List[str] = []
    u = 0
    for ch in s[start:]:
        cu = _wch(ch)
        if u + cu > budget:
            break
        out.append(ch)
        u += cu
    return "".join(out)


def _clip_back(s: str, end: int, budget: int) -> str:
    """截取 s[:end] 末尾至多 budget 显示宽度单位的子串。"""
    u = 0
    start = end
    while start > 0:
        cu = _wch(s[start - 1])
        if u + cu > budget:
            break
        u += cu
        start -= 1
    return s[start:end]


def _trim_line(text: str, word: Optional[str], budget: int = _LINE_BUDGET) -> str:
    """超长行截断；命中行（提供 word 且存在于行内）以敏感词为中心截取。"""
    if _units(text) <= budget:
        return text
    if word and word in text:
        i = text.index(word)
        w_u = _units(word)
        side = max(8, (budget - w_u) // 2 - 1)
        left = _clip_back(text, i, side)
        right = _clip_fwd(text, i + len(word), side)
        out = left + word + right
        if left:
            out = "…" + out
        if i + len(word) + len(right) < len(text):
            out += "…"
        return out
    return _clip_fwd(text, 0, budget - 1) + "…"


def _hl(line: str, word: str) -> str:
    """HTML 转义后高亮全部出现（split/join 保证不破坏实体边界）。"""
    if not word:
        return _html.escape(line)
    marked = '<mark class="w">' + _html.escape(word) + "</mark>"
    return marked.join(_html.escape(p) for p in line.split(word))


def _clip_cell(s: str, budget: int = _CELL_UNITS) -> str:
    """表格格子值截断（单行省略）。"""
    if _units(s) <= budget:
        return s
    return _clip_fwd(s, 0, budget - 1) + "…"


# ──────────────────────────────────────────────
# 视图构建与行渲染（R11：Markdown 风格结构）
# ──────────────────────────────────────────────

def _meta_chips(r: DetectionResult, loc_text: str) -> str:
    from .labels import SOURCE_LABELS, TYPE_LABELS

    chips = [
        f'<span class="mchip">{_html.escape(TYPE_LABELS.get(r.text_type, r.text_type.value))}</span>',
        f'<span class="mchip">{_html.escape(str(SOURCE_LABELS.get(r.source, r.source)))}</span>',
    ]
    if r.location and r.location.file:
        chips.append(
            '<span class="mchip">'
            + _html.escape(Path(r.location.file).name) + "</span>"
        )
    chips.append(f'<span class="mchip loc">{_html.escape(loc_text)}</span>')
    return "".join(chips)


def _line_shell(no, inner: str, hit: bool, blur: bool, part: str = "") -> str:
    return (
        f'<div class="cl-line{" hit" if hit else ""}{" blur" if blur else ""}">'
        f'<span class="no">{no}</span>'
        + (f'<span class="part">{_html.escape(part)}</span>' if part else "")
        + inner + "</div>"
    )


def _lines_html_docx(blocks: List[Dict], target: int, word: str) -> str:
    out: List[str] = []
    lo = max(0, target - (CONTEXT_RANGE + BLUR_PAD))
    hi = min(len(blocks) - 1, target + (CONTEXT_RANGE + BLUR_PAD))
    for i in range(lo, hi + 1):
        b = blocks[i]
        d = abs(i - target)
        is_hit = i == target
        kind = b["kind"]
        if kind in ("thead", "trow"):
            cells = b["cells"][:_MAX_CELLS]
            inner_cells = []
            for c in cells:
                v = _clip_cell(c["d"])
                tv = _hl(v, word) if (word and word in v) else _html.escape(v)
                inner_cells.append(f'<span class="tv">{tv}</span>')
            row = f'<span class="tcells{" thead-c" if kind == "thead" else ""}">{" ".join(inner_cells)}</span>'
            if len(b["cells"]) > _MAX_CELLS:
                row += '<span class="tmore">…</span>'
            out.append(_line_shell(i + 1, row, is_hit, d > CONTEXT_RANGE,
                                   "表格" if kind == "thead" else ""))
        else:
            disp = b["disp"]
            if kind == "li":
                disp = "• " + disp
            txt = _trim_line(disp, word if is_hit else None)
            cls = "tx"
            if kind == "h":
                cls += f" h h{b['level']}"
            elif kind == "li":
                cls += " li"
            out.append(_line_shell(i + 1, f'<span class="{cls}">{_hl(txt, word)}</span>',
                                   is_hit, d > CONTEXT_RANGE, b["part"]))
    return "".join(out)


def _lines_html_xlsx(rows: List[Dict], target_ri: int, hit_coord: str, word: str) -> str:
    out: List[str] = []
    lo = max(0, target_ri - (CONTEXT_RANGE + BLUR_PAD))
    hi = min(len(rows) - 1, target_ri + (CONTEXT_RANGE + BLUR_PAD))
    for ri in range(lo, hi + 1):
        row = rows[ri]
        d = abs(ri - target_ri)
        is_hit = ri == target_ri
        cells = row["cells"][:_MAX_CELLS]
        inner: List[str] = []
        for c in cells:
            col = "".join(ch for ch in c["coord"] if ch.isalpha())
            v = _clip_cell(c["d"])
            hit_cell = is_hit and c["coord"] == hit_coord
            if word and word in v:
                tv = _hl(v, word)
            elif hit_cell:
                # 千分位/格式化形态与检测词不同时，命中格整格标黄
                tv = f'<mark class="cellw">{_html.escape(v)}</mark>'
            else:
                tv = _html.escape(v)
            inner.append(
                f'<span class="tc{" hitc" if hit_cell else ""}">'
                f'<i class="col">{_html.escape(col)}</i>'
                f'<span class="tv">{tv}</span></span>'
            )
        row_html = f'<span class="tcells x{" thead-c" if row["kind"] == "thead" else ""}">{" ".join(inner)}</span>'
        if len(row["cells"]) > _MAX_CELLS:
            row_html += '<span class="tmore">…</span>'
        out.append(_line_shell(row["no"], row_html, is_hit, d > CONTEXT_RANGE, ""))
    return "".join(out)


def _build_locate_view(r: DetectionResult) -> Dict:
    """构建定位视图（docx/xlsx/path/miss 四型），供弹窗渲染。"""
    file_path = (r.location.file if r.location else None) or ""

    # 文件名来源：词不在文档正文中
    if r.source == "path" or not file_path:
        return {"kind": "path", "word": r.text, "file": file_path}

    path = Path(file_path)
    if not path.exists():  # tmp 目录已被清理（理论不可达：检测阶段始终存在）
        return {"kind": "miss", "word": r.text, "file": file_path, "r": r}

    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            blocks = _extract_docx_blocks(path)
            target = _find_target(blocks, r)
            if target is None:
                return {"kind": "miss", "word": r.text, "file": file_path, "r": r}
            b = blocks[target]
            where = b["part"] or "正文"
            detail = f"第 {target + 1} 行 · {where}"
            if b["kind"] == "h":
                detail += f"（{b['level']} 级标题）"
            elif b["kind"] in ("thead", "trow"):
                detail += "（表格行）"
            return {
                "kind": "docx", "word": r.text, "file": file_path,
                "loc_text": detail,
                "lines_html": _lines_html_docx(blocks, target, r.text),
            }
        if suffix == ".xlsx":
            sheets = _extract_xlsx_sheets(path)
            target = _find_xlsx_target(sheets, r)
            if target is None:
                return {"kind": "miss", "word": r.text, "file": file_path, "r": r}
            si, ri, coord = target
            sh = sheets[si]
            return {
                "kind": "xlsx", "word": r.text, "file": file_path,
                "loc_text": f"工作表「{sh['sheet']}」第 {sh['rows'][ri]['no']} 行 · 单元格 {coord}",
                "lines_html": _lines_html_xlsx(sh["rows"], ri, coord, r.text),
            }
    except Exception:
        return {"kind": "miss", "word": r.text, "file": file_path, "r": r}

    return {"kind": "miss", "word": r.text, "file": file_path, "r": r}


# ──────────────────────────────────────────────
# 弹窗（st.dialog：右上角 ✕ 原生关闭；on_dismiss 清标记防重开）
# ──────────────────────────────────────────────

def _dismiss_locate(*_args, **_kwargs) -> None:
    st.session_state.pop("ctx_locate_idx", None)


@st.dialog("📍 敏感词原文定位", width="large", on_dismiss=_dismiss_locate)
def _context_locate_dialog():
    """定位弹窗本体：meta 行 + ±10 行结构化视图（或降级卡片）。只读。"""
    idx = st.session_state.get("ctx_locate_idx")
    results = st.session_state.get("detection_results") or []
    if idx is None or not (0 <= idx < len(results)):
        st.session_state.pop("ctx_locate_idx", None)
        return
    r = results[idx]
    view = _build_locate_view(r)
    word = view["word"]

    meta = (
        f'<div class="cl-meta"><div class="cl-meta-row">'
        f'<span class="cl-word">{_html.escape(word)}</span>'
        + _meta_chips(r, view.get("loc_text", ""))
        + '</div><div class="cl-note">仅用于定位查看，不会修改文档 · '
          '显示敏感词上下各 10 行，更远处已模糊</div></div>'
    )

    if view["kind"] in ("docx", "xlsx"):
        st.markdown(
            meta + f'<div class="cl-lines">{view["lines_html"]}</div>',
            unsafe_allow_html=True,
        )
    elif view["kind"] == "path":
        name = Path(view["file"]).name if view["file"] else ""
        stem = Path(name).stem or name
        ext = Path(name).suffix
        st.markdown(
            meta
            + '<div class="cl-card"><div class="cl-f">📄 '
            + _hl(stem, word) + _html.escape(ext) + "</div>"
            + '<div class="cl-d">该敏感词来自<b>文件名</b>（非文档正文）：'
              "勾选执行后将对文件主名执行同套替换，扩展名保持不变。</div></div>",
            unsafe_allow_html=True,
        )
    else:  # miss：无法逐字定位（格式化数字/公式字面量/富文本块等）
        ctx = (r.context or "").strip() or "（无上下文快照）"
        st.markdown(
            meta
            + '<div class="cl-card miss"><div class="cl-d">'
            "未能按原文逐字定位（该词可能来自格式化数字、公式字面量或富文本"
            "块等特殊位置）。以下是检测时的上下文快照：</div>"
            f'<div class="cl-snap">{_hl(_trim_line(ctx, word, 3 * _LINE_BUDGET), word)}</div>'
            "</div>",
            unsafe_allow_html=True,
        )
