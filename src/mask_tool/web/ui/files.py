# -*- coding: utf-8 -*-
"""文件级工具：目录 zip 安全解压/分类、文档文本抽取、图标与置信度分级。"""
import os
import shutil
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Dict, List

from .labels import BLOCKED_EXTS, SUPPORTED_MASK_EXTS


# 目录 zip 上传防护（I6 问题2）：成员数与解压总大小上限（zip 炸弹防御）
ZIP_MAX_ENTRIES = 500
ZIP_MAX_TOTAL_BYTES = 500 * 1024 * 1024  # 500MB

# 未声明 UTF-8 的成员名依次尝试的回退解码顺序（首个成功者生效）。
# GB18030 为 GBK 超集，覆盖简体中文全部字符；顺序不可换（UTF-8 严格
# 校验通过时优先级更高，GBK 串通过严格 UTF-8 校验的概率可忽略）。
_ZIP_NAME_FALLBACK_ENCODINGS = ("utf-8", "gb18030")


def _decode_zip_member_name(info: zipfile.ZipInfo) -> str:
    """还原 zip 成员文件名的真实编码（GBK 名被 zipfile 按 CP437 误解码）。

    ZIP 规范允许成员名不带编码声明：中文 Windows 常见压缩工具
    （资源管理器、WinRAR 等）以本地 ANSI 代码页（GBK/CP936）写入
    文件名，且不设 UTF-8 标志（通用标志位 bit 0x800）；Python
    zipfile 对未设标志的名字一律按 CP437 解码，中文文件名自此变成
    "─Ω╢╚"式乱码，并随解压路径扩散到检测清单、确认弹窗与产物。

    恢复策略（仅对未声明 UTF-8 的名字生效）：
    - 名字按 CP437 回转成原始字节（zipfile 以 CP437 严格解码，
      回转必然无损；若失败则解码信息已损，保留原名）
    - 纯 ASCII 字节无编码歧义，直接原名返回
    - 按 _ZIP_NAME_FALLBACK_ENCODINGS 顺序尝试解码，失败回退原名
    """
    if info.flag_bits & 0x800:  # 已声明 UTF-8：zipfile 已正确解码
        return info.filename
    try:
        raw = info.filename.encode("cp437")
    except UnicodeEncodeError:  # 非严格 CP437 字符集：无法回转
        return info.filename
    if all(b < 0x80 for b in raw):
        return info.filename
    for enc in _ZIP_NAME_FALLBACK_ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return info.filename


def _safe_unzip(zip_bytes: bytes, dest: Path) -> Path:
    """安全解压目录 zip 到 dest/tree，返回解压根目录。

    防护（恶意/意外 zip 拒绝，抛 ValueError 由调用万 st.error）：
    - 路径穿越：成员路径 normpath 后必须仍在解压根内
      （拒绝 ../、绝对路径、盘符等逃逸形态）
    - zip 炸弹：成员数 ≤ 500、声明解压总大小 ≤ 500MB
    """
    root = dest / "tree"
    with zipfile.ZipFile(BytesIO(zip_bytes)) as zf:
        infos = zf.infolist()
        if len(infos) > ZIP_MAX_ENTRIES:
            raise ValueError(
                f"压缩包成员数 {len(infos)} 超过上限 {ZIP_MAX_ENTRIES}，已拒绝解压"
            )
        total = sum(i.file_size for i in infos)
        if total > ZIP_MAX_TOTAL_BYTES:
            raise ValueError(
                f"压缩包声明解压总大小 {total / 1024 / 1024:.1f}MB 超过上限 "
                f"{ZIP_MAX_TOTAL_BYTES // 1024 // 1024}MB，已拒绝解压"
            )
        root.mkdir(parents=True, exist_ok=True)
        base = root.resolve()
        for info in infos:
            name = _decode_zip_member_name(info)
            # 归一化后必须仍在解压根内（拒绝 ../ 与绝对路径穿越）
            member = Path(os.path.normpath(str(root / name)))
            try:
                member.relative_to(base)
            except ValueError:
                raise ValueError(
                    f"压缩包成员路径越界（疑似路径穿越），已拒绝解压: {name!r}"
                )
            if info.is_dir():
                member.mkdir(parents=True, exist_ok=True)
                continue
            member.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, open(member, "wb") as dst:
                shutil.copyfileobj(src, dst)
    return root


def _xlsm_macro_warnings(paths) -> List[str]:
    """对含 VBA 宏的 xlsm 输入生成 UI 警告（对抗审查 P1-2）。

    adapter 侧的 logger 告警在桌面版（pythonw 无控制台）不可见；
    检测阶段即向用户明示"宏代码内容不参与脱敏"边界。
    """
    import zipfile

    macro_files = []
    for p in paths:
        try:
            if Path(p).suffix.lower() != ".xlsm":
                continue
            with zipfile.ZipFile(str(p)) as z:
                if "xl/vbaProject.bin" in z.namelist():
                    macro_files.append(Path(p).name)
        except Exception:
            continue
    if macro_files:
        return [
            "以下带宏文件（" + "、".join(macro_files) + "）的宏代码内容不参与"
            "脱敏（宏部件原样保留），请人工确认宏内无敏感信息"
        ]
    return []


def _classify_tree_files(tree_root: Path) -> Dict[str, List[Path]]:
    """遍历解压根，按处理面分类：docs（docx/xlsx/xlsm/pdf 待脱敏）、
    blocked（pptx 等屏蔽类型，警告且不进产物）、others（其余原样拷入）。"""
    docs: List[Path] = []
    blocked: List[Path] = []
    others: List[Path] = []
    for p in sorted(tree_root.rglob("*")):
        if not p.is_file():
            continue
        suffix = p.suffix.lower()
        if suffix in SUPPORTED_MASK_EXTS:
            docs.append(p)
        elif suffix in BLOCKED_EXTS:
            blocked.append(p)
        else:
            others.append(p)
    return {"docs": docs, "blocked": blocked, "others": others}


def _extract_text(file_path: Path) -> str:
    """从文件中提取纯文本（R1-A1：委托 adapters/extract 公共实现，
    与 CLI/adapter 处理面同源——docx walker 全部件、xlsx 全部件，
    消除旧版 doc.paragraphs+tables / read_only str 单元格的检测盲区）。"""
    suffix = file_path.suffix.lower()
    try:
        if suffix in (".docx", ".xlsx", ".xlsm"):
            from mask_tool.adapters.extract import extract_texts
            return extract_texts(file_path)
        if suffix == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(file_path))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            for cell in row.cells:
                                texts.append(cell.text)
            return "\n".join(texts)
        if suffix == ".pdf":
            # 与 adapter 处理面同源（块级提取；extract_texts 内含加密守卫）
            from mask_tool.adapters.extract import extract_texts
            return extract_texts(file_path)
    except Exception:
        return ""
    return ""


def _file_icon(suffix: str) -> str:
    """返回文件类型图标"""
    icons = {
        ".docx": "📄", ".xlsx": "📊", ".pptx": "📽️", ".pdf": "📕",
    }
    return icons.get(suffix.lower(), "📁")


def _confidence_class(confidence: float) -> str:
    """返回置信度对应的CSS类名"""
    if confidence >= 0.85:
        return "confidence-high"
    elif confidence >= 0.60:
        return "confidence-mid"
    return "confidence-low"


